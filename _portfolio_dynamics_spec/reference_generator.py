# -*- coding: utf-8 -*-
"""
Генератор шаблона portfolio_limits_template.xlsx
Схема v3.0 — промежуточный слой между ручным вводом и BI.

Отличия от v2.0:
  - лист dim_type удалён; реестром типов служит fact_limit (одна строка на тип);
  - dim_portfolio: добавлен include_in_total, убран is_active;
  - fact_limit: грейн portfolio_type — лимит ставится на СУММАРНЫЙ объём типа;
  - fact_portfolio_snapshot: добавлена колонка note_text (ручная);
  - лист ref_note удалён;
  - view_monitor заканчивается колонкой «Лимит»; светофор живёт в view_by_type,
    где лимит и объём находятся на одном грейне.
"""
import datetime as dt
import random

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName

OUT = "/home/claude/work/portfolio_limits_template.xlsx"

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_CALC = Font(name=FONT, size=10, color="000000")
F_LINK = Font(name=FONT, size=10, color="008000")
F_H = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=13, bold=True)
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_BOLD = Font(name=FONT, size=10, bold=True)

FILL_MACHINE = PatternFill("solid", fgColor="1F3864")
FILL_MANUAL = PatternFill("solid", fgColor="7F6000")
FILL_VIEW = PatternFill("solid", fgColor="385723")
FILL_INPUTCELL = PatternFill("solid", fgColor="FFF2CC")
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")
FILL_RED = PatternFill("solid", fgColor="FFC7CE")
FILL_BREACH = PatternFill("solid", fgColor="FF9999")
FILL_TOTAL = PatternFill("solid", fgColor="EDEDED")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_AMT = "#,##0;(#,##0);-"
FMT_PCT = "0.0%;(0.0%);-"
FMT_PCT3 = "0.00%;(0.00%);-"
FMT_DUR = "0.00"
FMT_DATE = "YYYY-MM-DD"

UNIT = "млн RUB"
TOL = 0.005

SCHEMA = {
    "dim_portfolio": [
        ("portfolio_code", "Код портфеля", "text", "PK", "UPPER, A-Z0-9_", "Технический код детального портфеля. Ключ связи с fact_portfolio_snapshot."),
        ("portfolio_name", "Наименование портфеля", "text", "да", "", "Человекочитаемое название для отчётов и BI."),
        ("portfolio_type", "Тип портфеля", "text", "да", "FK -> fact_limit.portfolio_type", "Тип, к которому относится портфель. Лимит ставится на тип, а не на портфель."),
        ("include_in_total", "Входит в итог", "bool", "да", "TRUE/FALSE", "TRUE — объём портфеля включается в итог по банку. FALSE для портфелей, вложенных в другой, чтобы не задваивать."),
        ("is_limit_controlled", "Под лимитом", "bool", "да", "TRUE/FALSE", "TRUE — портфель учитывается в утилизации лимита своего типа."),
        ("sort_order", "Порядок", "int", "да", ">0", "Порядок вывода в витрине."),
    ],
    "fact_limit": [
        ("portfolio_type", "Тип портфеля", "text", "PK", "UPPER, A-Z0-9_", "Тип портфеля. Лимит установлен на СУММАРНЫЙ объём всех портфелей этого типа. Одна строка на тип."),
        ("limit_amount", "Установленный лимит", "num", "да", ">= 0", "Установленный лимит на суммарный объём типа, %s." % UNIT),
        ("green_max_util", "Зелёная зона, до", "num", "да", ">= 0, < yellow_max_util", "Верхняя граница зелёной зоны, АБСОЛЮТНАЯ сумма в %s." % UNIT),
        ("yellow_max_util", "Жёлтая зона, до", "num", "да", "> green_max_util, < red_max_util", "Верхняя граница жёлтой зоны, АБСОЛЮТНАЯ сумма в %s." % UNIT),
        ("red_max_util", "Красная зона, до", "num", "да", "> yellow_max_util, обычно = limit_amount", "Верхняя граница красной зоны, АБСОЛЮТНАЯ сумма в %s. Выше — превышение лимита." % UNIT),
        ("valid_from", "Действует с", "date", "да", "ISO дата", "Дата вступления лимита в силу — аудиторский след."),
        ("updated_by", "Кем изменено", "text", "да", "", "ФИО/логин сотрудника, изменившего строку."),
    ],
    "fact_type_daily": [
        ("business_date", "Дата", "date", "PK", "ISO дата, рабочий день", "Отчётная дата. Вместе с portfolio_type образует составной ключ."),
        ("portfolio_type", "Тип портфеля", "text", "PK", "FK -> fact_limit.portfolio_type", "Тип портфеля. Детализации до портфелей в этой таблице нет и быть не должно."),
        ("volume_amount", "Объём типа", "num", "да", ">= 0", "Суммарный объём всех портфелей этого типа на дату, %s." % UNIT),
    ],
    "fact_portfolio_snapshot": [
        ("business_date", "Дата среза", "date", "да", "ISO дата, одна на весь лист", "Дата, на которую построен срез. Одинакова во всех строках."),
        ("portfolio_code", "Код портфеля", "text", "PK", "FK -> dim_portfolio", "Ключ детального портфеля. Одна строка на портфель."),
        ("volume_t0", "Объём сегодня", "num", "да", ">= 0", "Объём портфеля на дату среза, %s." % UNIT),
        ("volume_t7", "Объём T-7", "num", "нет", ">= 0", "Объём портфеля на дату среза минус lookback (по умолчанию 7 календарных дней), %s. Приходит из выгрузки, а не считается — истории по портфелям нет." % UNIT),
        ("duration_current_yrs", "Текущая дюрация", "num", "нет", ">= 0", "Фактическая дюрация портфеля на дату среза, лет."),
        ("duration_target_yrs", "Целевая дюрация", "num", "нет", ">= 0", "Целевая (плановая) дюрация портфеля, лет."),
        ("note_text", "Цель / заметка", "text", "нет", "до 1000 символов, ЗАПОЛНЯЕТСЯ ВРУЧНУЮ", "Цель или комментарий по портфелю. ВНИМАНИЕ: единственная ручная колонка на машинном листе — загрузчик обязан вычитать её перед перезаписью и вернуть на место по portfolio_code."),
    ],
}

TABLE_KIND = {
    "dim_portfolio": "manual",
    "fact_limit": "manual",
    "fact_type_daily": "machine",
    "fact_portfolio_snapshot": "machine",
}

# тип: (лимит на суммарный объём типа)
TYPE_LIMITS = [
    ("AFS", 125_000),
    ("HTM", 320_000),
    ("HTM_KUAP", 130_000),
    ("TTS", 70_000),
]
TYPES = [t for t, _ in TYPE_LIMITS]

# code, name, type, include_in_total, limit_controlled, order, целевой объём, дюр.тек, дюр.цель, заметка
PORTFOLIOS = [
    ("AFS_OFZ_PD", "ОФЗ-ПД, портфель до продажи", "AFS", True, True, 10, 37_200, 3.10, 3.00,
     "Держим объём около 37 млрд, дюрацию тянем к 3.0 за счёт коротких выпусков."),
    ("AFS_CORP_RUB", "Корпоративные облигации RUB, AFS", "AFS", True, True, 20, 35_100, 2.40, 2.50,
     "Новые покупки только в обмен на погашения — тип близко к жёлтой зоне."),
    ("AFS_MUNI", "Субфедеральные и муниципальные, AFS", "AFS", True, True, 30, 11_000, 3.80, 3.50,
     "Портфель законсервирован до пересмотра лимита в ноябре."),
    ("HTM_SMTH_AMT", "Амортизируемые выпуски, HTM", "HTM", True, True, 40, 142_500, 6.20, 5.50,
     "Основной вклад в красную зону типа HTM. Готовим предложение на КУАП: расширение лимита либо продажа выпусков с погашением в 2027."),
    ("HTM_ASH_ASH", "Структурные выпуски ASH, HTM", "HTM", True, True, 50, 74_700, 4.90, 5.00,
     "Целевая дюрация 5.0; структурные выпуски не наращиваем."),
    ("HTM_OFZ_LONG", "ОФЗ длинные, HTM", "HTM", True, True, 60, 79_200, 7.40, 7.00,
     "Набор длины приостановлен до разгрузки типа HTM."),
    ("HTM_KUAP_CORE", "Ядро портфеля КУАП", "HTM_KUAP", True, True, 70, 70_400, 6.10, 6.00,
     "Решение КУАП от 28.08: лимит типа не увеличиваем до пересмотра в ноябре."),
    ("HTM_KUAP_EXT", "Расширение портфеля КУАП", "HTM_KUAP", True, True, 80, 18_000, 5.30, 5.50,
     "Резервный объём под сделки КУАП, использован меньше чем наполовину."),
    ("TTS_OFZ", "ОФЗ, торговый портфель", "TTS", True, True, 90, 25_200, 1.10, 1.20,
     "Держим запас под аукционы Минфина."),
    ("TTS_CORP", "Корпоративные облигации, торговый", "TTS", True, True, 100, 12_000, 0.90, 1.00,
     "Оборачиваемость в норме, ограничений нет."),
    ("TTS_FX", "Валютные облигации, торговый", "TTS", True, True, 110, 13_650, 1.60, 1.50,
     "До конца недели сокращаем позицию — тип в жёлтой зоне."),
]

START_DATE = dt.date(2026, 1, 1)
BUSINESS_DATE = dt.date(2026, 9, 1)
LOOKBACK = 7

MAXROW = 30000
SNAPROW = 500
LIMROW = 60


def business_days(d0, d1):
    out, d = [], d0
    while d <= d1:
        if d.weekday() < 5:
            out.append(d)
        d += dt.timedelta(days=1)
    return out


DATES = business_days(START_DATE, BUSINESS_DATE)
IDX = {d: i for i, d in enumerate(DATES)}


def gen_paths():
    random.seed(20260901)
    paths, dpaths = {}, {}
    for code, name, typ, inc, ctl, order, target, dc0, dt0, note in PORTFOLIOS:
        p = [float(target)]
        for _ in range(len(DATES) - 1):
            p.append(max(1.0, p[-1] / (1 + random.gauss(0.0004, 0.005))))
        p.reverse()
        d = [dc0]
        for _ in range(len(DATES) - 1):
            d.append(max(0.1, d[-1] - random.gauss(0.0008, 0.02)))
        d.reverse()
        paths[code] = [round(x, 1) for x in p]
        dpaths[code] = [round(x, 2) for x in d]
    return paths, dpaths


PATHS, DPATHS = gen_paths()

TYPE_FACTS = []
for d in DATES:
    i = IDX[d]
    for t in TYPES:
        s = sum(PATHS[p[0]][i] for p in PORTFOLIOS if p[2] == t)
        TYPE_FACTS.append([d, t, round(s, 1)])

i0 = IDX[BUSINESS_DATE]
d7 = BUSINESS_DATE - dt.timedelta(days=LOOKBACK)
i7 = IDX[d7]
SNAPSHOT = [[BUSINESS_DATE, p[0], PATHS[p[0]][i0], PATHS[p[0]][i7],
             DPATHS[p[0]][i0], p[8], p[9]] for p in PORTFOLIOS]

wb = Workbook()

# ============================================================ README
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
readme = [
    ("Портфельные лимиты и динамика — шаблон обмена v3.0", "title"),
    ("Промежуточный слой между ручным вводом и BI. Все суммы — %s." % UNIT, "sub"),
    ("", None),
    ("Два грейна — это главное, что нужно понять про файл", "h"),
    ("Лимит ставится на ТИП портфеля: это ограничение на суммарный объём всех портфелей этого типа.", None),
    ("fact_type_daily — история объёмов по типам: дата, тип, объём. Именно с ней сравнивается лимит.", None),
    ("fact_portfolio_snapshot — срез на отчётную дату по каждому ДЕТАЛЬНОМУ портфелю: объём сегодня, объём T-7, дюрации, заметка.", None),
    ("Ежедневная история по детальным портфелям не хранится: её выгрузка требовала бы запроса на каждый день за весь период и неоправданна.", None),
    ("Поэтому объём T-7 по портфелю — НЕ вычисляемая величина, а колонка, которая приходит из выгрузки вместе со срезом.", None),
    ("", None),
    ("Листы", "h"),
    ("dim_portfolio — справочник детальных портфелей и их привязка к типам.", None),
    ("fact_limit — лимиты и границы зон ПО ТИПАМ. Одна строка на тип; этот же лист служит реестром допустимых типов.", None),
    ("fact_type_daily — история объёмов по типам.", None),
    ("fact_portfolio_snapshot — срез по портфелям на отчётную дату.", None),
    ("view_monitor — витрина по портфелям: объём T0/T-7, дельты, дюрация, лимит типа.", None),
    ("view_by_type — свод по типам: динамика, утилизация, светофор, сверка с детальным срезом.", None),
    ("checks — автоматические проверки качества. Все строки должны быть OK перед отправкой в BI.", None),
    ("dict — словарь данных: типы, обязательность, правила, грейн. Машиночитаемый контракт.", None),
    ("_lists — служебные списки для выпадающих значений. Не редактировать без причины.", None),
    ("", None),
    ("Заметки живут в срезе — и это требует внимания", "h"),
    ("Колонка note_text на листе fact_portfolio_snapshot — единственная ручная колонка на машинном листе.", None),
    ("Полная перезапись листа сотрёт заметки, если загрузчик их не сохранит.", None),
    ("Правило для пайплайна: перед перезаписью прочитать note_text, сматчить по portfolio_code и вернуть значения на место.", None),
    ("Пустые заметки считает CHK_15 — если их число внезапно выросло, скорее всего забыли этот шаг.", None),
    ("", None),
    ("Сверка двух грейнов", "h"),
    ("Сумма объёмов портфелей по типу (из среза) и объём типа на ту же дату (из истории) приходят из разных выгрузок.", None),
    ("Колонка «Расхождение» в view_by_type показывает разницу, проверка CHK_16 не пропускает расхождение больше 0.5%.", None),
    ("Это единственное место, где два источника встречаются, — и единственное, где может вылезти рассинхрон выгрузок.", None),
    ("", None),
    ("Кто что пишет", "h"),
    ("Жёлтая шапка = ручной ввод. Код эти листы читает, но никогда не перезаписывает.", None),
    ("Синяя шапка = машинный лист. Код перезаписывает его целиком; исключение — колонка note_text (жёлтая шапка на синем листе).", None),
    ("Зелёная шапка = расчётная витрина на формулах. Не редактировать, только смотреть.", None),
    ("Цвет текста в ячейке: синий — введено руками, чёрный — формула, зелёный — ссылка на другой лист.", None),
    ("", None),
    ("Отчётная дата", "h"),
    ("Управляющие ячейки на листе view_monitor: жёлтая B3 — отчётная дата, D3 — сдвиг сравнения в календарных днях (по умолчанию 7).", None),
    ("B3 должна совпадать с датой среза в fact_portfolio_snapshot — это проверяет CHK_09.", None),
    ("", None),
    ("Зоны лимитов", "h"),
    ("Границы зон в fact_limit заданы АБСОЛЮТНЫМИ суммами: green_max_util < yellow_max_util < red_max_util.", None),
    ("Сравнивается с ними объём ТИПА из fact_type_daily: <= green — зелёная; <= yellow — жёлтая; <= red — красная; выше — ПРЕВЫШЕНИЕ.", None),
    ("Светофор находится на листе view_by_type, потому что лимит и объём типа лежат на одном грейне. В view_monitor колонка «Лимит» показывает лимит типа, к которому относится портфель.", None),
    ("", None),
    ("Правила, которые нельзя нарушать", "h"),
    ("1. Строка 1 каждого листа с данными — технические заголовки snake_case. Их читает код. Не переименовывать, не переставлять, не добавлять строку над ними.", None),
    ("2. В fact_type_daily и fact_limit — только типы. В fact_portfolio_snapshot и dim_portfolio — только детальные портфели.", None),
    ("3. Новый тип заводится строкой в fact_limit — она же открывает тип для выпадающих списков. Новый портфель — строкой в dim_portfolio.", None),
    ("4. Никаких объединённых ячеек, пустых строк-разделителей и строк «Итого» внутри таблиц.", None),
    ("5. Даты — настоящие даты Excel в формате YYYY-MM-DD. Не текст.", None),
    ("6. Числа — числа. Без пробелов-разделителей, без «млн» и знака валюты внутри ячейки. Единица одна на весь файл: %s." % UNIT, None),
    ("7. Пустая ячейка означает «нет данных». Не ставить прочерки, «н/д», «-» и нули вместо пропуска.", None),
    ("", None),
    ("Порядок работы", "h"),
    ("Шаг 1. Пайплайн дописывает свежие даты в fact_type_daily и перезаписывает fact_portfolio_snapshot, сохранив note_text.", None),
    ("Шаг 2. Казначейство ставит отчётную дату в view_monitor!B3, смотрит светофор на view_by_type, правит заметки и fact_limit при изменении лимитов.", None),
    ("Шаг 3. Лист checks должен быть весь OK. Любой FAIL блокирует загрузку в BI.", None),
    ("Шаг 4. Загрузчик читает dim_portfolio, fact_limit, fact_type_daily, fact_portfolio_snapshot. Витрины и checks в BI не грузятся.", None),
    ("", None),
    ("Внимание: файл заполнен демонстрационными данными за 01.01.2026–01.09.2026. Перед боевым запуском очистить строки данных.", "warn"),
]
r = 1
for text, kind in readme:
    c = ws.cell(row=r, column=1, value=text)
    if kind == "title":
        c.font = F_TITLE
    elif kind == "sub":
        c.font = F_SUB
    elif kind == "h":
        c.font = Font(name=FONT, size=11, bold=True, color="1F3864")
    elif kind == "warn":
        c.font = Font(name=FONT, size=10, bold=True, color="C00000")
    else:
        c.font = F_BASE
    c.alignment = Alignment(vertical="center")
    r += 1
ws.column_dimensions["A"].width = 165

# ============================================================ _lists
ws = wb.create_sheet("_lists")
for j, h in enumerate(["bool_value"], start=1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = F_H
    c.fill = FILL_MACHINE
for i, v in enumerate(["TRUE", "FALSE"], start=2):
    ws.cell(row=i, column=1, value=v).font = F_BASE
ws.column_dimensions["A"].width = 18
ws.sheet_state = "hidden"


# ============================================================ helpers
def write_table(sheet_name, cols, rows, kind, table_name, widths=None, formats=None,
                manual_cols=()):
    ws = wb.create_sheet(sheet_name)
    fill = {"manual": FILL_MANUAL, "machine": FILL_MACHINE}[kind]
    for j, (name, ru, typ, req, rule, desc) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=name)
        c.font = F_H
        c.fill = FILL_MANUAL if name in manual_cols else fill
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.comment = Comment("%s\nтип: %s | обязательность: %s\n%s\n%s" % (ru, typ, req, rule, desc), "schema v3.0")
    for i, row in enumerate(rows, start=2):
        for j, val in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=val)
            c.font = F_INPUT if (kind == "manual" or cols[j - 1][0] in manual_cols) else F_BASE
            if formats and cols[j - 1][0] in formats:
                c.number_format = formats[cols[j - 1][0]]
            elif isinstance(val, dt.date) and not isinstance(val, dt.datetime):
                c.number_format = FMT_DATE
    last = max(len(rows) + 1, 2)
    tab = Table(displayName=table_name, ref="A1:%s%d" % (get_column_letter(len(cols)), last))
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight9" if kind == "manual" else "TableStyleLight11",
        showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    # Автофильтр на уровне ЛИСТА здесь не ставится намеренно. Умная таблица
    # заводит свой autoFilter сама (openpyxl добавляет его при headerRowCount),
    # а два фильтра на одном диапазоне Excel считает испорченным содержимым:
    # открывает файл с предложением восстановить и выбрасывает таблицу целиком
    # («Удалённое свойство: Таблица из части /xl/tables/tableN.xml»). Выпадающие
    # списки фильтра при этом не теряются — их даёт сама таблица.
    for j, (name, *_r) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(name, max(14, len(name) + 3))
    ws.row_dimensions[1].height = 30
    return ws


def dv(ws, rng, **kw):
    d = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    ws.add_data_validation(d)
    d.add(rng)
    return d


# ============================================================ dim_portfolio
ws_dim = write_table("dim_portfolio", SCHEMA["dim_portfolio"],
                     [[p[0], p[1], p[2], p[3], p[4], p[5]] for p in PORTFOLIOS], "manual",
                     "tbl_dim_portfolio",
                     widths={"portfolio_code": 18, "portfolio_name": 38, "portfolio_type": 16,
                             "include_in_total": 17, "is_limit_controlled": 19})

# ============================================================ fact_limit
lim_rows = [[t, lim, round(lim * 0.70), round(lim * 0.90), lim,
             dt.date(2026, 1, 1), "Иванов И.И."] for t, lim in TYPE_LIMITS]
ws_lim = write_table("fact_limit", SCHEMA["fact_limit"], lim_rows, "manual", "tbl_fact_limit",
                     widths={"portfolio_type": 18, "limit_amount": 20, "green_max_util": 17,
                             "yellow_max_util": 17, "red_max_util": 17, "updated_by": 20,
                             "valid_from": 14},
                     formats={"limit_amount": FMT_AMT, "green_max_util": FMT_AMT,
                              "yellow_max_util": FMT_AMT, "red_max_util": FMT_AMT})

# ============================================================ fact_type_daily
ws_td = write_table("fact_type_daily", SCHEMA["fact_type_daily"], TYPE_FACTS, "machine",
                    "tbl_fact_type_daily",
                    widths={"business_date": 14, "portfolio_type": 16, "volume_amount": 18},
                    formats={"volume_amount": FMT_AMT})

# ============================================================ fact_portfolio_snapshot
ws_sn = write_table("fact_portfolio_snapshot", SCHEMA["fact_portfolio_snapshot"], SNAPSHOT,
                    "machine", "tbl_fact_portfolio_snapshot",
                    widths={"business_date": 14, "portfolio_code": 18, "volume_t0": 16,
                            "volume_t7": 16, "duration_current_yrs": 19,
                            "duration_target_yrs": 19, "note_text": 85},
                    formats={"volume_t0": FMT_AMT, "volume_t7": FMT_AMT,
                             "duration_current_yrs": FMT_DUR, "duration_target_yrs": FMT_DUR},
                    manual_cols=("note_text",))
for row in ws_sn.iter_rows(min_row=2, min_col=7, max_col=7):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")

# ============================================================ named ranges
wb.defined_names.add(DefinedName("PORTFOLIO_CODES", attr_text="dim_portfolio!$A$2:$A$%d" % (len(PORTFOLIOS) + 1)))
wb.defined_names.add(DefinedName("PORTFOLIO_TYPES", attr_text="fact_limit!$A$2:$A$%d" % (len(TYPES) + 1)))
wb.defined_names.add(DefinedName("BOOL_VALUES", attr_text="_lists!$A$2:$A$3"))
wb.defined_names.add(DefinedName("BUSINESS_DATE", attr_text="view_monitor!$B$3"))
wb.defined_names.add(DefinedName("LOOKBACK_DAYS", attr_text="view_monitor!$D$3"))

# ============================================================ validation
dv(ws_dim, "C2:C%d" % SNAPROW, type="list", formula1="=PORTFOLIO_TYPES",
   errorTitle="Неизвестный тип", error="Тип должен быть заведён строкой в fact_limit.")
dv(ws_dim, "D2:E%d" % SNAPROW, type="list", formula1="=BOOL_VALUES",
   errorTitle="Только TRUE/FALSE", error="Допустимо TRUE или FALSE.")

dv(ws_lim, "B2:E%d" % LIMROW, type="decimal", operator="greaterThanOrEqual", formula1=0,
   errorTitle="Отрицательная сумма", error="Лимит и границы зон задаются абсолютными суммами >= 0.")
dv(ws_lim, "F2:F%d" % LIMROW, type="date", operator="between",
   formula1=dt.date(2020, 1, 1), formula2=dt.date(2040, 12, 31),
   errorTitle="Некорректная дата", error="Введите дату в формате YYYY-MM-DD.")

dv(ws_td, "A2:A%d" % MAXROW, type="date", operator="between",
   formula1=dt.date(2020, 1, 1), formula2=dt.date(2040, 12, 31),
   errorTitle="Некорректная дата", error="Введите дату в формате YYYY-MM-DD.")
dv(ws_td, "B2:B%d" % MAXROW, type="list", formula1="=PORTFOLIO_TYPES",
   errorTitle="Неизвестный тип", error="В этой таблице допустимы только типы из fact_limit.")
dv(ws_td, "C2:C%d" % MAXROW, type="decimal", operator="greaterThanOrEqual", formula1=0,
   errorTitle="Отрицательный объём", error="Объём не может быть меньше нуля.")

dv(ws_sn, "A2:A%d" % SNAPROW, type="date", operator="between",
   formula1=dt.date(2020, 1, 1), formula2=dt.date(2040, 12, 31),
   errorTitle="Некорректная дата", error="Введите дату в формате YYYY-MM-DD.")
dv(ws_sn, "B2:B%d" % SNAPROW, type="list", formula1="=PORTFOLIO_CODES",
   errorTitle="Неизвестный портфель", error="Код должен быть из справочника dim_portfolio.")
dv(ws_sn, "C2:D%d" % SNAPROW, type="decimal", operator="greaterThanOrEqual", formula1=0,
   errorTitle="Отрицательный объём", error="Объём не может быть меньше нуля.")

# ============================================================ диапазоны
TD_DATE = "fact_type_daily!$A$2:$A$%d" % MAXROW
TD_TYPE = "fact_type_daily!$B$2:$B$%d" % MAXROW
TD_VOL = "fact_type_daily!$C$2:$C$%d" % MAXROW
S_CODE = "fact_portfolio_snapshot!$B$2:$B$%d" % SNAPROW
S_DATE = "fact_portfolio_snapshot!$A$2:$A$%d" % SNAPROW
S_T0 = "fact_portfolio_snapshot!$C$2:$C$%d" % SNAPROW
S_T7 = "fact_portfolio_snapshot!$D$2:$D$%d" % SNAPROW
S_DC = "fact_portfolio_snapshot!$E$2:$E$%d" % SNAPROW
S_DT = "fact_portfolio_snapshot!$F$2:$F$%d" % SNAPROW
S_NOTE = "fact_portfolio_snapshot!$G$2:$G$%d" % SNAPROW
L_TYPE = "fact_limit!$A$2:$A$%d" % LIMROW
L_AMT = "fact_limit!$B$2:$B$%d" % LIMROW
L_G = "fact_limit!$C$2:$C$%d" % LIMROW
L_Y = "fact_limit!$D$2:$D$%d" % LIMROW
L_R = "fact_limit!$E$2:$E$%d" % LIMROW
D_CODE = "dim_portfolio!$A$2:$A$%d" % SNAPROW
D_TYPE = "dim_portfolio!$C$2:$C$%d" % SNAPROW
D_INC = "dim_portfolio!$D$2:$D$%d" % SNAPROW

# ============================================================ view_monitor
ws = wb.create_sheet("view_monitor")
ws.sheet_view.showGridLines = False
ws["A1"] = "Мониторинг портфелей"
ws["A1"].font = F_TITLE
ws["A2"] = ("Срез по детальным портфелям из fact_portfolio_snapshot. Суммы в %s. "
            "Светофор — на листе view_by_type: лимит установлен на тип, а не на портфель." % UNIT)
ws["A2"].font = F_SUB
ws["A3"] = "Отчётная дата:"
ws["A3"].font = F_BOLD
c = ws["B3"]
c.value = BUSINESS_DATE
c.font = F_INPUT
c.fill = FILL_INPUTCELL
c.border = BORDER
c.number_format = FMT_DATE
c.protection = Protection(locked=False)
c.comment = Comment("Должна совпадать с датой среза в fact_portfolio_snapshot (CHK_09) "
                    "и с последней датой в fact_type_daily (CHK_02).", "schema v3.0")
ws["C3"] = "Сдвиг сравнения, дней:"
ws["C3"].font = F_BOLD
ws["C3"].alignment = Alignment(horizontal="right")
c = ws["D3"]
c.value = LOOKBACK
c.font = F_INPUT
c.fill = FILL_INPUTCELL
c.border = BORDER
c.protection = Protection(locked=False)
c.comment = Comment("Сдвиг для сравнения объёмов типов (календарных дней). "
                    "По портфелям объём T-7 приходит готовой колонкой в срезе — "
                    "убедитесь, что выгрузка использует тот же сдвиг.", "schema v3.0")

VH = ["Код", "Портфель", "Тип", "Объём T0", "Объём T-7", "Δ объёма", "Δ, %",
      "Дюрация тек.", "Дюрация цель", "Δ дюрации", "Лимит типа"]
HR = 5
for j, h in enumerate(VH, start=1):
    c = ws.cell(row=HR, column=j, value=h)
    c.font = F_H
    c.fill = FILL_VIEW
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[HR].height = 32
ws.cell(row=HR, column=11).comment = Comment(
    "Лимит типа, к которому относится портфель. Лимит установлен на суммарный объём типа, "
    "поэтому по отдельному портфелю утилизация не считается — смотрите view_by_type.", "schema v3.0")

n = len(PORTFOLIOS)
FIRST, LAST = HR + 1, HR + n
LK = "MATCH($A{r},%s,0)" % S_CODE
for i in range(n):
    r = FIRST + i
    dr = 2 + i
    f = {
        "A": "=dim_portfolio!$A$%d" % dr,
        "B": "=dim_portfolio!$B$%d" % dr,
        "C": "=dim_portfolio!$C$%d" % dr,
        "D": "=IFERROR(INDEX({v},{m}),\"\")".format(v=S_T0, m=LK.format(r=r)),
        "E": "=IFERROR(INDEX({v},{m}),\"\")".format(v=S_T7, m=LK.format(r=r)),
        "F": "=IF(OR($D{r}=\"\",$E{r}=\"\"),\"\",$D{r}-$E{r})".format(r=r),
        "G": "=IF(OR($E{r}=\"\",$E{r}=0),\"\",$D{r}/$E{r}-1)".format(r=r),
        "H": "=IFERROR(INDEX({v},{m}),\"\")".format(v=S_DC, m=LK.format(r=r)),
        "I": "=IFERROR(INDEX({v},{m}),\"\")".format(v=S_DT, m=LK.format(r=r)),
        "J": "=IF(OR($H{r}=\"\",$I{r}=\"\"),\"\",$H{r}-$I{r})".format(r=r),
        "K": "=IFERROR(INDEX({a},MATCH($C{r},{t},0)),\"\")".format(a=L_AMT, t=L_TYPE, r=r),
    }
    for col, formula in f.items():
        cc = ws["%s%d" % (col, r)]
        cc.value = formula
        cc.font = F_LINK if col in ("A", "B", "C") else F_CALC
        cc.border = BORDER
    for col, fmt in (("D", FMT_AMT), ("E", FMT_AMT), ("F", FMT_AMT), ("G", FMT_PCT),
                     ("H", FMT_DUR), ("I", FMT_DUR), ("J", FMT_DUR), ("K", FMT_AMT)):
        ws["%s%d" % (col, r)].number_format = fmt

for col, w in {"A": 18, "B": 36, "C": 12, "D": 14, "E": 14, "F": 13, "G": 10,
               "H": 13, "I": 13, "J": 12, "K": 14}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "C%d" % FIRST
delta_rng = "G%d:G%d" % (FIRST, LAST)
ws.conditional_formatting.add(delta_rng, CellIsRule(operator="lessThan", formula=["-0.05"], fill=FILL_RED))
ws.conditional_formatting.add(delta_rng, CellIsRule(operator="greaterThan", formula=["0.05"], fill=FILL_YELLOW))
ws.protection.sheet = True
ws.protection.selectLockedCells = False

V_TYPE = "view_monitor!$C$%d:$C$%d" % (FIRST, LAST)
V_T0 = "view_monitor!$D$%d:$D$%d" % (FIRST, LAST)

# ============================================================ view_by_type
ws = wb.create_sheet("view_by_type")
ws.sheet_view.showGridLines = False
ws["A1"] = "Свод по типам и контроль лимитов"
ws["A1"].font = F_TITLE
ws["A2"] = ("Лимит установлен на суммарный объём типа, поэтому светофор живёт здесь. "
            "Объёмы типов — из fact_type_daily, «Сумма по портфелям» — из детального среза.")
ws["A2"].font = F_SUB
ws["A3"] = "Отчётная дата:"
ws["A3"].font = F_BOLD
ws["B3"] = "=BUSINESS_DATE"
ws["B3"].font = F_LINK
ws["B3"].number_format = FMT_DATE

TH = ["Тип портфеля", "В итог", "Портфелей", "Объём типа T0", "Объём типа T-7", "Δ объёма",
      "Δ, %", "Сумма по портфелям", "Расхождение", "Лимит", "Утилизация", "Зона",
      "Свободный лимит"]
THR = 5
for j, h in enumerate(TH, start=1):
    c = ws.cell(row=THR, column=j, value=h)
    c.font = F_H
    c.fill = FILL_VIEW
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[THR].height = 32

nt = len(TYPES)
TFIRST, TLAST = THR + 1, THR + nt
for i in range(nt):
    r = TFIRST + i
    dr = 2 + i
    vals = {
        "A": "=fact_limit!$A$%d" % dr,
        "B": "=IF(COUNTIFS({t},$A{r},{inc},TRUE)>0,TRUE,FALSE)".format(t=D_TYPE, inc=D_INC, r=r),
        "C": "=COUNTIFS(%s,$A%d)" % (V_TYPE, r),
        "D": "=SUMIFS(%s,%s,BUSINESS_DATE,%s,$A%d)" % (TD_VOL, TD_DATE, TD_TYPE, r),
        "E": "=SUMIFS(%s,%s,BUSINESS_DATE-LOOKBACK_DAYS,%s,$A%d)" % (TD_VOL, TD_DATE, TD_TYPE, r),
        "F": "=IF(OR($D{r}=0,$E{r}=0),\"\",$D{r}-$E{r})".format(r=r),
        "G": "=IF($E{r}=0,\"\",$D{r}/$E{r}-1)".format(r=r),
        "H": "=SUMIFS(%s,%s,$A%d)" % (V_T0, V_TYPE, r),
        "I": "=IF(OR($D{r}=0,$H{r}=0),\"\",$H{r}/$D{r}-1)".format(r=r),
        "J": "=IFERROR(INDEX({a},MATCH($A{r},{t},0)),\"\")".format(a=L_AMT, t=L_TYPE, r=r),
        "K": "=IF(OR($J{r}=\"\",$J{r}=0),\"\",$D{r}/$J{r})".format(r=r),
        "L": ("=IF(OR($J{r}=\"\",$D{r}=0),\"НЕТ ДАННЫХ\","
              "IF($D{r}<=IFERROR(INDEX({g},MATCH($A{r},{t},0)),0),\"ЗЕЛЁНАЯ\","
              "IF($D{r}<=IFERROR(INDEX({y},MATCH($A{r},{t},0)),0),\"ЖЁЛТАЯ\","
              "IF($D{r}<=IFERROR(INDEX({rd},MATCH($A{r},{t},0)),0),\"КРАСНАЯ\",\"ПРЕВЫШЕНИЕ\"))))"
              ).format(r=r, g=L_G, y=L_Y, rd=L_R, t=L_TYPE),
        "M": "=IF($J{r}=\"\",\"\",$J{r}-$D{r})".format(r=r),
    }
    for col, v in vals.items():
        cc = ws["%s%d" % (col, r)]
        cc.value = v
        cc.font = F_LINK if col == "A" else F_CALC
        cc.border = BORDER
    for col, fmt in (("D", FMT_AMT), ("E", FMT_AMT), ("F", FMT_AMT), ("G", FMT_PCT),
                     ("H", FMT_AMT), ("I", FMT_PCT3), ("J", FMT_AMT), ("K", FMT_PCT),
                     ("M", FMT_AMT)):
        ws["%s%d" % (col, r)].number_format = fmt
    ws["B%d" % r].alignment = Alignment(horizontal="center")
    ws["L%d" % r].alignment = Alignment(horizontal="center")

TOT = TLAST + 1
T_INC = "$B$%d:$B$%d" % (TFIRST, TLAST)
tot_vals = {
    "A": "ИТОГО (include_in_total)",
    "B": "",
    "C": "=SUM($C$%d:$C$%d)" % (TFIRST, TLAST),
    "D": "=SUMIFS($D$%d:$D$%d,%s,TRUE)" % (TFIRST, TLAST, T_INC),
    "E": "=SUMIFS($E$%d:$E$%d,%s,TRUE)" % (TFIRST, TLAST, T_INC),
    "F": "=IF(OR($D{r}=0,$E{r}=0),\"\",$D{r}-$E{r})".format(r=TOT),
    "G": "=IF($E{r}=0,\"\",$D{r}/$E{r}-1)".format(r=TOT),
    "H": "=SUMIFS($H$%d:$H$%d,%s,TRUE)" % (TFIRST, TLAST, T_INC),
    "I": "=IF(OR($D{r}=0,$H{r}=0),\"\",$H{r}/$D{r}-1)".format(r=TOT),
    "J": "=SUMIFS($J$%d:$J$%d,%s,TRUE)" % (TFIRST, TLAST, T_INC),
    "K": "=IF($J{r}=0,\"\",$D{r}/$J{r})".format(r=TOT),
    "L": "",
    "M": "=IF($J{r}=0,\"\",$J{r}-$D{r})".format(r=TOT),
}
for col, v in tot_vals.items():
    cc = ws["%s%d" % (col, TOT)]
    cc.value = v
    cc.font = F_BOLD
    cc.fill = FILL_TOTAL
    cc.border = BORDER
for col, fmt in (("D", FMT_AMT), ("E", FMT_AMT), ("F", FMT_AMT), ("G", FMT_PCT),
                 ("H", FMT_AMT), ("I", FMT_PCT3), ("J", FMT_AMT), ("K", FMT_PCT), ("M", FMT_AMT)):
    ws["%s%d" % (col, TOT)].number_format = fmt
ws["A%d" % TOT].comment = Comment(
    "Итог считается по типам, у которых есть хотя бы один портфель с include_in_total = TRUE "
    "в dim_portfolio. Если тип вложен в другой (например, КУАП внутри HTM), снимите флаг у его "
    "портфелей, иначе объём задвоится.", "schema v3.0")
ws["I%d" % TFIRST].comment = Comment(
    "Сумма объёмов портфелей этого типа (детальный срез) против объёма типа за ту же дату "
    "(история). Источники разные, поэтому расхождение возможно. Порог — 0.5%%, CHK_16.", "schema v3.0")

for col, w in {"A": 24, "B": 9, "C": 12, "D": 16, "E": 16, "F": 14, "G": 11, "H": 19,
               "I": 14, "J": 16, "K": 13, "L": 15, "M": 17}.items():
    ws.column_dimensions[col].width = w
zone_rng = "L%d:L%d" % (TFIRST, TLAST)
for word, fill in (("ЗЕЛЁНАЯ", FILL_GREEN), ("ЖЁЛТАЯ", FILL_YELLOW),
                   ("КРАСНАЯ", FILL_RED), ("ПРЕВЫШЕНИЕ", FILL_BREACH)):
    ws.conditional_formatting.add(zone_rng, CellIsRule(operator="equal", formula=['"%s"' % word], fill=fill))
diff_rng = "I%d:I%d" % (TFIRST, TOT)
ws.conditional_formatting.add(diff_rng, CellIsRule(operator="greaterThan", formula=[str(TOL)], fill=FILL_RED))
ws.conditional_formatting.add(diff_rng, CellIsRule(operator="lessThan", formula=[str(-TOL)], fill=FILL_RED))
ws.protection.sheet = True
ws.protection.selectLockedCells = False

BT_DIFF = "view_by_type!$I$%d:$I$%d" % (TFIRST, TLAST)
BT_ZONE = "view_by_type!$L$%d:$L$%d" % (TFIRST, TLAST)
BT_UTIL = "view_by_type!$K$%d:$K$%d" % (TFIRST, TLAST)

# ============================================================ checks
ws = wb.create_sheet("checks")
ws.sheet_view.showGridLines = False
ws["A1"] = "Проверки качества данных"
ws["A1"].font = F_TITLE
ws["A2"] = "Все строки должны быть OK. Любой FAIL блокирует загрузку в BI."
ws["A2"].font = F_SUB
for j, h in enumerate(["check_id", "Проверка", "Значение", "Ожидание", "Статус"], start=1):
    c = ws.cell(row=4, column=j, value=h)
    c.font = F_H
    c.fill = FILL_VIEW
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

NTYPES = "COUNTA(%s)" % L_TYPE
NPORT = "COUNTA(%s)" % D_CODE
checks = [
    ("CHK_01", "Строк в fact_type_daily", "=COUNTA(%s)" % TD_TYPE, "больше 0",
     "=IF($C{r}>0,\"OK\",\"FAIL\")"),
    ("CHK_02", "Максимальная дата в fact_type_daily равна отчётной дате", "=MAX(%s)" % TD_DATE,
     "равна view_monitor!B3", "=IF($C{r}=BUSINESS_DATE,\"OK\",\"FAIL\")"),
    ("CHK_03", "Минимальная дата в fact_type_daily", "=MIN(%s)" % TD_DATE, "справочно", "=\"OK\""),
    ("CHK_04", "Строк на отчётную дату = число типов", "=COUNTIFS(%s,BUSINESS_DATE)" % TD_DATE,
     "равно числу типов в fact_limit", "=IF($C{r}=%s,\"OK\",\"FAIL\")" % NTYPES),
    ("CHK_05", "Строк на дату T-7 (иначе динамика типа считается от нуля)",
     "=COUNTIFS(%s,BUSINESS_DATE-LOOKBACK_DAYS)" % TD_DATE,
     "равно числу типов в fact_limit", "=IF($C{r}=%s,\"OK\",\"FAIL\")" % NTYPES),
    ("CHK_06", "Отрицательные объёмы в fact_type_daily", "=COUNTIFS(%s,\"<0\")" % TD_VOL,
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_07", "Типы в fact_type_daily, которых нет в fact_limit",
     "=SUMPRODUCT((%s<>\"\")*(COUNTIFS(%s,%s&\"\")=0))" % (TD_TYPE, L_TYPE, TD_TYPE),
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_08", "Строк в fact_portfolio_snapshot = число портфелей", "=COUNTA(%s)" % S_CODE,
     "равно числу строк dim_portfolio", "=IF($C{r}=%s,\"OK\",\"FAIL\")" % NPORT),
    ("CHK_09", "Дата среза совпадает с отчётной датой во всех строках",
     "=COUNTA({d})-COUNTIFS({d},BUSINESS_DATE)".format(d=S_DATE), "ровно 0",
     "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_10", "Дубли портфеля в fact_portfolio_snapshot",
     "=SUMPRODUCT((%s<>\"\")*(COUNTIFS(%s,%s&\"\")>1))" % (S_CODE, S_CODE, S_CODE),
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_11", "Портфели среза, которых нет в dim_portfolio",
     "=SUMPRODUCT((%s<>\"\")*(COUNTIFS(%s,%s&\"\")=0))" % (S_CODE, D_CODE, S_CODE),
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_12", "Пустой или нулевой объём T0 в срезе",
     "=COUNTA({c})-COUNTIFS({v},\">0\")".format(c=S_CODE, v=S_T0), "ровно 0",
     "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_13", "Пустой объём T-7 в срезе (тогда дельта не считается)",
     "=COUNTA({c})-COUNTIFS({v},\">0\")".format(c=S_CODE, v=S_T7), "ровно 0",
     "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_14", "Пустая текущая дюрация в срезе",
     "=COUNTA({c})-COUNTIFS({v},\">0\")".format(c=S_CODE, v=S_DC), "ровно 0",
     "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_15", "Портфели без заметки (загрузчик мог затереть note_text)",
     "=COUNTA({c})-COUNTA({n})".format(c=S_CODE, n=S_NOTE), "справочно, следить за ростом",
     "=\"OK\""),
    ("CHK_16", "Расхождение свода портфелей и объёма типа выше порога",
     "=COUNTIF({d},\">%s\")+COUNTIF({d},\"<-%s\")".format(d=BT_DIFF) % (TOL, TOL),
     "ровно 0 (порог 0.5%)", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_17", "Типы в dim_portfolio, для которых нет лимита в fact_limit",
     "=SUMPRODUCT((%s<>\"\")*(COUNTIFS(%s,%s&\"\")=0))" % (D_TYPE, L_TYPE, D_TYPE),
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_18", "Дубли типа в fact_limit",
     "=SUMPRODUCT((%s<>\"\")*(COUNTIFS(%s,%s&\"\")>1))" % (L_TYPE, L_TYPE, L_TYPE),
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_19", "Лимиты <= 0", "=COUNTIFS(%s,\"<=0\")" % L_AMT, "ровно 0",
     "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_20", "Границы зон нарушены (green < yellow < red)",
     "=SUMPRODUCT((%s<>\"\")*((%s>=%s)+(%s>=%s)>0))" % (L_TYPE, L_G, L_Y, L_Y, L_R),
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_21", "Красная граница выше установленного лимита",
     "=SUMPRODUCT((%s<>\"\")*(%s>%s))" % (L_TYPE, L_R, L_AMT),
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_22", "Типы в красной зоне", "=COUNTIF(%s,\"КРАСНАЯ\")" % BT_ZONE, "справочно", "=\"OK\""),
    ("CHK_23", "Типы с превышением лимита", "=COUNTIF(%s,\"ПРЕВЫШЕНИЕ\")" % BT_ZONE,
     "ровно 0", "=IF($C{r}=0,\"OK\",\"FAIL\")"),
    ("CHK_24", "Утилизация типа выше 100%", "=COUNTIF(%s,\">1\")" % BT_UTIL, "ровно 0",
     "=IF($C{r}=0,\"OK\",\"FAIL\")"),
]
for i, (cid, desc, val, exp, st) in enumerate(checks):
    r = 5 + i
    ws.cell(row=r, column=1, value=cid).font = F_BASE
    ws.cell(row=r, column=2, value=desc).font = F_BASE
    c = ws.cell(row=r, column=3, value=val)
    c.font = F_CALC
    if cid in ("CHK_02", "CHK_03"):
        c.number_format = FMT_DATE
    ws.cell(row=r, column=4, value=exp).font = F_SUB
    s = ws.cell(row=r, column=5, value=st.format(r=r))
    s.font = F_BOLD
    s.alignment = Alignment(horizontal="center")
    for j in range(1, 6):
        ws.cell(row=r, column=j).border = BORDER

st_rng = "E5:E%d" % (4 + len(checks))
ws.conditional_formatting.add(st_rng, CellIsRule(operator="equal", formula=['"OK"'], fill=FILL_GREEN))
ws.conditional_formatting.add(st_rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=FILL_RED))
for col, w in {"A": 12, "B": 64, "C": 18, "D": 32, "E": 12}.items():
    ws.column_dimensions[col].width = w
ws.protection.sheet = True
ws.protection.selectLockedCells = False

# ============================================================ dict
ws = wb.create_sheet("dict")
dhead = ["table_name", "column_name", "ru_label", "data_type", "required", "rule",
         "description", "grain", "owner"]
GRAIN = {
    "dim_portfolio": "portfolio_code",
    "fact_limit": "portfolio_type",
    "fact_type_daily": "business_date + portfolio_type",
    "fact_portfolio_snapshot": "portfolio_code (одна дата)",
}
for j, h in enumerate(dhead, start=1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = F_H
    c.fill = FILL_VIEW
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
r = 2
for tname, cols in SCHEMA.items():
    owner = "пайплайн" if TABLE_KIND[tname] == "machine" else "казначейство"
    for name, ru, typ, req, rule, desc in cols:
        own = "казначейство (ручная)" if name == "note_text" else owner
        for j, v in enumerate([tname, name, ru, typ, req, rule, desc, GRAIN[tname], own], start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_BASE
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
tab = Table(displayName="tbl_dict", ref="A1:I%d" % (r - 1))
tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight11", showRowStripes=True)
ws.add_table(tab)
ws.freeze_panes = "A2"
for col, w in {"A": 24, "B": 24, "C": 24, "D": 11, "E": 14, "F": 32, "G": 68,
               "H": 26, "I": 22}.items():
    ws.column_dimensions[col].width = w
ws.row_dimensions[1].height = 30

wb.save(OUT)
print("saved", OUT)
print("types:", len(TYPES), "| portfolios:", len(PORTFOLIOS),
      "| type_daily rows:", len(TYPE_FACTS), "| snapshot rows:", len(SNAPSHOT),
      "| dates:", len(DATES), "| T-7 =", d7)
