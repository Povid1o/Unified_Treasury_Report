"""ETL отчёта «Динамика портфелей»: два среза выгрузки позиций (T0 и T-7)
плюс предыдущий выпуск отчёта -> четыре таблицы схемы v3.0.

Запись .xlsx живёт отдельно, в workbook.py: здесь только разбор входных
файлов и слияние с предыдущим выпуском.

Почему предыдущий выпуск — это ВХОД. Выгрузка даёт ровно два среза, а лист
fact_type_daily требует ежедневной истории объёмов по типам с 01.01.2026.
Взять её неоткуда, кроме как накапливать: каждый запуск дописывает в историю
одну дату (T0) и переносит всё остальное — историю прошлых дат, лимиты
(их заводят руками), справочник портфелей и заметки — из предыдущего файла.

Разбор иерархии. В выгрузке нет колонки с портфелем: иерархия задана
значениями внутри колонки «Тип актива». Строка «Позиция: <CODE>» открывает
портфель, все последующие строки — его бумаги, вплоть до следующей такой
строки. Строки до первого маркера — шапка выгрузки.
"""
import datetime as dt
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd

BASE_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE_DIR))

import config  # noqa: E402
from common import excel_io, settings  # noqa: E402
from common.logging_utils import get_logger  # noqa: E402

logger = get_logger("portfolio_dynamics")


class PortfolioDynamicsError(RuntimeError):
    """Ошибка чтения или обработки входных файлов «Динамики портфелей»."""


# ── Колонки выходных таблиц (схема v3.0, см. CONTRACT.md) ────────────────────
DIM_COLUMNS = [
    "portfolio_code", "portfolio_name", "portfolio_type",
    "include_in_total", "is_limit_controlled", "sort_order",
]
# limit_remaining дописан ПОСЛЕ updated_by намеренно: формулы витрин и checks
# адресуют колонки fact_limit буквами ($B — лимит, $C..$E — границы зон), и
# вставка в середину сдвинула бы их все. В хвосте новая колонка не двигает
# ничего, и файл остаётся совместимым со всем, что читает лист по позициям.
LIMIT_COLUMNS = [
    "portfolio_type", "limit_amount", "green_max_util", "yellow_max_util",
    "red_max_util", "valid_from", "updated_by", "limit_remaining",
]
# Колонки, появившиеся позже схемы v3.0: в выпуске, сделанном до их появления,
# их нет, и требовать их от предыдущего файла нельзя — он входит в следующий.
OPTIONAL_LIMIT_COLUMNS = ["limit_remaining"]
TYPE_DAILY_COLUMNS = ["business_date", "portfolio_type", "volume_amount"]
SNAPSHOT_COLUMNS = [
    "business_date", "portfolio_code", "volume_t0", "volume_t7",
    "duration_current_yrs", "duration_target_yrs", "note_text",
]

# ── Колонки входной выгрузки ─────────────────────────────────────────────────
# Сравнение идёт по excel_io.normalize_label (без пробелов, casefold, ё->е):
# в выгрузках у заголовков висячие пробелы, неразрывные пробелы и переносы.
COL_ASSET_TYPE = "Тип актива"
COL_VALUE = "Чистая стоимость позиции (кон.)"
COL_DURATION_START = "Duration (нач.)"
COL_DURATION_END = "Дюрация"

# Маппинг дюраций на схему намеренно «перекрёстный»: «Duration (нач.)» ->
# duration_current_yrs, «Дюрация» (конечная) -> duration_target_yrs. Имена не
# совпадают по смыслу с «нач.»/«кон.» — это осознанное решение при согласовании
# схемы v3.0, колонки не переименовывать и маппинг соблюдать буквально.
#
# Конечная «Дюрация» временно не читается (настройка «Читать конечную дюрацию
# из выгрузки», по умолчанию выключена): колонка duration_target_yrs теперь
# заполняется дюрацией, установленной КУАП, — см. apply_kuap_durations.

POSITION_MARKER = excel_io.normalize_label("Позиция:")
TOTAL_MARKERS = tuple(excel_io.normalize_label(x) for x in ("Итого", "Всего", "Grand Total", "Total"))

# Вторая дата из "Позиция за период ... - ... - SECURITIES" — и в имени файла,
# и в первой строке самого листа. Квадратные скобки необязательны: выгрузка
# именует ФАЙЛ без них и с произвольным числом пробелов
# ("Позиция за период   18.09.2026  -  18.09.2026   - SECURITIES.xlsx"),
# а в ШАПКЕ ЛИСТА пишет со скобками. Разбирать надо оба вида.
PERIOD_DATES_PATTERN = re.compile(
    r"\[?(\d{2}\.\d{2}\.\d{4})\]?\s*-\s*\[?(\d{2}\.\d{2}\.\d{4})\]?"
)

# "30,000,000" — запятая как разделитель тысяч; "3,14" — как десятичная.
_THOUSANDS_GROUPED = re.compile(r"^-?\d{1,3}(,\d{3})+$")
_SPACES = re.compile(r"[\s   ]+")
_MISSING_TOKENS = {"", "-", "—", "–", "...", "…", "н/д", "нд", "n/a", "na", "#н/д"}


# ════════════════════════════════════════════════════════════════════════════
# Разбор одного среза
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class SliceStats:
    """Счётчики разбора — уходят в лог, чтобы молчаливая потеря строк была видна."""

    sheet_name: str = ""
    header_row: int = 0
    columns: Dict[str, int] = field(default_factory=dict)
    portfolios: int = 0
    securities: int = 0
    rows_without_portfolio: int = 0
    total_rows_skipped: int = 0
    subtotal_mismatches: List[str] = field(default_factory=list)


@dataclass
class PortfolioSlice:
    """Один срез выгрузки, свёрнутый до портфелей."""

    path: Path
    business_date: Optional[dt.date]
    frame: pd.DataFrame  # portfolio_code, volume, duration_current_yrs, duration_target_yrs
    stats: SliceStats


def parse_number(value) -> Optional[float]:
    """Число из ячейки выгрузки: и настоящее число, и текст с разделителями.

    В выгрузках встречается всё сразу: "1 234 567,89" (пробелы + запятая как
    десятичная), "30,000,000" (запятая как разделитель тысяч), неразрывные
    пробелы, прочерки и "..." вместо пропуска. Пустая ячейка и любой из
    прочерков -> None ("нет данных"), а не 0: ноль исказил бы и суммы, и веса.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return None if pd.isna(value) else float(value)

    text = _SPACES.sub("", str(value)).replace("−", "-")
    if text.casefold() in _MISSING_TOKENS:
        return None
    text = text.replace("₽", "").replace("руб.", "").replace("%", "")

    if "," in text and "." in text:
        # Смешанный формат: точка — десятичная, запятая — разделитель тысяч.
        text = text.replace(",", "")
    elif _THOUSANDS_GROUPED.match(text):
        text = text.replace(",", "")
    else:
        text = text.replace(",", ".")

    try:
        return float(text)
    except ValueError:
        return None


def _is_position_marker(normalized: str) -> bool:
    return normalized.startswith(POSITION_MARKER)


def _extract_portfolio_code(raw_value) -> str:
    """"Позиция: AFS_TR_RUR" -> "AFS_TR_RUR" (и "Позиция:AFS_TR_RUR" тоже)."""
    text = str(raw_value)
    _, _, tail = text.partition(":")
    return _SPACES.sub("", tail).upper()


def _find_sheet_and_header(path: Path) -> Tuple[str, pd.DataFrame, int, Dict[str, int]]:
    """Ищет лист по НАЛИЧИЮ колонки «Тип актива», а не по индексу.

    Выгрузки кладут перед нужным листом служебные (кэш сводной, параметры
    подключения), и «первый лист по умолчанию» читается пустым.
    """
    wanted = {
        COL_ASSET_TYPE: excel_io.normalize_label(COL_ASSET_TYPE),
        COL_VALUE: excel_io.normalize_label(COL_VALUE),
        COL_DURATION_START: excel_io.normalize_label(COL_DURATION_START),
        COL_DURATION_END: excel_io.normalize_label(COL_DURATION_END),
    }
    seen_sheets: List[str] = []

    for sheet_name, matrix in excel_io.iter_sheet_matrices(path):
        seen_sheets.append(sheet_name)
        if matrix.empty:
            continue
        for row_idx in range(len(matrix)):
            row = matrix.iloc[row_idx]
            normalized = [excel_io.normalize_label(v) for v in row]
            if wanted[COL_ASSET_TYPE] not in normalized:
                continue
            columns = {
                label: normalized.index(norm)
                for label, norm in wanted.items()
                if norm in normalized
            }
            return sheet_name, matrix, row_idx, columns

    explanation = excel_io.describe_uncached_formulas(path)
    message = (
        f"В файле {path.name} не найден лист с колонкой «{COL_ASSET_TYPE}». "
        f"Листы в файле: {seen_sheets}."
    )
    if explanation:
        message = f"{message}\n{explanation}"
    raise PortfolioDynamicsError(message)


def _business_date_from_name(path: Path) -> Optional[dt.date]:
    match = PERIOD_DATES_PATTERN.search(path.name)
    if not match:
        return None
    return dt.datetime.strptime(match.group(2), config.PORTFOLIO_DYNAMICS_DATE_FORMAT).date()


def _business_date_from_matrix(matrix: pd.DataFrame, header_row: int) -> Optional[dt.date]:
    """Та же строка «Позиция за период [..] - [..] - SECURITIES», но из шапки листа.

    Резерв на случай, если файл переименовали руками и дата из имени пропала.
    """
    for row_idx in range(min(header_row + 1, len(matrix))):
        for value in matrix.iloc[row_idx]:
            if value is None:
                continue
            match = PERIOD_DATES_PATTERN.search(str(value))
            if match:
                return dt.datetime.strptime(
                    match.group(2), config.PORTFOLIO_DYNAMICS_DATE_FORMAT
                ).date()
    return None


def read_business_date(path: Path) -> Optional[dt.date]:
    """Дата среза одного файла выгрузки: из имени файла, иначе из шапки листа.

    Разбор имени бесплатный, чтение книги — нет, поэтому лист читается только
    когда имя ничего не дало (файл переименовали руками).
    """
    path = Path(path)
    from_name = _business_date_from_name(path)
    if from_name is not None:
        return from_name
    try:
        _sheet, matrix, header_row, _cols = _find_sheet_and_header(path)
    except (PortfolioDynamicsError, excel_io.ExcelSourceError):
        return None
    return _business_date_from_matrix(matrix, header_row)


# Книги, которые вообще имеет смысл открывать при поиске по содержимому.
PROBE_SUFFIXES = (".xlsx", ".xlsm")
# Сколько верхних строк листа читать при таком поиске. Шапка выгрузки и строка
# «Позиция за период …» лежат в самом верху; читать весь лист, чтобы узнать
# только дату, незачем.
PROBE_ROWS = 60


def probe_business_date(path: Path) -> Optional[dt.date]:
    """Дата среза по СОДЕРЖИМОМУ книги — дёшево, не читая её целиком.

    Зачем отдельно от read_business_date: тот ради даты разбирает книгу
    полностью (все листы в DataFrame), и перебирать им чужую папку загрузок
    непозволительно долго. Здесь читается только верх каждого листа, и этого
    достаточно: и строка «Позиция за период [..] - [..]», и шапка таблицы
    находятся в первых строках.

    None — файл не похож на выгрузку позиций. Требуется И дата периода, И
    колонка «Тип актива»: одной даты мало, диапазон дат встречается в любом
    отчёте, и без второго признака поиск начал бы принимать за срез что попало.
    """
    path = Path(path)
    if path.suffix.lower() not in PROBE_SUFFIXES:
        return None
    wanted_header = excel_io.normalize_label(COL_ASSET_TYPE)
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # не книга, битая, нет прав — просто не наш файл
        return None
    try:
        for worksheet in workbook.worksheets:
            found_date: Optional[dt.date] = None
            has_header = False
            for row in worksheet.iter_rows(max_row=PROBE_ROWS, values_only=True):
                for value in row:
                    if value is None:
                        continue
                    text = str(value)
                    if found_date is None:
                        match = PERIOD_DATES_PATTERN.search(text)
                        if match:
                            try:
                                found_date = dt.datetime.strptime(
                                    match.group(2), config.PORTFOLIO_DYNAMICS_DATE_FORMAT
                                ).date()
                            except ValueError:
                                pass
                    if not has_header and excel_io.normalize_label(text) == wanted_header:
                        has_header = True
                if found_date is not None and has_header:
                    return found_date
            if found_date is not None and has_header:
                return found_date
    except Exception:
        return None
    finally:
        try:
            workbook.close()
        except Exception:
            pass
    return None


def is_limits_file(path: Path) -> bool:
    """Это выгрузка лимитов, а не срез позиций?

    В папке-дате лежат ТРИ файла: два среза позиций и файл лимитов. Последний
    к срезам отношения не имеет, и принимать его за срез нельзя — иначе разбор
    падает на «не удалось определить дату среза».
    """
    pattern = config.PORTFOLIO_DYNAMICS_LIMITS_REGEX
    return bool(pattern) and re.search(pattern, Path(path).name) is not None


def split_slice_files(files: List[Path], folder_label: str = "") -> Tuple[Path, Path]:
    """Из файлов одной папки-даты выбирает, какой срез T0, а какой T-7.

    Определяется по дате самой выгрузки (вторая дата периода), а не по имени
    файла как строке: поздний срез — T0, ранний — T-7. Так пользователю не нужно
    ничего переименовывать и раскладывать в правильном порядке — достаточно
    положить в папку два файла. Файл лимитов, если он тут же, в расчёт не идёт.
    """
    where = f" в {folder_label}" if folder_label else ""
    real_files = [f for f in files if f.is_file() and not f.name.startswith("~$")]
    limit_files = [f for f in real_files if is_limits_file(f)]
    real_files = [f for f in real_files if f not in limit_files]
    if limit_files:
        logger.debug("Файлы лимитов%s не участвуют в выборе срезов: %s",
                     where, ", ".join(f.name for f in limit_files))
    if len(real_files) < 2:
        raise PortfolioDynamicsError(
            f"Для отчёта нужны ДВА файла (срез T0 и срез T-7), а{where} найдено "
            f"{len(real_files)}: {', '.join(f.name for f in real_files) or 'ничего'}."
        )

    dated = [(read_business_date(f), f) for f in real_files]
    undated = [f.name for d, f in dated if d is None]
    if undated:
        raise PortfolioDynamicsError(
            f"Не удалось определить дату среза у файлов{where}: {', '.join(undated)}. "
            "В имени файла (или в шапке листа) должна быть строка вида "
            "«Позиция за период [дд.мм.гггг] - [дд.мм.гггг]»."
        )

    dated.sort(key=lambda item: item[0], reverse=True)
    if len(dated) > 2:
        ignored = ", ".join(f.name for _d, f in dated[2:])
        logger.warning(
            "В папке%s больше двух файлов — взяты два самых свежих среза (%s и %s), "
            "остальные проигнорированы: %s",
            where, dated[0][1].name, dated[1][1].name, ignored,
        )

    (t0_date, t0_path), (t7_date, t7_path) = dated[0], dated[1]
    if t0_date == t7_date:
        raise PortfolioDynamicsError(
            f"Оба файла{where} — срезы на одну и ту же дату {t0_date.isoformat()} "
            f"({t0_path.name}, {t7_path.name}). Нужны два РАЗНЫХ среза: T0 и T-7."
        )
    return t0_path, t7_path


def _weighted_duration(pairs: List[Tuple[float, float]]) -> Optional[float]:
    """Дюрация портфеля — средневзвешенная по стоимости, а не сумма и не среднее.

    pairs — (дюрация, вес) только по бумагам с заполненной дюрацией: бумага без
    дюрации исключается из весов, но её объём остаётся в сумме портфеля.
    Нулевая сумма весов -> None (делить на ноль нечем и нечего).
    """
    weight_sum = sum(weight for _duration, weight in pairs)
    if not weight_sum:
        return None
    return sum(duration * weight for duration, weight in pairs) / weight_sum


def parse_slice(path: Path, label: str, value_scale: Optional[float] = None) -> PortfolioSlice:
    """Разбирает один срез выгрузки и сворачивает его до портфелей."""
    path = Path(path)
    if not path.exists():
        raise PortfolioDynamicsError(f"[{label}] Файл не найден: {path}")

    scale = config.PORTFOLIO_DYNAMICS_VALUE_SCALE if value_scale is None else value_scale
    sheet_name, matrix, header_row, columns = _find_sheet_and_header(path)

    if not config.PORTFOLIO_DYNAMICS_READ_DURATION_END:
        # Колонка есть в выгрузке, но её значения намеренно не используются:
        # убираем её из найденных, и дальше она ведёт себя как отсутствующая.
        columns.pop(COL_DURATION_END, None)

    stats = SliceStats(sheet_name=sheet_name, header_row=header_row + 1, columns=dict(columns))
    expected = [COL_VALUE, COL_DURATION_START]
    if config.PORTFOLIO_DYNAMICS_READ_DURATION_END:
        expected.append(COL_DURATION_END)
    for missing in expected:
        if missing not in columns:
            logger.warning(
                "[%s] %s: колонка «%s» не найдена на листе %r — соответствующие "
                "значения останутся пустыми.", label, path.name, missing, sheet_name,
            )
    if COL_VALUE not in columns:
        raise PortfolioDynamicsError(
            f"[{label}] В файле {path.name} (лист {sheet_name!r}) нет колонки «{COL_VALUE}» — "
            "объёмы портфелей брать неоткуда."
        )

    type_col = columns[COL_ASSET_TYPE]
    value_col = columns[COL_VALUE]

    order: List[str] = []
    securities: Dict[str, List[Tuple[Optional[float], Optional[float], Optional[float]]]] = {}
    stated: Dict[str, Dict[str, Optional[float]]] = {}
    current: Optional[str] = None

    for row_idx in range(header_row + 1, len(matrix)):
        row = matrix.iloc[row_idx]
        raw_type = row.iloc[type_col]
        normalized = excel_io.normalize_label(raw_type)
        if not normalized:
            continue

        if _is_position_marker(normalized):
            code = _extract_portfolio_code(raw_type)
            if not code:
                logger.warning(
                    "[%s] %s: строка %d — маркер портфеля без кода (%r), пропущена.",
                    label, path.name, row_idx + 1, raw_type,
                )
                continue
            if code not in securities:
                securities[code] = []
                order.append(code)
            current = code
            # Подытоги из самой строки маркера: если выгрузка их даёт, они
            # авторитетнее посчитанных по бумагам (см. сверку ниже).
            stated[code] = {
                "volume": parse_number(row.iloc[value_col]),
                "duration_current_yrs": (
                    parse_number(row.iloc[columns[COL_DURATION_START]])
                    if COL_DURATION_START in columns else None
                ),
                "duration_target_yrs": (
                    parse_number(row.iloc[columns[COL_DURATION_END]])
                    if COL_DURATION_END in columns else None
                ),
            }
            continue

        if normalized.startswith(TOTAL_MARKERS):
            stats.total_rows_skipped += 1
            continue

        if current is None:
            # Шапка выгрузки до первого маркера — в данные не попадает.
            stats.rows_without_portfolio += 1
            continue

        securities[current].append((
            parse_number(row.iloc[value_col]),
            parse_number(row.iloc[columns[COL_DURATION_START]]) if COL_DURATION_START in columns else None,
            parse_number(row.iloc[columns[COL_DURATION_END]]) if COL_DURATION_END in columns else None,
        ))
        stats.securities += 1

    stats.portfolios = len(order)
    if not order:
        raise PortfolioDynamicsError(
            f"[{label}] В файле {path.name} (лист {sheet_name!r}) не найдено ни одной строки "
            f"«{COL_ASSET_TYPE}» вида «Позиция: <КОД>» — портфели определить не по чему."
        )

    records = []
    for code in order:
        rows = securities[code]
        computed_volume = sum(value for value, _dc, _dt in rows if value is not None)
        computed = {
            "volume": computed_volume,
            "duration_current_yrs": _weighted_duration(
                [(dc, value) for value, dc, _dt in rows if dc is not None and value is not None]
            ),
            "duration_target_yrs": _weighted_duration(
                [(dtg, value) for value, _dc, dtg in rows if dtg is not None and value is not None]
            ),
        }
        chosen = {
            key: (stated[code][key] if stated[code][key] is not None else computed[key])
            for key in computed
        }

        declared = stated[code]["volume"]
        if declared is not None and _relative_diff(computed_volume, declared) > config.PORTFOLIO_DYNAMICS_TOLERANCE:
            message = (
                f"{code}: подытог в строке «Позиция: ...» = {declared:,.2f}, "
                f"сумма по бумагам = {computed_volume:,.2f}"
            )
            stats.subtotal_mismatches.append(message)
            logger.warning("[%s] %s: расхождение подытога и суммы по бумагам — %s", label, path.name, message)

        records.append({
            "portfolio_code": code,
            "volume": None if chosen["volume"] is None else chosen["volume"] / scale,
            "duration_current_yrs": chosen["duration_current_yrs"],
            "duration_target_yrs": chosen["duration_target_yrs"],
        })

    business_date = _business_date_from_name(path) or _business_date_from_matrix(matrix, header_row)

    logger.info(
        "[%s] %s: лист %r, шапка в строке %d, портфелей %d, бумаг %d, "
        "строк вне иерархии %d, строк «Итого» отброшено %d, дата среза %s",
        label, path.name, sheet_name, header_row + 1, stats.portfolios, stats.securities,
        stats.rows_without_portfolio, stats.total_rows_skipped,
        business_date.isoformat() if business_date else "не определена",
    )

    return PortfolioSlice(
        path=path,
        business_date=business_date,
        frame=pd.DataFrame(records, columns=["portfolio_code", "volume", "duration_current_yrs", "duration_target_yrs"]),
        stats=stats,
    )


def _relative_diff(computed: Optional[float], declared: Optional[float]) -> float:
    if declared is None or computed is None:
        return 0.0
    if declared == 0:
        return 0.0 if computed == 0 else float("inf")
    return abs(computed / declared - 1)


# ════════════════════════════════════════════════════════════════════════════
# Предыдущий выпуск отчёта
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class PreviousRelease:
    """Четыре таблицы, перенос которых и делает отчёт инкрементальным."""

    path: Optional[Path]
    dim_portfolio: pd.DataFrame
    fact_limit: pd.DataFrame
    fact_type_daily: pd.DataFrame
    notes: pd.DataFrame  # portfolio_code, note_text
    # Объёмы прошлого среза: нужны, чтобы у портфелей, ведущихся вручную,
    # получалась настоящая дельта, а не ноль.
    snapshot_volumes: Optional[Dict[str, float]] = None

    @classmethod
    def empty(cls) -> "PreviousRelease":
        return cls(
            path=None,
            dim_portfolio=pd.DataFrame(columns=DIM_COLUMNS),
            fact_limit=pd.DataFrame(columns=LIMIT_COLUMNS),
            fact_type_daily=pd.DataFrame(columns=TYPE_DAILY_COLUMNS),
            notes=pd.DataFrame(columns=["portfolio_code", "note_text"]),
            snapshot_volumes={},
        )


OUTPUT_FILENAME_PREFIX = "dinamika_portfeley_"


def find_previous_release(output_dir: Path) -> Optional[Path]:
    """Самый свежий выпуск отчёта в выходной папке (по дате изменения).

    Файл за сегодняшний T0 тоже считается предыдущим выпуском — именно на этом
    держится идемпотентность: повторный прогон читает его и ЗАМЕНЯЕТ строки
    истории за свою дату, а не дописывает вторые такие же.
    """
    output_dir = Path(output_dir)
    if not output_dir.exists():
        return None
    candidates = [
        f for f in output_dir.glob(f"{OUTPUT_FILENAME_PREFIX}*.xlsx")
        if f.is_file() and not f.name.startswith("~$")
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda f: f.stat().st_mtime)


def _read_sheet(path: Path, sheet_name: str, columns: List[str],
                optional: Optional[List[str]] = None) -> pd.DataFrame:
    """optional — колонки, которых у старого выпуска может не быть: они
    дозаполняются пустыми. Требовать их значило бы, что первый же запуск после
    обновления схемы падает на файле, сделанном вчера."""
    optional = optional or []
    try:
        frame = pd.read_excel(path, sheet_name=sheet_name)
    except ValueError as exc:  # листа нет в книге
        raise PortfolioDynamicsError(
            f"В предыдущем выпуске {path.name} нет листа {sheet_name!r}: {exc}. "
            "Похоже, это не файл «Динамики портфелей» — укажите другой через --previous "
            "или запустите с --bootstrap."
        ) from exc
    except Exception as exc:
        raise PortfolioDynamicsError(
            f"Не удалось прочитать лист {sheet_name!r} из предыдущего выпуска {path}: {exc}"
        ) from exc

    missing = [c for c in columns if c not in frame.columns]
    added = [c for c in missing if c in optional]
    for column in added:
        frame[column] = None
    if added:
        logger.info(
            "В предыдущем выпуске %s, лист %r нет колонок %s — они появились позже "
            "и заполнены пустыми значениями.", path.name, sheet_name, ", ".join(added),
        )
    missing = [c for c in missing if c not in optional]
    if missing:
        raise PortfolioDynamicsError(
            f"В предыдущем выпуске {path.name}, лист {sheet_name!r}: нет колонок {missing}. "
            "Строка 1 листа должна содержать технические заголовки snake_case."
        )
    return frame[columns].copy()


def load_previous_release(path: Path) -> PreviousRelease:
    """Читает из предыдущего выпуска ровно то, что скрипт не умеет получить из выгрузки."""
    path = Path(path)
    if not path.exists():
        raise PortfolioDynamicsError(f"Предыдущий выпуск не найден: {path}")

    dim = _read_sheet(path, "dim_portfolio", DIM_COLUMNS)
    limit = _read_sheet(path, "fact_limit", LIMIT_COLUMNS,
                        optional=OPTIONAL_LIMIT_COLUMNS)
    history = _read_sheet(path, "fact_type_daily", TYPE_DAILY_COLUMNS)
    snapshot = _read_sheet(path, "fact_portfolio_snapshot", SNAPSHOT_COLUMNS)

    dim = dim[dim["portfolio_code"].notna()].copy()
    limit = limit[limit["portfolio_type"].notna()].copy()
    history = history[history["business_date"].notna() & history["portfolio_type"].notna()].copy()
    history["business_date"] = pd.to_datetime(history["business_date"], errors="coerce").dt.date
    history = history[history["business_date"].notna()]

    notes = snapshot[["portfolio_code", "note_text"]]
    notes = notes[notes["portfolio_code"].notna()].copy()

    _canonicalise_types(dim, limit, history, path)

    volumes = snapshot[snapshot["portfolio_code"].notna()]
    volumes = dict(zip(volumes["portfolio_code"].astype(str),
                       pd.to_numeric(volumes["volume_t0"], errors="coerce")))

    return PreviousRelease(path=path, dim_portfolio=dim, fact_limit=limit,
                           fact_type_daily=history, notes=notes,
                           snapshot_volumes={k: v for k, v in volumes.items() if pd.notna(v)})


def _canonicalise_types(*frames_and_path) -> None:
    """Приводит имена типов из предыдущего выпуска к каноническому написанию.

    Выпуски, сделанные до переименования, содержат TTS. Без приведения в
    отчёте оказались бы ДВА торговых типа сразу — старый в перенесённой истории
    и новый в свежей строке, — и история торгового портфеля разорвалась бы
    надвое.
    """
    *frames, path = frames_and_path
    renamed = 0
    for frame in frames:
        if frame is None or frame.empty or "portfolio_type" not in frame.columns:
            continue
        original = frame["portfolio_type"].copy()
        frame["portfolio_type"] = [
            canonical_type(v) if pd.notna(v) and str(v).strip() else v
            for v in frame["portfolio_type"]
        ]
        renamed += int((original.astype(str) != frame["portfolio_type"].astype(str)).sum())
    if renamed:
        logger.info(
            "Предыдущий выпуск %s: %d значений типа приведено к каноническому написанию "
            "(TTS -> TSS) — иначе история торгового портфеля разорвалась бы надвое.",
            Path(path).name, renamed,
        )


# ════════════════════════════════════════════════════════════════════════════
# Сборка четырёх таблиц
# ════════════════════════════════════════════════════════════════════════════
@dataclass
class PortfolioDynamicsData:
    """То, что etl отдаёт в workbook: четыре таблицы схемы v3.0 и параметры витрин."""

    business_date: dt.date
    lookback_days: int
    dim_portfolio: pd.DataFrame
    fact_limit: pd.DataFrame
    fact_type_daily: pd.DataFrame
    fact_portfolio_snapshot: pd.DataFrame
    new_portfolio_codes: List[str] = field(default_factory=list)
    missing_portfolio_codes: List[str] = field(default_factory=list)
    history_rows_carried: int = 0
    history_rows_added: int = 0
    history_rows_replaced: int = 0
    history_rows_imported: int = 0
    notes_restored: int = 0


def parse_type_parents(raw: Optional[str] = None) -> Dict[str, str]:
    """«HTM_KUAP=HTM» -> {вложенный тип: объемлющий}.

    Лимит в выгрузке СОВОКУПНЫЙ: лимит HTM ограничивает HTM вместе с HTM_KUAP.
    Отсюда три следствия, которые дальше и реализуются: объём объемлющего типа
    считается с вложенными, у портфелей вложенного типа снимается флаг «входит
    в итог» (иначе ИТОГО задвоится), а свой лимит вложенного типа остаётся
    подлимитом внутри общего.
    """
    raw = config.PORTFOLIO_DYNAMICS_TYPE_PARENTS if raw is None else raw
    parents: Dict[str, str] = {}
    for item in str(raw or "").split(","):
        child, _, parent = item.partition("=")
        child, parent = child.strip().upper(), parent.strip().upper()
        if child and parent and child != parent:
            parents[child] = parent
    return parents


def parse_nested_limits(raw: Optional[str] = None) -> Dict[str, float]:
    """«HTM_KUAP=100» -> {вложенный тип: выделенная ему сумма, млн RUB}.

    Это ручное решение казначейства: из совокупного лимита (900 на HTM) часть
    отводится вложенному типу (100 на КУАП), остаток достаётся объемлющему
    (800 на весь остальной HTM). Меняется настройкой, без правки кода.
    """
    raw = config.PORTFOLIO_DYNAMICS_NESTED_LIMITS if raw is None else raw
    allocations: Dict[str, float] = {}
    for item in str(raw or "").split(","):
        name, _, amount = item.partition("=")
        if not name.strip() or not amount.strip():
            continue
        try:
            value = float(amount.strip().replace(" ", "").replace(",", "."))
        except ValueError as exc:
            raise PortfolioDynamicsError(
                f"Не удалось разобрать подлимит {item.strip()!r}: ожидается «тип=сумма», "
                "сумма в млн RUB. Поправьте настройку «Сколько отдано вложенным типам»."
            ) from exc
        if value <= 0:
            raise PortfolioDynamicsError(
                f"Подлимит для {name.strip()} должен быть больше нуля, задано {value}."
            )
        allocations[name.strip().upper()] = value
    return allocations


def aggregated_parents() -> Dict[str, str]:
    """Вложенности, объёмы которых СКЛАДЫВАЮТСЯ с объемлющим типом.

    Выделили вложенному типу собственный подлимит — его объём и лимит стоят
    отдельно, складывать нечего. Не выделили — объём идёт внутрь объемлющего
    и сравнивается с общим лимитом: так превышение общего лимита не потеряется.
    """
    allocations = parse_nested_limits()
    return {child: parent for child, parent in parse_type_parents().items()
            if child not in allocations}


def descendants_of(portfolio_type: str, parents: Dict[str, str]) -> List[str]:
    """Типы, объём которых входит в объём этого типа (на любую глубину)."""
    result, frontier = [], [str(portfolio_type).upper()]
    while frontier:
        current = frontier.pop()
        children = [c for c, p in parents.items() if p == current and c not in result]
        result.extend(children)
        frontier.extend(children)
    return result


def types_counted_in(portfolio_type: str, parents: Dict[str, str]) -> List[str]:
    """Сам тип плюс все вложенные — то, из чего складывается его объём."""
    return [str(portfolio_type).upper()] + descendants_of(portfolio_type, parents)


# Торговый тип называется TSS — так его зовут в казначействе и в отчёте
# старого формата. В CONTRACT.md и эталонном шаблоне он назван TTS, поэтому
# старое написание продолжает приниматься везде (в настройках, в файлах, в
# уже выпущенных отчётах) и приводится к каноническому в одном месте.
TYPE_SPELLINGS = {"TTS": "TSS"}

# Четыре согласованных типа (CONTRACT.md, п. 8.1). Известны всегда, даже когда
# файла лимитов нет: без этого AFS_OFZ_PD попадал бы под правило «OFZ_PD -> TTS»
# и уезжал в торговый портфель, хотя начинается с типа AFS.
KNOWN_PORTFOLIO_TYPES = ("AFS", "HTM", "HTM_KUAP", "TSS")


def canonical_type(name) -> str:
    """Каноническое написание типа: TSS и TTS — один и тот же торговый тип."""
    text = str(name).strip().upper()
    return TYPE_SPELLINGS.get(text, text)


def manual_portfolios() -> List[dict]:
    """Портфели, которых нет в выгрузке, но объём по ним ведётся вручную.

    Задаются в настройках («Дополнительные портфели»). Попадают и в справочник,
    и в срез, и в объём своего типа — иначе их объём выпал бы из светофора.
    """
    records = config.PORTFOLIO_DYNAMICS_MANUAL_PORTFOLIOS or []
    return [dict(r, type=canonical_type(r["type"])) for r in records]


def parse_type_rules(raw: Optional[str] = None) -> Dict[str, str]:
    """«HTM=HTM, OFZ_PD=TSS» -> {подстрока в коде: тип}, В ПОРЯДКЕ ЗАПИСИ.

    Порядок важен: срабатывает первое подходящее правило, поэтому более общие
    правила должны стоять ниже частных.
    """
    raw = config.PORTFOLIO_DYNAMICS_TYPE_RULES if raw is None else raw
    rules: Dict[str, str] = {}
    for item in str(raw or "").split(","):
        marker, _, portfolio_type = item.partition("=")
        if marker.strip() and portfolio_type.strip():
            rules[marker.strip().upper()] = canonical_type(portfolio_type)
    return rules


def parse_portfolio_overrides(raw: Optional[str] = None) -> Dict[str, str]:
    """«OFZ_SPECIAL=AFS» -> {код портфеля: тип}. Точечные исключения."""
    raw = config.PORTFOLIO_DYNAMICS_PORTFOLIO_TYPES if raw is None else raw
    overrides: Dict[str, str] = {}
    for item in str(raw or "").split(","):
        code, _, portfolio_type = item.partition("=")
        if code.strip() and portfolio_type.strip():
            overrides[code.strip().upper()] = canonical_type(portfolio_type)
    return overrides


def guess_type(portfolio_code: str, known_types: Iterable[str] = ()) -> str:
    """Тип портфеля по его коду. Порядок правил — от частного к общему.

    0. Файл разметки портфелей по типам (portfolio_types.json) — главный
       источник: что в нём указано, то и тип, без всяких догадок.
    1. Точечное исключение из настроек («Разметка отдельных портфелей»).
    2. Самый ДЛИННЫЙ известный тип, которым код начинается: HTM_KUAP_CORE ->
       HTM_KUAP, а не HTM. Это должно идти раньше правил по подстроке, иначе
       правило «HTM» перехватывало бы КУАП и уводило его объём в чужой тип.
    3. Правила по ВХОЖДЕНИЮ подстроки («Правила разметки по имени портфеля»):
       OFZ_HTM -> HTM, OFZ_PD/OFZ_PK/OFZ_CNY -> торговый. Срабатывает первое
       подходящее, поэтому порядок в настройке — это приоритет.
    4. Если ничего не подошло — префикс до первого «_» (AFS_TR_RUR -> AFS).
    """
    code = str(portfolio_code).upper()

    from reports.portfolio_dynamics import type_map
    mapped = type_map.explicit_type(code)
    if mapped:
        return mapped

    override = parse_portfolio_overrides().get(code)
    if override:
        return override

    candidates = {canonical_type(t) for t in known_types if str(t).strip()}
    candidates.update(KNOWN_PORTFOLIO_TYPES)
    matches = [t for t in candidates if code == t or code.startswith(t + "_")]
    if matches:
        return max(matches, key=len)

    for marker, portfolio_type in parse_type_rules().items():
        if marker in code:
            return portfolio_type

    # Канонизируем и здесь: код TTS_OFZ, оставшийся с прежнего написания,
    # должен дать тот же тип, что и TSS_OFZ.
    return canonical_type(code.split("_", 1)[0])


def _allowed_types() -> set:
    """Типы, которые бывают в отчёте: ключи файла разметки (TSS, AFS, HTM, HTM_KUAP)."""
    from reports.portfolio_dynamics import type_map
    return set(type_map.registry())


def _drop_unknown_types(previous: PreviousRelease) -> None:
    """Убирает из перенесённых лимитов и истории типы, которых в отчёте не бывает.

    Прежнее угадывание заводило типы по кодам портфелей («OFZ», «BOND») —
    с нулевым лимитом и своей строкой в истории. Раз перечень типов теперь
    задан файлом разметки, такие строки — мусор, из-за которого свод по типам
    показывает лишние строки, а проверки лимитов падают.
    """
    allowed = _allowed_types()
    dropped = set()
    for name in ("fact_limit", "fact_type_daily"):
        frame = getattr(previous, name)
        if frame.empty:
            continue
        types = frame["portfolio_type"].map(
            lambda t: canonical_type(t) if pd.notna(t) and str(t).strip() else "")
        keep = types.isin(allowed)
        if keep.all():
            continue
        dropped.update(types[~keep])
        setattr(previous, name, frame[keep].reset_index(drop=True))
    dropped.discard("")
    if dropped:
        logger.warning(
            "Из лимитов и истории предыдущего выпуска убраны типы, которых нет в файле "
            "разметки (%s): %s. Если какой-то из них настоящий — добавьте его ключом "
            "в файл разметки.", ", ".join(sorted(allowed)), ", ".join(sorted(dropped)),
        )


def _record_unmapped(dim: pd.DataFrame) -> None:
    """Портфели без явной разметки — в раздел «_не_размечены» файла разметки."""
    from reports.portfolio_dynamics import type_map
    guesses = {str(code): (str(t) if pd.notna(t) and str(t).strip() else None)
               for code, t in zip(dim["portfolio_code"], dim["portfolio_type"])
               if not type_map.explicit_type(code)}
    added = type_map.record_unmapped(guesses)
    if guesses:
        logger.warning(
            "Портфелей без явной разметки в файле %s: %d (новых %d) — тип им угадан "
            "по коду. Они перечислены в разделе «%s» с угаданным типом: перенесите "
            "каждый код в список своего типа.",
            type_map.types_file(), len(guesses), len(added), type_map.UNMAPPED_KEY,
        )


def _next_sort_order(dim: pd.DataFrame) -> int:
    values = pd.to_numeric(dim["sort_order"], errors="coerce").dropna()
    return int(values.max()) + 10 if len(values) else 10


def _build_dim(t0: PortfolioSlice, previous: PreviousRelease, bootstrap: bool,
               extra_types: Iterable[str] = ()) -> Tuple[pd.DataFrame, List[str]]:
    """Справочник портфелей: предыдущий как есть + новые коды из T0.

    Тип нового кода угадывается по префиксу, и только если такой тип уже заведён
    в fact_limit — иначе остаётся пустым. В любом случае привязку подтверждает
    человек, поэтому список новых кодов уходит в лог предупреждением.
    """
    dim = previous.dim_portfolio.copy()
    if dim.empty:
        dim = pd.DataFrame(columns=DIM_COLUMNS)
    known_types = set(previous.fact_limit["portfolio_type"].astype(str)) if not previous.fact_limit.empty else set()
    # Типы из файла лимитов — самый надёжный источник: это и есть реестр типов.
    known_types.update(str(t).upper() for t in extra_types if str(t).strip())
    # Какие типы вообще бывают, решает файл разметки (TSS, AFS, HTM, HTM_KUAP):
    # угаданное по коду «OFZ» или «BOND» типом не становится никогда.
    allowed = _allowed_types()
    if bootstrap:
        # Первый выпуск: реестра типов ещё нет, поэтому недостающие типы
        # достраиваются по кодам портфелей. Именно ДОБАВЛЯЮТСЯ, а не заменяют:
        # типы из настроек и файла лимитов точнее угадывания по префиксу.
        known_types.update(t for t in (guess_type(code, known_types)
                                       for code in t0.frame["portfolio_code"])
                           if t in allowed)

    known_codes = set(dim["portfolio_code"].astype(str))
    # Дополнительные портфели ведутся руками и в выгрузке не встречаются —
    # в справочник они должны попасть наравне с выгруженными.
    codes_in_report = list(t0.frame["portfolio_code"]) + [
        r["code"] for r in manual_portfolios()
    ]
    new_codes = []
    for code in codes_in_report:
        if code not in known_codes and code not in new_codes:
            new_codes.append(code)

    sort_order = _next_sort_order(dim)
    # У портфеля, заведённого вручную, название и тип известны точно — их
    # указал человек, гадать по коду не нужно.
    manual_by_code = {r["code"]: r for r in manual_portfolios()}
    additions = []
    for code in new_codes:
        manual = manual_by_code.get(code)
        guessed = guess_type(code, known_types)
        additions.append({
            "portfolio_code": code,
            "portfolio_name": manual["name"] if manual else code,
            "portfolio_type": manual["type"] if manual else (
                guessed if guessed in known_types and guessed in allowed else None),
            "include_in_total": True,
            "is_limit_controlled": True,
            "sort_order": sort_order,
        })
        sort_order += 10

    if additions:
        dim = pd.concat([dim, pd.DataFrame(additions, columns=DIM_COLUMNS)], ignore_index=True)

    dim["sort_order"] = pd.to_numeric(dim["sort_order"], errors="coerce")
    dim = dim.sort_values(["sort_order", "portfolio_code"], na_position="last").reset_index(drop=True)
    _apply_type_rules(dim, known_types)
    _apply_portfolio_overrides(dim)
    _apply_nesting_to_total_flag(dim)
    return dim[DIM_COLUMNS], new_codes


def _apply_type_rules(dim: pd.DataFrame, known_types) -> None:
    """Пересчитывает тип по правилам для ВСЕХ строк справочника, не только новых.

    Справочник переносится из предыдущего выпуска целиком, поэтому портфель,
    размеченный до появления правил, так и оставался бы с прежним типом:
    правило есть, а разметка «не чинится». Пересчёт делает правила
    действующими задним числом.

    Тип, которого нет среди известных, не проставляется: иначе странный код
    (SOMETHING_ELSE -> «SOMETHING») затирал бы уже верную разметку мусором.
    Точечные исключения применяются ПОСЛЕ и перебивают этот пересчёт.
    """
    if dim.empty:
        return
    allowed = _allowed_types()

    changes = []
    for index, row in dim.iterrows():
        derived = guess_type(row["portfolio_code"], known_types)
        current = canonical_type(row["portfolio_type"]) if str(row["portfolio_type"] or "").strip() else ""
        if derived not in allowed:
            # Правила ничего осмысленного не дали. Верную разметку не трогаем,
            # а тип, которого в отчёте не бывает, снимаем: пусть портфель
            # честно висит неразмеченным, чем уводит объём в несуществующий тип.
            if not current or current in allowed:
                continue
            derived = None
        if current == (derived or ""):
            continue
        changes.append(f"{row['portfolio_code']}: {current or 'пусто'} -> {derived or 'пусто'}")
        dim.at[index, "portfolio_type"] = derived

    if changes:
        logger.warning(
            "Разметка портфелей пересчитана по правилам, изменено %d: %s. "
            "Если какая-то строка размечена неверно, задайте её в настройке "
            "«Разметка отдельных портфелей» — она перебивает правила.",
            len(changes), "; ".join(changes),
        )


def _apply_portfolio_overrides(dim: pd.DataFrame) -> None:
    """Проставляет типы из точечных исключений, в том числе уже размеченным строкам.

    Иначе настройка работала бы только для новых кодов, а исправить уже
    неверно размеченный портфель можно было бы только правкой файла руками.
    Колонка ручная, поэтому каждое изменение пишется в лог.
    """
    overrides = parse_portfolio_overrides()
    if not overrides or dim.empty:
        return
    from reports.portfolio_dynamics import type_map
    for index, row in dim.iterrows():
        code = str(row["portfolio_code"]).upper()
        wanted = overrides.get(code)
        if wanted is None or str(row["portfolio_type"] or "").upper() == wanted:
            continue
        if type_map.explicit_type(code):
            continue  # файл разметки главнее настройки
        logger.warning(
            "Портфель %s размечен как %s по настройке «Разметка отдельных портфелей» "
            "(было: %s).", row["portfolio_code"], wanted, row["portfolio_type"] or "пусто",
        )
        dim.at[index, "portfolio_type"] = wanted


def _apply_nesting_to_total_flag(dim: pd.DataFrame) -> None:
    """Снимает «входит в итог» у портфелей вложенных типов.

    Объём вложенного типа уже посчитан внутри объемлющего, поэтому в ИТОГО по
    банку он попадать не должен — иначе КУАП сложится дважды (CONTRACT.md, п. 5
    про include_in_total и п. 8.2). Колонка ручная, поэтому каждое изменение
    пишется в лог: человек должен видеть, что скрипт тронул его лист.
    """
    parents = aggregated_parents()
    if not parents or dim.empty:
        return
    nested = dim["portfolio_type"].astype(str).str.upper().isin(parents)
    changed = dim.loc[nested & (dim["include_in_total"] != False)]  # noqa: E712
    if len(changed):
        logger.warning(
            "Снят флаг «Входит в итог» у %d портфелей вложенных типов (%s): их объём "
            "уже учтён в объемлющем типе, иначе ИТОГО по банку задвоится.",
            len(changed), ", ".join(str(c) for c in changed["portfolio_code"]),
        )
    dim.loc[nested, "include_in_total"] = False


def _build_limits(dim: pd.DataFrame, previous: PreviousRelease, business_date: dt.date,
                  bootstrap: bool) -> pd.DataFrame:
    """Лимиты переносятся из предыдущего выпуска ЦЕЛИКОМ: скрипт их не трогает.

    Исключение — первый запуск: там заводятся строки типов, встреченных в T0, с
    нулевыми лимитами, чтобы файл читался и витрины считались. Заполнить их
    руками обязан человек (об этом пишется предупреждение в лог).
    """
    if not bootstrap or not previous.fact_limit.empty:
        return previous.fact_limit[LIMIT_COLUMNS].copy()

    # Тип берётся из справочника, а не угадывается по коду заново: у портфеля,
    # заведённого вручную, он указан человеком, и повторное угадывание завело бы
    # в fact_limit лишний тип (OFZ_EXTRA -> «OFZ») с нулевым объёмом.
    types = {canonical_type(t) for t in dim["portfolio_type"].dropna() if str(t).strip()}
    allowed = _allowed_types()
    for code, portfolio_type in zip(dim["portfolio_code"], dim["portfolio_type"]):
        if not str(portfolio_type or "").strip():
            guessed = guess_type(code, types)
            if guessed in allowed:
                types.add(guessed)
    types = sorted(types)
    return pd.DataFrame(
        [{
            "portfolio_type": t, "limit_amount": 0, "green_max_util": 0,
            "yellow_max_util": 0, "red_max_util": 0,
            "valid_from": business_date, "updated_by": None, "limit_remaining": None,
        } for t in types],
        columns=LIMIT_COLUMNS,
    )


# Сколько дат в истории предыдущего выпуска означает «история уже накоплена».
# Тогда файл старого формата больше не читается: история переносится из
# выпуска в выпуск, а импорт был разовым. Порог — в датах, а не в строках:
# строк на дату столько, сколько типов, и «10 строк» — это всего 2–3 дня.
HISTORY_ACCUMULATED_DATES = 10


def history_dates(frame: pd.DataFrame) -> int:
    """Сколько разных дат в истории по типам."""
    return 0 if frame is None or frame.empty else int(frame["business_date"].nunique())


def _import_history(previous: PreviousRelease, history_path, force: bool = False) -> int:
    """Дополняет историю предыдущего выпуска строками из отчёта старого формата.

    Меняет previous.fact_type_daily на месте: дальше он идёт в _build_type_daily
    обычным путём, и импортированные даты ничем не отличаются от накопленных.

    Импорт РАЗОВЫЙ: когда в предыдущем выпуске уже накоплено
    HISTORY_ACCUMULATED_DATES дат, история берётся оттуда, а файл старого
    формата не читается вовсе — его можно спокойно удалить из загрузок, и
    отчёт не станет каждый день предупреждать, что файла нет. force — путь
    задан явно (--history или ответом в диалоге): тогда импорт выполняется
    всегда, а отсутствие файла — ошибка.

    Импорт внутри функции: history.py читает константы из etl.py, и импорт на
    уровне модуля замкнул бы их в кольцо.
    """
    if history_path is None or not str(history_path).strip():
        return 0

    accumulated = history_dates(previous.fact_type_daily)
    if not force and accumulated >= HISTORY_ACCUMULATED_DATES:
        logger.info(
            "История: в предыдущем выпуске уже %d дат — она переносится оттуда, файл "
            "старого формата (%s) не читается. Подтянуть его ещё раз: --history <файл>.",
            accumulated, Path(str(history_path)).name,
        )
        return 0

    from reports.portfolio_dynamics import history as history_module

    path = history_module.resolve_path(history_path, history_module.search_dirs())
    if path is None:
        if force:
            raise PortfolioDynamicsError(f"Файл с историей не найден: {history_path}.")
        logger.warning(
            "Файл с историей не найден: %s — история из отчёта старого формата не "
            "подтягивается, отчёт собирается без неё. Поправьте путь в «Настройки» → "
            "«история из старого отчёта» или очистите его.", history_path,
        )
        return 0
    history_path = path

    from reports.portfolio_dynamics import history as history_module

    imported = history_module.parse_history_file(Path(history_path))
    merged, added, skipped = history_module.merge_into(previous.fact_type_daily, imported)
    previous.fact_type_daily = merged
    if skipped:
        logger.info(
            "История: %d импортированных строк пропущено — эти даты уже есть в отчёте "
            "и перезаписи не подлежат.", skipped,
        )
    logger.info("История: из %s добавлено строк %d", Path(history_path).name, added)
    return added


def _parse_limits(limits_path: Optional[Path]):
    """Разбирает файл лимитов, если он есть. None — файла нет.

    Импорт внутри функции намеренно: limits.py читает константы из etl.py, и
    импорт на уровне модуля замкнул бы их в кольцо.
    """
    if limits_path is None:
        return None
    from reports.portfolio_dynamics import limits as limits_module
    return limits_module.parse_limits_file(Path(limits_path))


def _resolve_limits(dim: pd.DataFrame, previous: PreviousRelease, business_date: dt.date,
                    bootstrap: bool, limits_path: Optional[Path], parsed) -> pd.DataFrame:
    """Лимиты: из файла «Состояние лимитов», иначе — из предыдущего выпуска."""
    if parsed is None:
        return _build_limits(dim, previous, business_date, bootstrap)

    from reports.portfolio_dynamics import limits as limits_module

    known_types = [t for t in dim["portfolio_type"].dropna().unique() if str(t).strip()]
    if not previous.fact_limit.empty:
        known_types += [t for t in previous.fact_limit["portfolio_type"].dropna()]
    return limits_module.build_fact_limit(
        parsed, known_types, previous.fact_limit, business_date, Path(limits_path).stem,
    )


def _build_snapshot(t0: PortfolioSlice, t7: PortfolioSlice, business_date: dt.date,
                    previous: PreviousRelease) -> Tuple[pd.DataFrame, int]:
    """Срез пересобирается целиком; note_text возвращается на место по portfolio_code.

    Это единственная ручная колонка на машинном листе — без merge полная
    перезапись листа стёрла бы заметки казначейства (CONTRACT.md, п. 2).
    """
    snapshot = t0.frame.rename(columns={"volume": "volume_t0"}).copy()
    volumes_t7 = t7.frame[["portfolio_code", "volume"]].rename(columns={"volume": "volume_t7"})
    snapshot = snapshot.merge(volumes_t7, on="portfolio_code", how="left")
    snapshot = _add_manual_portfolios(snapshot, previous)
    snapshot = apply_kuap_durations(snapshot)
    snapshot.insert(0, "business_date", business_date)

    notes = previous.notes.copy()
    if not notes.empty:
        notes = notes.drop_duplicates(subset=["portfolio_code"], keep="last")
        notes["note_text"] = notes["note_text"].where(notes["note_text"].notna(), None)
        snapshot = snapshot.merge(notes, on="portfolio_code", how="left")
    else:
        snapshot["note_text"] = None

    restored = int(snapshot["note_text"].notna().sum())
    return snapshot[SNAPSHOT_COLUMNS], restored


def _add_manual_portfolios(snapshot: pd.DataFrame, previous: PreviousRelease) -> pd.DataFrame:
    """Дописывает в срез портфели, которые ведутся вручную.

    volume_t7 берётся из предыдущего выпуска: тогда правка объёма в настройках
    видна в отчёте как настоящая дельта. Предыдущего значения нет — берём тот же
    объём, чтобы дельта была нулевой, а не выдуманной (и чтобы CHK_13 не падал
    на пустом volume_t7).
    """
    records = manual_portfolios()
    if not records:
        return snapshot

    previous_t0 = {}
    if previous.snapshot_volumes is not None:
        previous_t0 = previous.snapshot_volumes

    existing = set(snapshot["portfolio_code"].astype(str))
    additions = []
    for record in records:
        if record["code"] in existing:
            logger.warning(
                "Портфель %s задан в «Дополнительных портфелях», но он есть и в выгрузке "
                "позиций — значение из настроек не применяется, взято из выгрузки.",
                record["code"],
            )
            continue
        additions.append({
            "portfolio_code": record["code"],
            "volume": record["volume"],
            "volume_t0": record["volume"],
            "volume_t7": previous_t0.get(record["code"], record["volume"]),
            "duration_current_yrs": record["duration"],
            # Пока конечная дюрация не читается, «Дюрация цель» — только КУАП;
            # копировать в неё текущую дюрацию значило бы выдать её за целевую.
            "duration_target_yrs": (
                record["duration"] if config.PORTFOLIO_DYNAMICS_READ_DURATION_END else None
            ),
        })

    if not additions:
        return snapshot
    logger.info("Дополнительных портфелей добавлено в срез: %d (%s)",
                len(additions), ", ".join(a["portfolio_code"] for a in additions))
    return pd.concat([snapshot, pd.DataFrame(additions)], ignore_index=True)


def parse_kuap_durations(raw: Optional[str] = None) -> Dict[str, float]:
    """«HTM_KUAP_CORE=3.5» -> {код портфеля: дюрация по КУАП, лет}."""
    raw = config.PORTFOLIO_DYNAMICS_KUAP_DURATIONS if raw is None else raw
    try:
        return settings.parse_durations(raw, "portfolio_dynamics_kuap_durations")
    except settings.SettingsError as exc:
        raise PortfolioDynamicsError(
            f"{exc} Поправьте настройку «Дюрации по КУАП»."
        ) from exc


def apply_kuap_durations(snapshot: pd.DataFrame) -> pd.DataFrame:
    """Записывает дюрацию, установленную КУАП, в duration_target_yrs.

    Значение из настройки перебивает то, что пришло из выгрузки или из
    «Дополнительных портфелей»: это явное решение КУАП. Портфели, которых
    в настройке нет, не трогаются.
    """
    durations = parse_kuap_durations()
    if not durations:
        return snapshot

    snapshot = snapshot.copy()
    codes = snapshot["portfolio_code"].astype(str).str.upper()
    applied = []
    for code, duration in durations.items():
        mask = codes == code
        if mask.any():
            snapshot.loc[mask, "duration_target_yrs"] = duration
            applied.append(code)
    unknown = [code for code in durations if code not in applied]

    if applied:
        logger.info("Дюрация по КУАП проставлена портфелям: %d (%s)",
                    len(applied), ", ".join(applied))
    if unknown:
        logger.warning(
            "В «Дюрациях по КУАП» заданы портфели, которых нет в срезе T0: %s — пропущены. "
            "Проверьте коды в настройке.", ", ".join(unknown),
        )
    return snapshot


def _build_type_daily(snapshot: pd.DataFrame, dim: pd.DataFrame, limits: pd.DataFrame,
                      previous: PreviousRelease, business_date: dt.date
                      ) -> Tuple[pd.DataFrame, int, int, int]:
    """История по типам: прошлые даты как есть + свежая свёртка T0 по типам.

    Идемпотентность: строки за business_date из предыдущего выпуска ЗАМЕНЯЮТСЯ,
    а не дополняются. История за прошлые даты не переписывается никогда — она
    существует только в этом файле, восстановить её из выгрузок невозможно.
    """
    history = previous.fact_type_daily.copy()
    replaced = 0
    if not history.empty:
        same_date = history["business_date"] == business_date
        replaced = int(same_date.sum())
        history = history[~same_date]
    carried = len(history)

    types_frame = snapshot.merge(
        dim[["portfolio_code", "portfolio_type"]], on="portfolio_code", how="left"
    )
    own = (
        types_frame[types_frame["portfolio_type"].notna()]
        .groupby("portfolio_type")["volume_t0"].sum()
    )
    # Лимит совокупный, поэтому и объём объемлющего типа — совокупный:
    # volume(HTM) = HTM + HTM_KUAP. Иначе светофор HTM занижал бы использование
    # и не показывал бы превышение общего лимита.
    parents = aggregated_parents()
    sums = {}
    for portfolio_type in set(own.index) | set(parents) | set(parents.values()):
        sums[portfolio_type] = float(sum(
            own.get(t, 0.0) for t in types_counted_in(portfolio_type, parents)
        ))
    sums = pd.Series(sums, dtype=float)

    # Одна строка на КАЖДЫЙ тип из fact_limit (в том числе на тип без портфелей):
    # так CHK_04 «строк на отчётную дату = число типов» остаётся выполнимой.
    known_types = [str(t) for t in limits["portfolio_type"]] if not limits.empty else []
    for extra in sums.index:
        if str(extra) not in known_types and sums.get(extra):
            known_types.append(str(extra))

    fresh = pd.DataFrame(
        [{
            "business_date": business_date,
            "portfolio_type": t,
            "volume_amount": float(sums.get(t, 0.0)),
        } for t in known_types],
        columns=TYPE_DAILY_COLUMNS,
    )

    combined = pd.concat([history, fresh], ignore_index=True) if carried else fresh
    combined = combined.sort_values(["business_date", "portfolio_type"]).reset_index(drop=True)
    return combined[TYPE_DAILY_COLUMNS], carried, len(fresh), replaced


def build_data(t0_path: Path, t7_path: Path, previous_path: Optional[Path] = None,
               bootstrap: bool = False, limits_path: Optional[Path] = None,
               history_path: Optional[Path] = None,
               force_history: bool = False) -> PortfolioDynamicsData:
    """Полный цикл ETL: два среза + предыдущий выпуск -> четыре таблицы схемы v3.0.

    limits_path — выгрузка «Состояние лимитов» на отчётную дату. Есть файл —
    лимиты и границы зон берутся из него; нет — переносятся из предыдущего
    выпуска, как было раньше.

    history_path — отчёт СТАРОГО формата, из которого один раз подтягивается
    уже накопленная история объёмов по типам. Импорт только дополняет: даты,
    накопленные своими запусками, не перезаписываются. Когда история в
    предыдущем выпуске уже накоплена, файл не читается (см. _import_history);
    force_history — путь указан явно, импорт выполняется всё равно.
    """
    t0 = parse_slice(t0_path, "T0")
    t7 = parse_slice(t7_path, "T-7")

    if t0.business_date is None:
        raise PortfolioDynamicsError(
            f"Не удалось определить дату среза T0: ни в имени файла {Path(t0_path).name}, "
            "ни в шапке листа нет строки вида «Позиция за период [дд.мм.гггг] - [дд.мм.гггг]». "
            "Переименуйте файл по образцу выгрузки."
        )
    business_date = t0.business_date

    if t7.business_date is not None and t7.business_date >= business_date:
        logger.warning(
            "Дата среза T-7 (%s) не раньше даты T0 (%s) — проверьте, что файлы не перепутаны.",
            t7.business_date.isoformat(), business_date.isoformat(),
        )
    lookback = (
        (business_date - t7.business_date).days
        if t7.business_date is not None and t7.business_date < business_date
        else config.PORTFOLIO_DYNAMICS_DEFAULT_LOOKBACK
    )

    if bootstrap:
        previous = PreviousRelease.empty()
    elif previous_path is not None:
        previous = load_previous_release(previous_path)
        logger.info("Предыдущий выпуск: %s", previous.path)
    else:
        raise PortfolioDynamicsError(
            "Предыдущий выпуск отчёта не найден. История объёмов по типам, лимиты и "
            "заметки существуют только в нём и из выгрузок не восстанавливаются. "
            "Укажите файл через --previous либо запустите первый выпуск с --bootstrap."
        )

    # Файл лимитов разбирается ДО справочника: он задаёт перечень типов, и без
    # него HTM_KUAP_CORE был бы отнесён к HTM (префикс до первого "_").
    parsed_limits = _parse_limits(limits_path)
    limit_types = list(parsed_limits["portfolio_type"]) if parsed_limits is not None else []
    # Типы, объявленные настройками, тоже известны: подлимит HTM_KUAP задаётся
    # руками, и строки HTM_KUAP в файле лимитов может не быть вовсе — без этого
    # HTM_KUAP_CORE был бы отнесён к HTM и его объём ушёл бы в чужой тип.
    limit_types += list(parse_nested_limits())
    limit_types += list(parse_type_parents()) + list(parse_type_parents().values())

    _drop_unknown_types(previous)
    imported_rows = _import_history(previous, history_path, force=force_history)
    dim, new_codes = _build_dim(t0, previous, bootstrap, extra_types=limit_types)
    _record_unmapped(dim)
    limits = _resolve_limits(dim, previous, business_date, bootstrap, limits_path, parsed_limits)
    snapshot, notes_restored = _build_snapshot(t0, t7, business_date, previous)
    history, carried, added, replaced = _build_type_daily(snapshot, dim, limits, previous, business_date)

    if parsed_limits is not None:
        from reports.portfolio_dynamics import limits as limits_module
        today_rows = history[history["business_date"] == business_date]
        limits_module.check_utilisation(
            parsed_limits,
            dict(zip(today_rows["portfolio_type"].astype(str), today_rows["volume_amount"])),
        )

    if imported_rows:
        from reports.portfolio_dynamics import history as history_module
        today = history[history["business_date"] == business_date]
        history_module.warn_if_scale_looks_wrong(
            previous.fact_type_daily,
            dict(zip(today["portfolio_type"].astype(str), today["volume_amount"])),
        )

    missing = sorted(set(dim["portfolio_code"].astype(str)) - set(snapshot["portfolio_code"].astype(str)))

    if bootstrap:
        logger.warning(
            "Первый выпуск (--bootstrap): лист fact_limit заведён с НУЛЕВЫМИ лимитами по "
            "%d типам (%s). Заполните лимиты и границы зон руками — до этого светофор на "
            "view_by_type и проверки CHK_19/CHK_23/CHK_24 работать не будут.",
            len(limits), ", ".join(str(t) for t in limits["portfolio_type"]),
        )
    if new_codes:
        logger.warning(
            "Новых портфелей в T0: %d (%s) — дописаны в dim_portfolio. Проверьте и "
            "подтвердите привязку к типу руками.",
            len(new_codes), ", ".join(new_codes),
        )
    if missing:
        logger.warning(
            "Портфелей в dim_portfolio, которых нет в срезе T0: %d (%s) — CHK_08 покажет FAIL, "
            "пока они не будут удалены из справочника или не вернутся в выгрузку.",
            len(missing), ", ".join(missing),
        )
    empty_types = dim["portfolio_type"].isna().sum() + (dim["portfolio_type"].astype(str) == "").sum()
    if empty_types:
        logger.warning(
            "Портфелей без типа в dim_portfolio: %d — их объём не попадёт в fact_type_daily.",
            int(empty_types),
        )

    logger.info(
        "История fact_type_daily: перенесено %d строк, дописано %d за %s (заменено %d), "
        "заметок сохранено %d, портфелей в срезе %d, сдвиг сравнения %d дн.",
        carried, added, business_date.isoformat(), replaced, notes_restored, len(snapshot), lookback,
    )

    return PortfolioDynamicsData(
        business_date=business_date,
        lookback_days=lookback,
        dim_portfolio=dim,
        fact_limit=limits,
        fact_type_daily=history,
        fact_portfolio_snapshot=snapshot,
        new_portfolio_codes=new_codes,
        missing_portfolio_codes=missing,
        history_rows_carried=carried,
        history_rows_added=added,
        history_rows_imported=imported_rows,
        history_rows_replaced=replaced,
        notes_restored=notes_restored,
    )
