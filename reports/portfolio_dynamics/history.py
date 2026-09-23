"""Импорт накопленной истории объёмов по типам из отчёта СТАРОГО формата.

Зачем. Лист fact_type_daily наполняется по одной дате за запуск, поэтому у
свежего файла истории нет — а она уже накоплена в отчёте, который вели до
перехода на схему v3.0. Этот модуль её оттуда переносит.

Старый формат: по листу на тип портфеля — «Динамика AFS», «Динамика HTM»,
«Динамика TSS» — и в каждом три колонки: «Дата», «Текущий объём»,
«Изменение». Нужны только первые две: изменение отчёт считает сам, хранить
его в данных незачем (и незачем спорить с ним при расхождении).

Объём в старом файле — СУММАРНЫЙ по типу за день, то есть ровно тот же грейн,
что и в fact_type_daily: (дата, тип) -> объём. Поэтому перенос — это перенос,
без пересчётов.

Правило слияния: импорт только ДОПОЛНЯЕТ. Дата, которая уже есть в истории,
не трогается никогда — накопленное своими запусками всегда важнее.
"""
import datetime as dt
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import pandas as pd

BASE_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE_DIR))

import config  # noqa: E402
from common import excel_io  # noqa: E402
from reports.portfolio_dynamics.etl import (  # noqa: E402
    TYPE_DAILY_COLUMNS, PortfolioDynamicsError, canonical_type, logger, parse_number,
)

SHEET_PREFIX = "Динамика"
COL_DATE = "Дата"
COL_VOLUME = "Текущий объём"
COL_CHANGE = "Изменение"  # намеренно не читается: дельту отчёт считает сам

# Форматы даты, встречающиеся в старом файле, если дата лежит текстом.
_DATE_FORMATS = ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%d_%m_%Y")


def parse_aliases(raw: Optional[str] = None) -> Dict[str, str]:
    """«TSS=TTS» -> {нормализованное имя из файла: тип портфеля}.

    Нужны, только если лист назван не так, как называется тип: написание TSS
    отчёт понимает и без псевдонимов.
    """
    raw = config.PORTFOLIO_DYNAMICS_HISTORY_ALIASES if raw is None else raw
    aliases: Dict[str, str] = {}
    for item in str(raw or "").split(","):
        name, _, portfolio_type = item.partition("=")
        if name.strip() and portfolio_type.strip():
            aliases[excel_io.normalize_label(name)] = portfolio_type.strip().upper()
    return aliases


def _type_from_sheet(sheet_name: str, aliases: Dict[str, str]) -> Optional[str]:
    """«Динамика AFS» -> AFS. Не наш лист -> None."""
    normalized = excel_io.normalize_label(sheet_name)
    prefix = excel_io.normalize_label(SHEET_PREFIX)
    if not normalized.startswith(prefix):
        return None
    raw_type = sheet_name.strip()[len(SHEET_PREFIX):].strip(" -_:")
    if not raw_type:
        return None
    return canonical_type(aliases.get(excel_io.normalize_label(raw_type), raw_type))


def _parse_date(value) -> Optional[dt.date]:
    """Дата из ячейки: настоящая дата Excel либо текст вида 09.09.2026."""
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _find_columns(matrix: pd.DataFrame) -> Optional[Tuple[int, int, int]]:
    """(строка шапки, колонка даты, колонка объёма) — по содержимому, не по номеру."""
    wanted_date = excel_io.normalize_label(COL_DATE)
    wanted_volume = excel_io.normalize_label(COL_VOLUME)
    for row_idx in range(len(matrix)):
        normalized = [excel_io.normalize_label(v) for v in matrix.iloc[row_idx]]
        if wanted_date in normalized and wanted_volume in normalized:
            return row_idx, normalized.index(wanted_date), normalized.index(wanted_volume)
    return None


def parse_history_file(path: Path, scale: Optional[float] = None,
                       aliases: Optional[Dict[str, str]] = None) -> pd.DataFrame:
    """Читает отчёт старого формата -> DataFrame[business_date, portfolio_type, volume_amount]."""
    found = resolve_path(path)
    if found is None:
        raise PortfolioDynamicsError(f"Файл с историей не найден: {path}")
    path = found

    scale = config.PORTFOLIO_DYNAMICS_HISTORY_SCALE if scale is None else scale
    aliases = parse_aliases() if aliases is None else aliases

    records: List[dict] = []
    seen_sheets: List[str] = []
    skipped_rows = 0
    for sheet_name, matrix in excel_io.iter_sheet_matrices(path):
        seen_sheets.append(sheet_name)
        portfolio_type = _type_from_sheet(sheet_name, aliases)
        if portfolio_type is None or matrix.empty:
            continue
        found = _find_columns(matrix)
        if found is None:
            logger.warning(
                "История: на листе %r нет колонок «%s» и «%s» — лист пропущен.",
                sheet_name, COL_DATE, COL_VOLUME,
            )
            continue

        header_row, date_col, volume_col = found
        rows_here = 0
        for row_idx in range(header_row + 1, len(matrix)):
            row = matrix.iloc[row_idx]
            business_date = _parse_date(row.iloc[date_col])
            volume = parse_number(row.iloc[volume_col])
            if business_date is None or volume is None:
                if str(row.iloc[date_col] or "").strip():
                    skipped_rows += 1
                continue
            records.append({
                "business_date": business_date,
                "portfolio_type": portfolio_type,
                "volume_amount": volume / scale,
            })
            rows_here += 1
        logger.info("История: лист %r -> тип %s, строк %d (шапка в строке %d)",
                    sheet_name, portfolio_type, rows_here, header_row + 1)

    if not records:
        raise PortfolioDynamicsError(
            f"В файле {path.name} не нашлось листов вида «{SHEET_PREFIX} <ТИП>» с "
            f"колонками «{COL_DATE}» и «{COL_VOLUME}». Листы в файле: {seen_sheets}."
        )

    frame = pd.DataFrame(records, columns=TYPE_DAILY_COLUMNS)
    duplicates = int(frame.duplicated(["business_date", "portfolio_type"]).sum())
    if duplicates:
        # Одна дата дважды на одном листе — берём последнюю: в старом отчёте
        # дописывали вниз, и нижняя строка свежее.
        frame = frame.drop_duplicates(["business_date", "portfolio_type"], keep="last")
        logger.warning("История: в файле %d дублирующихся пар (дата, тип) — взяты последние.",
                       duplicates)
    if skipped_rows:
        logger.info("История: пропущено строк без даты или объёма: %d", skipped_rows)

    frame = frame.sort_values(["business_date", "portfolio_type"]).reset_index(drop=True)
    logger.info(
        "История из %s: строк %d, типов %d (%s), период %s..%s",
        path.name, len(frame), frame["portfolio_type"].nunique(),
        ", ".join(sorted(frame["portfolio_type"].unique())),
        frame["business_date"].min().isoformat(), frame["business_date"].max().isoformat(),
    )
    return frame


def merge_into(existing: pd.DataFrame, imported: pd.DataFrame) -> Tuple[pd.DataFrame, int, int]:
    """Дополняет историю импортированными строками. Возвращает (итог, добавлено, пропущено).

    Импорт только ДОПОЛНЯЕТ: пара (дата, тип), уже накопленная своими запусками,
    не перезаписывается никогда — иначе один старый файл мог бы затереть то, что
    отчёт считал сам.
    """
    if imported is None or imported.empty:
        return existing, 0, 0
    if existing is None or existing.empty:
        return imported.copy(), len(imported), 0

    have = set(zip(existing["business_date"], existing["portfolio_type"].astype(str)))
    mask = [(d, str(t)) not in have
            for d, t in zip(imported["business_date"], imported["portfolio_type"])]
    fresh = imported[mask]
    skipped = len(imported) - len(fresh)
    if fresh.empty:
        return existing, 0, skipped

    combined = (pd.concat([existing, fresh], ignore_index=True)
                  .sort_values(["business_date", "portfolio_type"])
                  .reset_index(drop=True))
    return combined, len(fresh), skipped


def warn_if_scale_looks_wrong(imported: pd.DataFrame, current: Dict[str, float]) -> None:
    """Сверяет порядок величин импорта с объёмами, посчитанными по выгрузке.

    Единицы старого файла — догадка (настройка «Делитель истории»), и ошибка в
    ней в миллион раз не видна глазом в таблице на тысячи строк, зато мгновенно
    видна при сравнении с сегодняшним объёмом того же типа.
    """
    if imported is None or imported.empty or not current:
        return
    latest = (imported.sort_values("business_date")
                      .drop_duplicates("portfolio_type", keep="last")
                      .set_index("portfolio_type")["volume_amount"])
    for portfolio_type, value in latest.items():
        today = current.get(str(portfolio_type))
        if not today or not value:
            continue
        ratio = value / today
        if ratio > 100 or ratio < 0.01:
            logger.warning(
                "История: последний объём типа %s в импорте (%s) отличается от "
                "посчитанного по свежей выгрузке (%s) в %.0f раз — похоже, в файле "
                "другие единицы. Проверьте настройку «Делитель истории».",
                portfolio_type, f"{value:,.0f}", f"{today:,.0f}",
                ratio if ratio > 1 else 1 / ratio,
            )


# ── Поиск файла с историей ───────────────────────────────────────────────────
# Открывать книгу целиком, чтобы понять, наша ли она, незачем: имена листов
# лежат в оглавлении файла и читаются мгновенно. Поэтому кандидатов можно
# искать по СОДЕРЖИМОМУ, а не по имени — имя у старого отчёта произвольное
# («Динамика портфелей.xlsx», «Отчёт КУАП итог (2).xlsx»), и требовать от
# человека вспомнить его точно — ровно та причина, по которой возможность
# импорта до сих пор было не найти.
_EXCEL_SUFFIXES = (".xlsx", ".xlsm")
_MAX_FILES_TO_PROBE = 400


def history_sheets(path: Path) -> List[str]:
    """Листы вида «Динамика <ТИП>» в книге. Пусто — файл не тот (или не книга)."""
    path = Path(path)
    if path.suffix.lower() not in _EXCEL_SUFFIXES:
        return []
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            names = list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception:  # битый файл, чужой формат, нет прав — просто не наш
        return []
    prefix = excel_io.normalize_label(SHEET_PREFIX)
    return [name for name in names
            if excel_io.normalize_label(name).startswith(prefix) and len(name.strip()) > len(SHEET_PREFIX)]


def search_dirs() -> List[Path]:
    """Где искать отчёт старого формата: загрузки, папка отчёта, папка выгрузки."""
    return [Path(config.DOWNLOADS_DIR), Path(config.PORTFOLIO_DYNAMICS_DIR),
            Path(config.PORTFOLIO_DYNAMICS_OUTPUT_DIR)]


def resolve_path(raw, directories: Optional[List[Path]] = None) -> Optional[Path]:
    """Путь из настройки или аргумента -> существующий файл. None — не нашёлся.

    Проводник Windows по умолчанию скрывает расширения, и путь, набранный по
    тому, что видно на экране, приходит без «.xlsx» («…\\Лимиты портфелей
    казначейства (1)»). Поэтому, если файла ровно по такому пути нет, пробуем
    дописать расширение. Если нет и так — ищем файл с тем же именем в
    directories: книгу могли переложить из загрузок в папку отчёта.

    Расширение ДОПИСЫВАЕТСЯ, а не заменяется: у имени «отчёт 22.09» with_suffix
    отрезал бы «.09».
    """
    text = str(raw or "").strip().strip('"').strip("'")
    if not text:
        return None
    path = Path(text).expanduser()
    names = [path.name]
    if path.suffix.lower() not in _EXCEL_SUFFIXES:
        names += [path.name + suffix for suffix in _EXCEL_SUFFIXES]

    for name in names:
        candidate = path.with_name(name)
        if candidate.is_file():
            return candidate
    for directory in directories or []:
        if directory is None or not Path(directory).is_dir():
            continue
        for name in names:
            candidate = Path(directory) / name
            if candidate.is_file():
                return candidate
    return None


def find_candidates(directories: List[Path], limit: int = 10) -> List[Tuple[Path, List[str]]]:
    """Книги с листами «Динамика <ТИП>» в указанных папках, свежие сначала.

    Папки перебираются без рекурсии и в порядке, в котором переданы: загрузки,
    папка отчёта, папка выгрузки. Недоступная папка пропускается молча — это
    поиск, а не проверка путей.
    """
    seen = set()
    files: List[Path] = []
    for directory in directories:
        if directory is None:
            continue
        directory = Path(directory)
        if not directory.is_dir():
            continue
        try:
            entries = [f for f in directory.iterdir()
                       if f.is_file() and not f.name.startswith("~$")
                       and f.suffix.lower() in _EXCEL_SUFFIXES]
        except OSError:
            continue
        entries.sort(key=lambda f: f.stat().st_mtime, reverse=True)
        for entry in entries:
            resolved = entry.resolve()
            if resolved in seen:
                continue
            seen.add(resolved)
            files.append(entry)

    found: List[Tuple[Path, List[str]]] = []
    for path in files[:_MAX_FILES_TO_PROBE]:
        sheets = history_sheets(path)
        if sheets:
            found.append((path, sheets))
            if len(found) >= limit:
                break
    return found
