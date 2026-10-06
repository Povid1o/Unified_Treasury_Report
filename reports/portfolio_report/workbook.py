"""xlsx-витрина «Отчёта по портфелям» — место, где пишутся комментарии.

Устроено как в «Динамике портфелей»: рядом с CSV для BI каждый запуск пишет
otchet_po_portfelyam_<дата>.xlsx, строка на портфель, в конце жёлтая колонка
«Комментарий». Её правят руками. Следующий запуск берёт комментарии из самого
свежего xlsx в папке результатов (с датой позиций не позже своей — значит, и
из выпуска за ту же дату) и ставит их в новый выпуск и в CSV. Писать заново
не нужно; чтобы правка попала в CSV сразу, отчёт перезапускают за ту же дату.

Комментарий привязан к коду портфеля, а не к номеру строки: порядок строк
между выпусками меняется, а портфели появляются и пропадают.
"""
import datetime as dt
import re
import sys
from pathlib import Path
from typing import Dict, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(BASE_DIR))

from common.logging_utils import get_logger  # noqa: E402
from reports.portfolio_report import etl  # noqa: E402

logger = get_logger("portfolio_report")

SHEET = "Портфели"
CODE_HEADER = "Код"
COMMENT_HEADER = "Комментарий"
HEADER_ROW = 4

# (заголовок, колонка свода, делитель, формат). Витрина для людей: суммы в тех
# единицах, в которых их смотрят на дашборде. В CSV единицы свои (настройка
# «Округление»), xlsx на них не влияет.
COLUMNS = [
    ("Open QTY, шт", "open_qty", 1, "#,##0"),
    ("Изменение Open QTY, шт", "open_qty_change", 1, "+#,##0;-#,##0;0"),
    ("Чистая стоимость, млрд руб", "net_value", 1e9, "#,##0.000"),
    ("Total Full PL with Funding, млн руб", "total_pl", 1e6, "#,##0.0"),
    ("DV01, руб", "dv01", 1, "#,##0"),
    ("Yield, %", "yield", 1, "0.00"),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_COMMENT_HEADER_FILL = PatternFill("solid", fgColor="BF8F00")
_COMMENT_FILL = PatternFill("solid", fgColor="FFF2CC")
_WORKBOOK_DATE = re.compile(re.escape(etl.OUTPUT_FILENAME_PREFIX) + r"(\d{4}-\d{2}-\d{2})\.xlsx$")


def workbook_path(csv_path: Path) -> Path:
    return Path(csv_path).with_suffix(".xlsx")


def find_comments_release(output_dir: Path, upto: dt.date) -> Optional[Path]:
    """Самый свежий xlsx с датой позиций НЕ ПОЗЖЕ upto.

    Не позже, а не строго раньше (в отличие от сравнения Open QTY): при
    повторном прогоне за ту же дату комментарии, дописанные в сегодняшний
    выпуск, должны сохраниться.
    """
    output_dir = Path(output_dir)
    if not output_dir.is_dir():
        return None
    dated = []
    for path in output_dir.glob(f"{etl.OUTPUT_FILENAME_PREFIX}*.xlsx"):
        match = _WORKBOOK_DATE.search(path.name)
        if not match or not path.is_file() or path.name.startswith("~$"):
            continue
        try:
            when = dt.date.fromisoformat(match.group(1))
        except ValueError:
            continue
        if when <= upto:
            dated.append((when, path))
    return max(dated, key=lambda item: item[0])[1] if dated else None


def read_comments(path: Path) -> Dict[str, str]:
    """{код портфеля: комментарий} из витрины. Пустые комментарии не попадают."""
    path = Path(path)
    try:
        matrix = pd.read_excel(path, sheet_name=SHEET, header=None, dtype=object)
    except Exception as exc:
        raise etl.PortfolioReportError(
            f"Не удалось прочитать комментарии из {path.name}: {exc}") from exc

    def text(value) -> str:
        return "" if value is None or pd.isna(value) else str(value).strip()

    for row_idx in range(len(matrix)):
        row = [text(v) for v in matrix.iloc[row_idx]]
        if CODE_HEADER in row and COMMENT_HEADER in row:
            code_col, comment_col = row.index(CODE_HEADER), row.index(COMMENT_HEADER)
            break
    else:
        logger.warning("%s: на листе «%s» нет строки заголовков с «%s» и «%s» — "
                       "комментарии не перенесены.", path.name, SHEET, CODE_HEADER, COMMENT_HEADER)
        return {}

    comments: Dict[str, str] = {}
    for r in range(row_idx + 1, len(matrix)):
        code = text(matrix.iat[r, code_col]).upper()
        note = text(matrix.iat[r, comment_col])
        if code and note:
            comments[code] = note
    return comments


def write_workbook(data: "etl.PortfolioReportData", path: Path) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    ws["A1"] = f"Отчёт по портфелям — позиции на {data.business_date:%d.%m.%Y}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Комментарий пишите в жёлтой колонке: следующий запуск перенесёт его в новый "
                "выпуск и в CSV для BI. Чтобы правка попала в CSV этого выпуска, "
                "перезапустите отчёт за ту же дату.")
    ws["A2"].font = Font(italic=True, color="595959")

    headers = [CODE_HEADER, "Тип"] + [title for title, *_ in COLUMNS] + [COMMENT_HEADER]
    comment_col = len(headers)
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _COMMENT_HEADER_FILL if col == comment_col else _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    frame = data.frame.sort_values(["portfolio_type", "portfolio_code"]).reset_index(drop=True)
    for i, record in enumerate(frame.to_dict("records")):
        r = HEADER_ROW + 1 + i
        ws.cell(row=r, column=1, value=record["portfolio_code"])
        ws.cell(row=r, column=2, value=record["portfolio_type"])
        for j, (_title, column, divisor, fmt) in enumerate(COLUMNS, start=3):
            value = record[column]
            if value is None or pd.isna(value):
                continue
            cell = ws.cell(row=r, column=j, value=float(value) / divisor)
            cell.number_format = fmt
        note = ws.cell(row=r, column=comment_col,
                       value=data.comments.get(str(record["portfolio_code"]).upper()))
        note.fill = _COMMENT_FILL
        note.alignment = Alignment(wrap_text=True, vertical="top")

    widths = [22, 10] + [16] * len(COLUMNS) + [60]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[HEADER_ROW].height = 32
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=3)
    last = HEADER_ROW + max(len(frame), 1)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(comment_col)}{last}"

    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError as exc:
        raise etl.PortfolioReportError(
            f"Нет доступа для записи в {path} (файл открыт в Excel?): {exc}. "
            "Закройте файл и повторите запуск.") from exc
    logger.info("Витрина с комментариями сохранена: %s", path)
    return path
