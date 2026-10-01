"""Тесты настройки «Округление» (common/rounding.py) во всех отчётах.

Главное, что проверяется: пустая настройка ничего не меняет, а изменённая
меняет ровно свой показатель — значение, подпись единицы и знаки. Отдельно —
места, где отчёт читает собственный прошлый выпуск: значение, записанное в
другой единице, обязано вернуться в исходную, иначе изменение Open QTY или
история объёмов уедут на порядки.
"""
import datetime as dt
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import pandas as pd
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "tests"))

import config  # noqa: E402
from common import rounding, settings, settings_ui  # noqa: E402


class SettingsSandbox(unittest.TestCase):
    """Свой settings.json во временной папке — настройки разработчика не трогаем."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self._saved_env = os.environ.get(settings.SETTINGS_FILE_ENV)
        settings_file = self.tmp / "settings.json"
        settings_file.write_text("{}", encoding="utf-8")
        os.environ[settings.SETTINGS_FILE_ENV] = str(settings_file)
        config.reload()

    def tearDown(self):
        if self._saved_env is None:
            os.environ.pop(settings.SETTINGS_FILE_ENV, None)
        else:
            os.environ[settings.SETTINGS_FILE_ENV] = self._saved_env
        config.reload()
        self._tmp.cleanup()

    def set_rounding(self, report, value):
        settings.set_value(report + settings.ROUNDING_SUFFIX, value)
        config.reload()


# ════════════════════════════════════════════════════════════════════════════
# Правило и значение настройки
# ════════════════════════════════════════════════════════════════════════════
class RuleTests(unittest.TestCase):
    def cat(self, report, key):
        return rounding.CATEGORIES[report][key]

    def test_empty_entry_is_the_default_rule(self):
        for report, items in rounding.CATALOG.items():
            for category in items:
                self.assertTrue(rounding.make_rule(category, None).is_default)
                self.assertTrue(rounding.make_rule(category, {}).is_default)

    def test_default_rule_keeps_the_reports_own_rounding(self):
        rule = rounding.DEFAULT_RULE
        self.assertEqual(rule.apply(2.345, current=lambda v: round(v, 1)), 2.3)
        self.assertEqual(rule.apply(7), 7)
        self.assertEqual(rule.apply("—"), "—")

    def test_money_changes_value_and_label(self):
        rule = rounding.make_rule(self.cat("chpd", "volume"), {"scale": 6})
        self.assertEqual(rule.shift, 3)
        self.assertEqual(rule.label("млрд руб"), "млн руб")
        # «Как сейчас» у ЧПД — до целых: теперь целых миллионов.
        self.assertEqual(rule.apply(1.2345, current=lambda v: int(round(v))), 1234)

    def test_money_down_to_rubles_divides_exactly(self):
        rule = rounding.make_rule(self.cat("portfolio_report", "net_value"), {"scale": 3})
        self.assertEqual(rule.label("руб"), "тыс. руб")
        self.assertEqual(rule.apply(412_345_678.0), 412_345.678)

    def test_relative_divisor_and_label(self):
        rule = rounding.make_rule(self.cat("portfolio_report", "open_qty"), {"scale": 3})
        self.assertEqual(rule.label("шт"), "тыс. шт")
        self.assertEqual(rule.apply(1500.0), 1.5)
        # У ОВП подписи единицы нет — и не появляется.
        rule = rounding.make_rule(self.cat("ovp", "curr_balance"), {"scale": 6})
        self.assertIsNone(rule.label(None))

    def test_decimals(self):
        rule = rounding.make_rule(self.cat("nim", "nim"), {"decimals": 0})
        self.assertEqual(rule.apply(2.567, current=lambda v: round(v, 2)), 3)
        self.assertIsInstance(rule.apply(2.567), int)
        rule = rounding.make_rule(self.cat("nim", "nim"), {"decimals": 3})
        self.assertEqual(rule.apply(2.56789), 2.568)

    def test_text_and_nan_are_left_alone(self):
        rule = rounding.make_rule(self.cat("transfert", "rate"), {"decimals": 1})
        self.assertEqual(rule.apply("-"), "-")
        self.assertIsNone(rule.apply(None))
        self.assertTrue(pd.isna(rule.apply(float("nan"))))

    def test_label_is_read_back_to_the_base_unit(self):
        self.assertEqual(rounding.scale_of_label("portfolio_report", "open_qty", "тыс. шт"), 3)
        self.assertEqual(rounding.scale_of_label("portfolio_report", "open_qty", "шт"), 0)
        self.assertEqual(rounding.scale_of_label("portfolio_report", "rwa", "млн руб"), 6)
        self.assertEqual(rounding.scale_of_label("portfolio_report", "rgbi", "пункты"), 0)
        self.assertIsNone(rounding.scale_of_label("portfolio_report", "open_qty", "кг"))


class ParseTests(unittest.TestCase):
    def test_as_now_entries_are_dropped(self):
        self.assertEqual(rounding.parse("chpd", {"volume": {"scale": 9}, "share": {}}), {})
        self.assertEqual(rounding.parse("chpd", ""), {})

    def test_unknown_categories_are_ignored(self):
        self.assertEqual(rounding.parse("chpd", {"no_such": {"decimals": 1},
                                                 "share": {"decimals": 1}}),
                         {"share": {"decimals": 1}})

    def test_json_string_is_accepted(self):
        self.assertEqual(rounding.parse("chpd", '{"volume": {"scale": 6}}'),
                         {"volume": {"scale": 6}})

    def test_invalid_values_are_rejected(self):
        bad = [
            {"volume": {"scale": 5}},              # нет такой единицы
            {"share": {"scale": 3}},               # у процентов единица не меняется
            {"volume": {"decimals": 11}},
            {"volume": {"decimals": -1}},
            {"volume": {"decimals": True}},
            {"volume": {"unit": "млн"}},
            {"volume": 6},
            [1, 2],
            "{not json",
        ]
        for value in bad:
            with self.subTest(value=value), self.assertRaises(rounding.RoundingError):
                rounding.parse("chpd", value)


class SettingsRegistryTests(SettingsSandbox):
    def test_every_report_has_a_rounding_item_and_it_is_empty_by_default(self):
        for group in settings.GROUPS:
            if group.key == "common":
                continue
            key = group.key + settings.ROUNDING_SUFFIX
            self.assertIn(key, settings.SETTINGS_BY_KEY, group.key)
            self.assertEqual(settings.get(key), {})
            self.assertEqual(config.ROUNDING[group.key], {})

    def test_every_category_is_default_without_settings(self):
        for report, items in rounding.CATALOG.items():
            for category in items:
                self.assertTrue(rounding.rule(report, category.key).is_default,
                                (report, category.key))

    def test_saved_value_reaches_config_and_survives_reload(self):
        self.set_rounding("chpd", {"volume": {"scale": 6, "decimals": 1}})
        settings.reload()
        config.reload()
        self.assertEqual(rounding.rule("chpd", "volume"),
                         rounding.Rule(shift=3, decimals=1, unit="млн руб"))

    def test_invalid_value_is_a_settings_error(self):
        with self.assertRaises(settings.SettingsError):
            settings.set_value("chpd_rounding", {"volume": {"scale": 7}})


class RoundingScreenTests(SettingsSandbox):
    def run_screen(self, answers):
        setting = settings.SETTINGS_BY_KEY["chpd_rounding"]
        with mock.patch.object(settings_ui.ui, "ask", side_effect=answers), \
                mock.patch.object(settings_ui.ui.console, "print"), \
                mock.patch.object(settings_ui.ui, "success"):
            return settings_ui._rounding_screen(setting)

    def test_unit_and_decimals_are_saved(self):
        # Показатель 1 (объёмы) -> единица 3 (млн руб) -> 1 знак; выход.
        self.assertTrue(self.run_screen(["1", "3", "1", "0"]))
        self.assertEqual(settings.get("chpd_rounding"), {"volume": {"scale": 6, "decimals": 1}})

    def test_percent_asks_only_decimals(self):
        self.run_screen(["2", "2", "0"])
        self.assertEqual(settings.get("chpd_rounding"), {"share": {"decimals": 2}})

    def test_reset_one_and_all(self):
        self.run_screen(["1", "3", "", "2", "1", "0"])
        self.assertEqual(set(settings.get("chpd_rounding")), {"volume", "share"})
        self.run_screen(["-1", "0"])
        self.assertEqual(settings.get("chpd_rounding"), {"share": {"decimals": 1}})
        self.run_screen(["с", "0"])
        self.assertEqual(settings.get("chpd_rounding"), {})
        self.assertFalse(settings.is_overridden("chpd_rounding"))


# ════════════════════════════════════════════════════════════════════════════
# Отчёты
# ════════════════════════════════════════════════════════════════════════════
def write_chpd(path: Path) -> Path:
    from reports.chpd import etl as chpd
    wb = Workbook()
    ws = wb.active
    ws.title = chpd.SHEET_NAME
    dates = ["20.05.2026", None, "10.06.2026", None, "Изм", None]
    for j in range(12):
        ws.cell(row=1, column=2 + j, value=dates[j % 6])
    r = 4
    for k, (leaf, *_rest) in enumerate(chpd.INDEX_PATHS):
        ws.cell(row=r, column=1, value=leaf)
        for j in range(12):
            ws.cell(row=r, column=2 + j, value=(k + 1.2345) if j % 2 == 0 else 0.1834)
        r += 1
    wb.save(path)
    return path


class CsvReportsTests(SettingsSandbox):
    def test_chpd(self):
        from reports.chpd import etl as chpd
        path = write_chpd(self.tmp / "ЧПД 2026 06 10.xlsx")
        before = chpd.build_report(path)
        self.set_rounding("chpd", {"volume": {"scale": 6}, "share": {"decimals": 1}})
        after = chpd.build_report(path)
        volumes = after[after["axis_5"] != "%"]
        self.assertEqual(set(volumes["axis_5"]), {"млн руб"})
        self.assertEqual(volumes["value"].iloc[0], 1234)        # 1.2345 млрд -> 1234 млн
        self.assertEqual(before[before["axis_5"] != "%"]["value"].iloc[0], 1)
        self.assertEqual(after[after["axis_5"] == "%"]["value"].iloc[0], 18.3)

    def test_integers_stay_integers_in_csv(self):
        """0 знаков у одного показателя и дробные у соседнего в той же колонке —
        целые не должны уйти в CSV как «6660.0»."""
        from reports.chpd import etl as chpd
        path = write_chpd(self.tmp / "ЧПД 2026 06 10.xlsx")
        self.set_rounding("chpd", {"volume": {"scale": 6}, "share": {"decimals": 1}})
        out = chpd.save_report(chpd.build_report(path), self.tmp / "chpd.csv")
        lines = out.read_text(encoding="utf-8").splitlines()
        self.assertTrue(lines[1].endswith(",1234,,млн руб"), lines[1])
        self.assertTrue(lines[2].endswith(",18.3,,%"), lines[2])

    def test_nim(self):
        from reports.nim import etl as nim
        wb = Workbook()
        ws = wb.active
        rows = [(nim.MARKER_OSNOVA, "202501"), ("NIM", 0.025678), ("NIM", 95.5),
                ("NIM", 0.03), ("NIM", 0.01)]
        for i, (label, value) in enumerate(rows, start=1):
            ws.cell(row=i, column=1, value=label)
            ws.cell(row=i, column=2, value=value)
        path = self.tmp / "NIM_2025_01.xlsx"
        wb.save(path)
        self.assertEqual(nim.build_report(path)["value"].iloc[0], 2.57)
        self.set_rounding("nim", {"nim": {"decimals": 4}})
        frame = nim.build_report(path)
        self.assertEqual(frame["value"].iloc[0], 2.5678)
        self.assertEqual(frame[frame["axis_1"] == "% результат"]["value"].iloc[0], 95.5)

    def test_balance_struct(self):
        from reports.balance_struct import etl as bs
        wb = Workbook()
        ws = wb.active
        ws.title = bs.SHEET_NAME
        ws.cell(row=1, column=1, value="Структура баланса, млрд руб.")
        ws.cell(row=2, column=2, value="01.01.2026")
        ws.cell(row=3, column=1, value="Активы")
        ws.cell(row=3, column=2, value=100.4567)
        path = self.tmp / "ПФ.xlsx"
        wb.save(path)
        self.assertEqual(bs.build_report(path)["value"].iloc[0], 100)
        self.set_rounding("balance_struct", {"value": {"scale": 6}})
        self.assertEqual(bs.build_report(path)["value"].iloc[0], 100457)

    def test_ovp(self):
        from reports.ovp import etl as ovp
        import test_ovp_etl as tovp
        before = ovp._convert_sheet_matrices([("Свод", tovp._consolidated_frame())])
        self.set_rounding("ovp", {"curr_balance": {"scale": 3, "decimals": 2}})
        after = ovp._convert_sheet_matrices([("Свод", tovp._consolidated_frame())])
        expected = [round(v / 1000, 2) for v in before["curr_balance"]]
        self.assertEqual(list(after["curr_balance"]), expected)
        self.assertEqual(list(after["reserve_msfo"]), list(before["reserve_msfo"]))

    def test_transfert_leaves_text_alone(self):
        from reports.transfert_stavka import etl as ts
        frame = pd.DataFrame({c: ["x", "x"] for c in ts.OUT_COLUMNS})
        frame["fvalue"] = [15.12345, "-"]
        self.set_rounding("transfert", {"rate": {"decimals": 2}})
        with mock.patch.object(ts, "build_short_report", return_value=frame.iloc[:0]), \
                mock.patch.object(ts, "build_long_report", return_value=frame):
            result = ts.build_report(Path("a"), Path("b"))
        self.assertEqual(list(result["fvalue"]), [15.12, "-"])

    def test_ofz(self):
        from reports.ofz_rates import etl as ofz

        class Api:
            def get_index_values(self, *args, **kwargs):
                return [{"date": "2026-07-06", "value": "16.7149"}]

        self.assertEqual(ofz.fetch_yield_curve_group(Api(), "a", "b")["fvalue"].iloc[0], 16.71)
        self.set_rounding("ofz", {"yield": {"decimals": 3}})
        self.assertEqual(ofz.fetch_yield_curve_group(Api(), "a", "b")["fvalue"].iloc[0], 16.715)


class PortfolioReportTests(SettingsSandbox):
    def test_units_change_and_open_qty_survives_the_previous_release(self):
        from reports.portfolio_report import etl as pr
        import test_portfolio_report as tpr
        d1 = tpr.write_export(self.tmp / tpr.export_name("01.09.2026"), tpr.DAY1, "01.09.2026")
        d2 = tpr.write_export(self.tmp / tpr.export_name("02.09.2026"), tpr.DAY2, "02.09.2026")
        self.set_rounding("portfolio_report", {"open_qty": {"scale": 3},
                                               "net_value": {"scale": 6, "decimals": 1}})
        r1 = pr.save_report(pr.build_data(d1), self.tmp / "otchet_po_portfelyam_2026-09-01.csv")
        flat = pd.read_csv(r1, encoding="utf-8-sig", keep_default_na=False)
        row = flat[(flat["axis_2"] == "AFS_TR_RUR") & (flat["axis_3"] == "Open QTY")].iloc[0]
        self.assertEqual((float(row["value"]), row["axis_4"]), (0.15, "тыс. шт"))
        row = flat[(flat["axis_2"] == "AFS_TR_RUR") & (flat["axis_3"] == "Чистая стоимость")].iloc[0]
        self.assertEqual((float(row["value"]), row["axis_4"]), (400.0, "млн руб"))

        # Вчерашний Open QTY записан в тыс. шт — изменение всё равно в штуках: 170 - 150.
        data = pr.build_data(d2, previous_path=r1)
        change = data.frame.set_index("portfolio_code").loc["AFS_TR_RUR", "open_qty_change"]
        self.assertEqual(change, 20)


class PortfolioDynamicsTests(SettingsSandbox):
    def setUp(self):
        super().setUp()
        from test_workbook_formulas import build_demo_data
        self.data = build_demo_data()

    def test_csv_units_are_independent_of_xlsx(self):
        from reports.portfolio_dynamics import flat
        self.set_rounding("portfolio_dynamics", {"flat_t0": {"scale": 6},
                                                 "flat_delta_pct": {"decimals": 2}})
        frame = flat.to_flat(self.data)
        t0 = frame[(frame["axis_2"] == "AFS_OFZ") & (frame["axis_3"] == flat.M_T0)].iloc[0]
        self.assertEqual((t0["value"], t0["axis_4"]), (30_000.0, "млн RUB"))
        t7 = frame[(frame["axis_2"] == "AFS_OFZ") & (frame["axis_3"] == flat.M_T7)].iloc[0]
        self.assertEqual(t7["axis_4"], "млрд RUB")
        pct = frame[(frame["axis_2"] == "AFS_OFZ") & (frame["axis_3"] == flat.M_DELTA_PCT)].iloc[0]
        self.assertEqual(pct["value"], 2.04)

    def test_xlsx_in_thousands_round_trips_through_the_next_release(self):
        from reports.portfolio_dynamics import etl, workbook
        self.set_rounding("portfolio_dynamics", {"xlsx_amounts": {"scale": 3, "decimals": 0},
                                                 "xlsx_percent": {"decimals": 2}})
        path = self.tmp / "dinamika_portfeley_2026-09-18.xlsx"
        workbook.save_workbook(self.data, path)
        wb = load_workbook(path)
        self.assertEqual(wb.defined_names["AMOUNT_UNIT"].attr_text, '"тыс. RUB"')
        snap = wb["fact_portfolio_snapshot"]
        self.assertEqual(snap["C2"].value, 30_000_000)            # 30 000 млн = 30 000 000 тыс.
        self.assertEqual(snap["C2"].number_format, "#,##0;(#,##0);-")
        self.assertEqual(wb["view_monitor"]["G6"].number_format, "0.00%;(0.00%);-")
        self.assertIn("тыс. RUB", wb["README"]["A2"].value)
        self.assertNotIn("млрд RUB", wb["README"]["A2"].value)

        # Следующий запуск читает выпуск обратно в млн — история не съезжает в 1000 раз.
        previous = etl.load_previous_release(path)
        restored = previous.fact_type_daily.sort_values(["business_date", "portfolio_type"])
        original = self.data.fact_type_daily.sort_values(["business_date", "portfolio_type"])
        for got, want in zip(restored["volume_amount"], original["volume_amount"]):
            self.assertAlmostEqual(got, want, places=6)

    def test_default_xlsx_formats_are_untouched(self):
        from reports.portfolio_dynamics import workbook
        path = self.tmp / "d.xlsx"
        workbook.save_workbook(self.data, path)
        wb = load_workbook(path)
        self.assertEqual(wb.defined_names["AMOUNT_UNIT"].attr_text, '"млрд RUB"')
        self.assertEqual(wb["fact_portfolio_snapshot"]["C2"].number_format, workbook.FMT_AMT)
        self.assertEqual(wb["view_monitor"]["G6"].number_format, workbook.FMT_PCT)


if __name__ == "__main__":
    unittest.main()
