"""Тесты портфелей, которых нет в выгрузке, но объём по ним ведётся вручную.

Задаются в настройках («Дополнительные портфели»). Проверяется, что они
доходят до всех трёх мест сразу — справочник, срез и объём своего типа, — и
что правка объёма даёт настоящую дельту, а не ноль.
"""
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "tests"))

import config  # noqa: E402
from common import settings  # noqa: E402
from reports.portfolio_dynamics import etl, workbook  # noqa: E402
from test_portfolio_dynamics_etl import MLN, export_name, write_export  # noqa: E402

EXTRA = {"code": "OFZ_EXTRA", "name": "Внебиржевой ОФЗ", "type": "HTM",
         "volume": 150, "duration": 4.2}


class ManualPortfolioTestCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.out = self.tmp / "out"
        self._saved_env = os.environ.get(settings.SETTINGS_FILE_ENV)
        settings_file = self.tmp / "settings.json"
        settings_file.write_text(json.dumps({
            "portfolio_dynamics_dir": str(self.tmp / "data"),
            "portfolio_dynamics_output_dir": str(self.out),
            "downloads_dir": str(self.tmp / "downloads"),
        }), encoding="utf-8")
        os.environ[settings.SETTINGS_FILE_ENV] = str(settings_file)
        config.reload()

        codes = [("AFS_TR_RUR", 300), ("HTM_GOV", 600), ("OFZ_PD", 200)]

        def rows(scale):
            out = []
            for code, volume in codes:
                out += [(f"Позиция: {code}", None, None, None),
                        ("Bond", volume * scale * MLN, 3.0, 2.5)]
            return out

        self.t0 = write_export(self.tmp / export_name("21.09.2026"), rows(1.0),
                               period_end="21.09.2026")
        self.t7 = write_export(self.tmp / export_name("14.09.2026"), rows(0.97),
                               period_end="14.09.2026")

    def tearDown(self):
        if self._saved_env is None:
            os.environ.pop(settings.SETTINGS_FILE_ENV, None)
        else:
            os.environ[settings.SETTINGS_FILE_ENV] = self._saved_env
        config.reload()
        self._tmp.cleanup()

    def set_extra(self, *records):
        settings.set_value("portfolio_dynamics_manual_portfolios", list(records))
        config.reload()

    def build(self, previous=None):
        return etl.build_data(self.t0, self.t7, previous_path=previous,
                              bootstrap=previous is None)

    def volumes_by_type(self, data):
        today = data.fact_type_daily[
            data.fact_type_daily["business_date"] == data.business_date]
        return dict(zip(today["portfolio_type"], today["volume_amount"]))


class InTheReportTests(ManualPortfolioTestCase):
    def test_manual_portfolio_reaches_the_directory(self):
        self.set_extra(EXTRA)
        row = self.build().dim_portfolio.set_index("portfolio_code").loc["OFZ_EXTRA"]

        self.assertEqual(row["portfolio_name"], "Внебиржевой ОФЗ")
        self.assertEqual(row["portfolio_type"], "HTM", "тип указан человеком, гадать не нужно")

    def test_manual_portfolio_reaches_the_snapshot(self):
        self.set_extra(EXTRA)
        row = self.build().fact_portfolio_snapshot.set_index("portfolio_code").loc["OFZ_EXTRA"]

        self.assertEqual(row["volume_t0"], 150)
        self.assertEqual(row["duration_current_yrs"], 4.2)
        self.assertTrue(pd.isna(row["duration_target_yrs"]),
                        "текущая дюрация не выдаётся за целевую (КУАП)")

    def test_kuap_duration_reaches_a_manual_portfolio(self):
        self.set_extra(EXTRA)
        settings.set_value("portfolio_dynamics_kuap_durations", "OFZ_EXTRA=5")
        config.reload()
        row = self.build().fact_portfolio_snapshot.set_index("portfolio_code").loc["OFZ_EXTRA"]
        self.assertEqual(row["duration_target_yrs"], 5.0)

    def test_manual_volume_counts_towards_its_type(self):
        """Иначе объём выпал бы из светофора своего типа."""
        self.set_extra(EXTRA)
        self.assertEqual(self.volumes_by_type(self.build())["HTM"], 750)  # 600 + 150

    def test_without_manual_portfolios_nothing_changes(self):
        data = self.build()
        self.assertNotIn("OFZ_EXTRA", set(data.dim_portfolio["portfolio_code"]))
        self.assertEqual(self.volumes_by_type(data)["HTM"], 600)

    def test_no_phantom_type_is_created(self):
        """Тип берётся из справочника: повторное угадывание по коду завело бы
        в fact_limit лишний «OFZ» с нулевым объёмом."""
        self.set_extra(EXTRA)
        data = self.build()
        self.assertNotIn("OFZ", set(data.fact_limit["portfolio_type"]))
        self.assertNotIn("OFZ", set(data.fact_type_daily["portfolio_type"]))

    def test_several_manual_portfolios(self):
        self.set_extra(EXTRA, {"code": "AFS_EXTRA", "name": "Прочее AFS", "type": "AFS",
                               "volume": 25, "duration": 1.5})
        volumes = self.volumes_by_type(self.build())
        self.assertEqual(volumes["HTM"], 750)
        self.assertEqual(volumes["AFS"], 325)

    def test_grain_check_stays_clean(self):
        self.set_extra(EXTRA)
        data = self.build()
        failed = {cid for cid, status, _v in workbook.evaluate_checks(data) if status == "FAIL"}
        self.assertNotIn("CHK_16", failed)
        self.assertNotIn("CHK_08", failed)

    def test_code_present_in_the_export_is_not_overridden(self):
        """Выгрузка авторитетнее ручного значения — но человек должен об этом узнать."""
        self.set_extra({"code": "HTM_GOV", "name": "Дубль", "type": "HTM",
                        "volume": 1, "duration": 1})

        with self.assertLogs("portfolio_dynamics", level="WARNING") as captured:
            data = self.build()

        row = data.fact_portfolio_snapshot.set_index("portfolio_code").loc["HTM_GOV"]
        self.assertEqual(row["volume_t0"], 600)
        self.assertTrue(any("HTM_GOV" in line for line in captured.output))


class VolumeChangeTests(ManualPortfolioTestCase):
    def test_changed_volume_gives_a_real_delta(self):
        """Объём T-7 берётся из прошлого выпуска, поэтому правка настройки видна
        в отчёте как настоящее изменение, а не как ноль."""
        self.set_extra(EXTRA)
        first = self.build()
        released = workbook.save_workbook(first, self.out / "dinamika_portfeley_prev.xlsx")

        self.set_extra(dict(EXTRA, volume=180))
        second = self.build(previous=released)

        row = second.fact_portfolio_snapshot.set_index("portfolio_code").loc["OFZ_EXTRA"]
        self.assertEqual(row["volume_t0"], 180)
        self.assertEqual(row["volume_t7"], 150)

    def test_first_release_has_no_previous_value(self):
        """Прошлого значения нет — дельта нулевая, а не выдуманная."""
        self.set_extra(EXTRA)
        row = self.build().fact_portfolio_snapshot.set_index("portfolio_code").loc["OFZ_EXTRA"]
        self.assertEqual(row["volume_t7"], 150)

    def test_note_still_survives_a_rerun(self):
        self.set_extra(EXTRA)
        released = workbook.save_workbook(self.build(), self.out / "dinamika_portfeley_prev.xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(released)
        ws = wb["fact_portfolio_snapshot"]
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=2).value == "OFZ_EXTRA":
                ws.cell(row=row, column=7, value="Ведём вручную")
        wb.save(released)

        data = self.build(previous=released)
        row = data.fact_portfolio_snapshot.set_index("portfolio_code").loc["OFZ_EXTRA"]
        self.assertEqual(row["note_text"], "Ведём вручную")


class ValidationTests(ManualPortfolioTestCase):
    def set_raw(self, records):
        settings.set_value("portfolio_dynamics_manual_portfolios", records)

    def test_code_and_type_are_required(self):
        with self.assertRaisesRegex(settings.SettingsError, "код"):
            self.set_raw([{"type": "HTM", "volume": 1}])
        with self.assertRaisesRegex(settings.SettingsError, "тип"):
            self.set_raw([{"code": "X", "volume": 1}])

    def test_volume_must_be_a_number(self):
        with self.assertRaisesRegex(settings.SettingsError, "объём"):
            self.set_raw([{"code": "X", "type": "HTM", "volume": "много"}])

    def test_duplicate_codes_are_refused(self):
        with self.assertRaisesRegex(settings.SettingsError, "дважды"):
            self.set_raw([dict(EXTRA), dict(EXTRA)])

    def test_name_defaults_to_the_code(self):
        self.set_raw([{"code": "X_ONE", "type": "HTM", "volume": 5}])
        self.assertEqual(settings.get("portfolio_dynamics_manual_portfolios")[0]["name"], "X_ONE")

    def test_numbers_with_spaces_and_commas_are_understood(self):
        self.set_raw([{"code": "X", "type": "HTM", "volume": "1 500,5", "duration": "3,25"}])
        record = settings.get("portfolio_dynamics_manual_portfolios")[0]
        self.assertEqual(record["volume"], 1500.5)
        self.assertEqual(record["duration"], 3.25)

    def test_duration_is_optional(self):
        self.set_raw([{"code": "X", "type": "HTM", "volume": 5}])
        self.assertIsNone(settings.get("portfolio_dynamics_manual_portfolios")[0]["duration"])

    def test_old_spelling_of_the_trading_type_is_accepted(self):
        self.set_extra(dict(EXTRA, type="TTS"))
        row = self.build().dim_portfolio.set_index("portfolio_code").loc["OFZ_EXTRA"]
        self.assertEqual(row["portfolio_type"], "TSS")

    def test_records_survive_a_reload_from_disk(self):
        self.set_extra(EXTRA)
        settings.reload()
        self.assertEqual(settings.get("portfolio_dynamics_manual_portfolios")[0]["code"],
                         "OFZ_EXTRA")


if __name__ == "__main__":
    unittest.main()
