"""Тесты разбора выгрузки «Состояние лимитов» и сборки листа fact_limit.

Проверяется то, что тихо испортит светофор: выбор из двух строк AFS, перевод
«Облигации» в TTS, единицы (рубли -> млн), проценты зон и то, что неполный файл
не обнуляет уже согласованные лимиты.
"""
import datetime as dt
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "tests"))

import config  # noqa: E402
from common import settings  # noqa: E402
from reports.portfolio_dynamics import etl, limits  # noqa: E402

RUB = 1_000_000  # рублей в одном млн: выгрузка лимитов приходит в рублях

LIMITS_HEADER = ["Подразделение", "Тип лимита", "Валюта",
                 "Лимит снизу", "Лимит сверху", "Использование"]
# Проценты зон из согласованной таблицы казначейства (зелёная, жёлтая, красная).
EXPECTED_PERCENTS = {
    "TTS": (0.7808, 0.8509, 0.9510),
    "AFS": (0.9032, 0.9355, 0.9677),
    "HTM": (0.8929, 0.9286, 0.9643),
    "HTM_KUAP": (0.7800, 0.8500, 0.9500),
}


def write_limits_file(path: Path, rows, header_row: int = 6) -> Path:
    """Выгрузка лимитов: шапка на 6-й строке, под ней ПУСТАЯ строка, потом данные.

    rows — список (тип лимита, лимит сверху в рублях).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Результат"
    ws["A1"] = "Состояние лимитов"
    ws["A3"] = "Служебная строка выгрузки"
    for j, title in enumerate(LIMITS_HEADER, start=1):
        ws.cell(row=header_row, column=j, value=title)
    # header_row + 1 намеренно остаётся пустой
    for i, (limit_type, amount) in enumerate(rows, start=header_row + 2):
        ws.cell(row=i, column=1, value="Казначейство")
        ws.cell(row=i, column=2, value=limit_type)
        ws.cell(row=i, column=3, value="RUB")
        ws.cell(row=i, column=5, value=amount)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


DEFAULT_ROWS = [
    ("AFS", 40_000 * RUB),
    ("AFS", 12_000 * RUB),          # второй AFS, меньше — должен быть отброшен
    ("HTM", 220_000 * RUB),
    ("HTM_KUAP", 70_000 * RUB),
    ("Облигации", 45_000 * RUB),    # торговый портфель
    ("Прочий лимит", 9_000 * RUB),  # не портфельный тип
]


class LimitsTestCase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self._saved_env = os.environ.get(settings.SETTINGS_FILE_ENV)
        settings_file = self.tmp / "settings.json"
        settings_file.write_text(json.dumps({
            "portfolio_dynamics_dir": str(self.tmp / "data"),
            "portfolio_dynamics_output_dir": str(self.tmp / "out"),
            "downloads_dir": str(self.tmp / "downloads"),
        }), encoding="utf-8")
        os.environ[settings.SETTINGS_FILE_ENV] = str(settings_file)
        config.reload()

    def tearDown(self):
        if self._saved_env is None:
            os.environ.pop(settings.SETTINGS_FILE_ENV, None)
        else:
            os.environ[settings.SETTINGS_FILE_ENV] = self._saved_env
        config.reload()
        self._tmp.cleanup()

    def parse(self, rows=None, **kwargs) -> pd.DataFrame:
        path = write_limits_file(self.tmp / "Состояние лимитов на дату 21_09_2026 - Результат.xlsx",
                                 DEFAULT_ROWS if rows is None else rows, **kwargs)
        return limits.parse_limits_file(path)


class ParseTests(LimitsTestCase):
    def test_header_is_found_despite_the_blank_row_under_it(self):
        parsed = self.parse()
        self.assertEqual(set(parsed["portfolio_type"]), {"AFS", "HTM", "HTM_KUAP", "TTS", "ПРОЧИЙ ЛИМИТ"})

    def test_header_row_is_found_by_content_not_by_number(self):
        """Форму выгрузки подвинут на строку — разбор не должен развалиться."""
        parsed = self.parse(header_row=9)
        self.assertEqual(len(parsed), 5)

    def test_duplicate_afs_takes_the_larger_limit(self):
        parsed = self.parse().set_index("portfolio_type")
        self.assertAlmostEqual(parsed.loc["AFS", "limit_amount"], 40_000)

    def test_trading_portfolio_is_renamed_to_its_type(self):
        """В файле лимитов торговый портфель называется «Облигации», в отчёте — TTS."""
        parsed = self.parse().set_index("portfolio_type")
        self.assertIn("TTS", parsed.index)
        self.assertAlmostEqual(parsed.loc["TTS", "limit_amount"], 45_000)
        self.assertNotIn("ОБЛИГАЦИИ", parsed.index)

    def test_roubles_are_converted_to_millions(self):
        parsed = self.parse().set_index("portfolio_type")
        self.assertAlmostEqual(parsed.loc["HTM", "limit_amount"], 220_000)

    def test_rows_without_an_amount_are_skipped(self):
        parsed = self.parse([("AFS", 40_000 * RUB), ("HTM", None), ("HTM", 220_000 * RUB)])
        self.assertEqual(set(parsed["portfolio_type"]), {"AFS", "HTM"})

    def test_numbers_as_text_are_understood(self):
        parsed = self.parse([("AFS", "40 000 000 000"), ("HTM", "220 000 000 000")])
        self.assertAlmostEqual(parsed.set_index("portfolio_type").loc["AFS", "limit_amount"], 40_000)

    def test_missing_columns_are_reported_clearly(self):
        wb = Workbook()
        wb.active["A6"] = "Совсем другая колонка"
        path = self.tmp / "чужой.xlsx"
        wb.save(path)
        with self.assertRaisesRegex(etl.PortfolioDynamicsError, "Тип лимита"):
            limits.parse_limits_file(path)

    def test_aliases_are_configurable(self):
        settings.set_value("portfolio_dynamics_limit_aliases", "Облигации=HFT, Ценные бумаги=AFS")
        config.reload()
        parsed = self.parse().set_index("portfolio_type")
        self.assertIn("HFT", parsed.index)
        self.assertNotIn("TTS", parsed.index)


class ZoneTests(LimitsTestCase):
    def test_zone_percents_match_the_agreed_table(self):
        for portfolio_type, (green, yellow, red) in EXPECTED_PERCENTS.items():
            with self.subTest(portfolio_type=portfolio_type):
                limit = 100_000
                got_green, got_yellow, got_red = limits.zones_for(portfolio_type, limit)
                self.assertAlmostEqual(got_green / limit, green, places=6)
                self.assertAlmostEqual(got_yellow / limit, yellow, places=6)
                self.assertAlmostEqual(got_red / limit, red, places=6)

    def test_zones_are_monotonic_and_below_the_limit(self):
        """CHK_20 требует green < yellow < red, CHK_21 — red <= лимита."""
        for portfolio_type in EXPECTED_PERCENTS:
            with self.subTest(portfolio_type=portfolio_type):
                green, yellow, red = limits.zones_for(portfolio_type, 100_000)
                self.assertLess(green, yellow)
                self.assertLess(yellow, red)
                self.assertLessEqual(red, 100_000)

    def test_unknown_type_falls_back_and_warns(self):
        with self.assertLogs("portfolio_dynamics", level="WARNING"):
            green, yellow, red = limits.zones_for("НОВЫЙ_ТИП", 100_000)
        self.assertEqual((green, yellow, red), (70_000, 90_000, 100_000))


class BuildFactLimitTests(LimitsTestCase):
    def build(self, known_types, previous=None, rows=None):
        return limits.build_fact_limit(
            self.parse(rows), known_types,
            previous if previous is not None else pd.DataFrame(columns=etl.LIMIT_COLUMNS),
            dt.date(2026, 9, 21), "Состояние лимитов на дату 21_09_2026",
        )

    def test_only_known_portfolio_types_get_into_fact_limit(self):
        frame = self.build(["AFS", "HTM", "HTM_KUAP", "TTS"])
        self.assertEqual(sorted(frame["portfolio_type"]), ["AFS", "HTM", "HTM_KUAP", "TTS"])

    def test_unrelated_limit_rows_are_ignored(self):
        frame = self.build(["AFS"])
        self.assertNotIn("ПРОЧИЙ ЛИМИТ", list(frame["portfolio_type"]))

    def test_type_absent_from_the_file_keeps_its_previous_limit(self):
        """Неполный файл не должен обнулять уже согласованные лимиты."""
        previous = pd.DataFrame([
            ["ОСОБЫЙ", 5_000, 3_500, 4_500, 5_000, dt.date(2026, 1, 1), "Иванов И.И."],
        ], columns=etl.LIMIT_COLUMNS)

        with self.assertLogs("portfolio_dynamics", level="WARNING") as captured:
            frame = self.build(["AFS", "ОСОБЫЙ"], previous=previous)

        row = frame.set_index("portfolio_type").loc["ОСОБЫЙ"]
        self.assertEqual(row["limit_amount"], 5_000)
        self.assertEqual(row["updated_by"], "Иванов И.И.")
        self.assertTrue(any("ОСОБЫЙ" in line for line in captured.output))

    def test_changed_limit_is_logged(self):
        previous = pd.DataFrame([
            ["AFS", 10_000, 7_000, 9_000, 10_000, dt.date(2026, 1, 1), "Иванов И.И."],
        ], columns=etl.LIMIT_COLUMNS)

        with self.assertLogs("portfolio_dynamics", level="WARNING") as captured:
            self.build(["AFS"], previous=previous)

        self.assertTrue(any("изменён" in line and "AFS" in line for line in captured.output))

    def test_audit_trail_points_at_the_source_file(self):
        frame = self.build(["AFS"]).set_index("portfolio_type")
        self.assertIn("Состояние лимитов", str(frame.loc["AFS", "updated_by"]))
        self.assertEqual(frame.loc["AFS", "valid_from"], dt.date(2026, 9, 21))


class TypeGuessTests(unittest.TestCase):
    def test_longest_known_type_wins(self):
        """HTM_KUAP_CORE — это КУАП, а не HTM: резка по первому «_» уводила бы
        объём КУАП в чужой тип."""
        known = ["AFS", "HTM", "HTM_KUAP", "TTS"]
        self.assertEqual(etl.guess_type("HTM_KUAP_CORE", known), "HTM_KUAP")
        self.assertEqual(etl.guess_type("HTM_ALCO", known), "HTM")
        self.assertEqual(etl.guess_type("HTM_KUAP", known), "HTM_KUAP")

    def test_falls_back_to_the_prefix_without_known_types(self):
        self.assertEqual(etl.guess_type("HTM_KUAP_CORE"), "HTM")
        self.assertEqual(etl.guess_type("AFS_TR_RUR"), "AFS")

    def test_unknown_code_keeps_the_prefix_rule(self):
        self.assertEqual(etl.guess_type("XXX_NEW", ["AFS", "HTM"]), "XXX")


class NestedFixture(LimitsTestCase):
    """Общая песочница для вложенности: HTM собственный 150 000 + КУАП 50 000.

    Тестов не содержит — только фикстуру: два режима вложенности проверяются
    наследниками, и наследовать друг у друга тесты нельзя, они противоречат.
    """

    def setUp(self):
        super().setUp()
        from test_portfolio_dynamics_etl import export_name, write_export, MLN
        self.downloads = self.tmp / "downloads"
        self.downloads.mkdir(exist_ok=True)
        # HTM собственный 150 000 + КУАП 50 000; AFS 55 000; TTS 32 000
        self.portfolios = [("AFS_A", 30_000), ("AFS_B", 25_000), ("HTM_GOV", 90_000),
                           ("HTM_ALCO", 60_000), ("HTM_KUAP_CORE", 50_000),
                           ("TTS_OFZ", 20_000), ("TTS_CORP", 12_000)]

        def rows(scale):
            out = []
            for code, volume in self.portfolios:
                out += [(f"Позиция: {code}", None, None, None),
                        ("Bond", volume * scale * MLN, 3.0, 2.5)]
            return out

        self.t0 = write_export(self.downloads / export_name("21.09.2026"), rows(1.0),
                               period_end="21.09.2026")
        self.t7 = write_export(self.downloads / export_name("14.09.2026"), rows(0.97),
                               period_end="14.09.2026")
        self.limits_file = write_limits_file(
            self.downloads / "Состояние лимитов на дату 21_09_2026 - Результат.xlsx",
            [("AFS", 60_000 * RUB), ("HTM", 220_000 * RUB),
             ("HTM_KUAP", 70_000 * RUB), ("Облигации", 45_000 * RUB)])

    def build(self):
        return etl.build_data(self.t0, self.t7, bootstrap=True, limits_path=self.limits_file)

    def volumes(self, data):
        today = data.fact_type_daily[data.fact_type_daily["business_date"] == data.business_date]
        return dict(zip(today["portfolio_type"], today["volume_amount"]))

    def limits_of(self, data):
        return dict(zip(data.fact_limit["portfolio_type"], data.fact_limit["limit_amount"]))


class NestedTypesTests(NestedFixture):
    """Подлимит НЕ выделен: объём вложенного типа складывается с объемлющим и
    сравнивается с общим лимитом — так превышение общего лимита не потеряется."""

    def test_parent_volume_includes_the_nested_type(self):
        volumes = self.volumes(self.build())
        self.assertEqual(volumes["HTM"], 200_000)       # 90 000 + 60 000 + 50 000 КУАП
        self.assertEqual(volumes["HTM_KUAP"], 50_000)   # свой подлимитный объём
        self.assertEqual(volumes["AFS"], 55_000)        # без вложенных — как было

    def test_nested_portfolios_are_excluded_from_the_bank_total(self):
        """Иначе КУАП сложится дважды: сам по себе и внутри HTM."""
        with self.assertLogs("portfolio_dynamics", level="WARNING") as captured:
            data = self.build()

        row = data.dim_portfolio.set_index("portfolio_code").loc["HTM_KUAP_CORE"]
        self.assertFalse(bool(row["include_in_total"]))
        self.assertTrue(all(bool(v) for v in
                            data.dim_portfolio[data.dim_portfolio.portfolio_type != "HTM_KUAP"]
                            ["include_in_total"]))
        self.assertTrue(any("Входит в итог" in line for line in captured.output))

    def test_grain_check_stays_clean_with_nesting(self):
        """CHK_16 сравнивает свод по портфелям с объёмом типа — при вложенности
        обе стороны должны складываться одинаково."""
        from reports.portfolio_dynamics import workbook
        data = self.build()
        failed = {cid for cid, status, _v in workbook.evaluate_checks(data) if status == "FAIL"}
        self.assertNotIn("CHK_16", failed)

    def test_sub_limit_of_the_nested_type_still_applies(self):
        limits_frame = self.build().fact_limit.set_index("portfolio_type")
        self.assertEqual(limits_frame.loc["HTM_KUAP", "limit_amount"], 70_000)
        self.assertEqual(limits_frame.loc["HTM", "limit_amount"], 220_000)

    def test_nesting_can_be_switched_off(self):
        settings.set_value("portfolio_dynamics_type_parents", "нет=нет")
        config.reload()
        volumes = self.volumes(self.build())
        self.assertEqual(volumes["HTM"], 150_000)  # без КУАП


class AllocatedSubLimitTests(NestedFixture):
    """Часть совокупного лимита отводится вложенному типу настройкой.

    Лимит HTM в выгрузке — общий на HTM и КУАП. Сколько из него отдано КУАП,
    выгрузка не знает: это решение казначейства, оно живёт в настройках.
    """

    def setUp(self):
        super().setUp()
        # В файле только совокупный лимит HTM = 900, строки HTM_KUAP нет.
        self.limits_file = write_limits_file(
            self.downloads / "Состояние лимитов на дату 21_09_2026 - Результат.xlsx",
            [("AFS", 400_000 * RUB), ("HTM", 900_000 * RUB), ("Облигации", 250_000 * RUB)])
        settings.set_value("portfolio_dynamics_nested_limits", "HTM_KUAP=100000")
        config.reload()

    def test_aggregate_limit_is_split_between_parent_and_child(self):
        """900 всего, 100 отдано КУАП -> на весь остальной HTM остаётся 800."""
        got = self.limits_of(self.build())
        self.assertEqual(got["HTM_KUAP"], 100_000)
        self.assertEqual(got["HTM"], 800_000)
        self.assertEqual(got["AFS"], 400_000)  # без вложенных — как пришло

    def test_volumes_stay_separate_when_a_sub_limit_is_allocated(self):
        """Лимиты разделены — значит и объёмы сравниваются раздельно."""
        volumes = self.volumes(self.build())
        self.assertEqual(volumes["HTM"], 150_000)       # только свои, без КУАП
        self.assertEqual(volumes["HTM_KUAP"], 50_000)

    def test_nested_portfolios_stay_in_the_bank_total(self):
        """Объёмы не пересекаются, поэтому задвоения нет и флаг снимать не надо."""
        data = self.build()
        row = data.dim_portfolio.set_index("portfolio_code").loc["HTM_KUAP_CORE"]
        self.assertTrue(bool(row["include_in_total"]))

    def test_type_is_known_even_when_absent_from_the_limits_file(self):
        """Строки HTM_KUAP в файле нет — тип объявлен настройкой, и портфель
        КУАП не должен уехать в HTM."""
        data = self.build()
        row = data.dim_portfolio.set_index("portfolio_code").loc["HTM_KUAP_CORE"]
        self.assertEqual(row["portfolio_type"], "HTM_KUAP")

    def test_allocation_larger_than_the_aggregate_is_refused(self):
        settings.set_value("portfolio_dynamics_nested_limits", "HTM_KUAP=950000")
        config.reload()
        with self.assertRaisesRegex(etl.PortfolioDynamicsError, "больше совокупного"):
            self.build()

    def test_allocation_wins_over_the_row_in_the_file(self):
        self.limits_file = write_limits_file(
            self.downloads / "Состояние лимитов на дату 21_09_2026 - Результат.xlsx",
            [("HTM", 900_000 * RUB), ("HTM_KUAP", 70_000 * RUB)])
        got = self.limits_of(self.build())
        self.assertEqual(got["HTM_KUAP"], 100_000)  # из настроек, не 70 000
        self.assertEqual(got["HTM"], 800_000)

    def test_zones_follow_the_split_limit(self):
        limits_frame = self.build().fact_limit.set_index("portfolio_type")
        htm = limits_frame.loc["HTM"]
        self.assertAlmostEqual(htm["red_max_util"] / htm["limit_amount"], 0.9643, places=6)
        kuap = limits_frame.loc["HTM_KUAP"]
        self.assertAlmostEqual(kuap["green_max_util"] / kuap["limit_amount"], 0.78, places=6)

    def test_allocation_for_a_type_that_is_not_nested_is_reported(self):
        settings.set_value("portfolio_dynamics_nested_limits", "AFS=10000")
        config.reload()
        with self.assertLogs("portfolio_dynamics", level="WARNING") as captured:
            got = self.limits_of(self.build())
        self.assertEqual(got["AFS"], 10_000)
        self.assertTrue(any("не объявлен вложенным" in line for line in captured.output))

    def test_bad_allocation_value_is_reported(self):
        settings.set_value("portfolio_dynamics_nested_limits", "HTM_KUAP=сто")
        config.reload()
        with self.assertRaisesRegex(etl.PortfolioDynamicsError, "тип=сумма"):
            etl.parse_nested_limits()


class NestingHelpersTests(unittest.TestCase):
    def test_descendants_and_counted_types(self):
        parents = etl.parse_type_parents("HTM_KUAP=HTM, HTM_SUB=HTM_KUAP")
        self.assertEqual(sorted(etl.descendants_of("HTM", parents)), ["HTM_KUAP", "HTM_SUB"])
        self.assertEqual(sorted(etl.types_counted_in("HTM", parents)),
                         ["HTM", "HTM_KUAP", "HTM_SUB"])
        self.assertEqual(etl.types_counted_in("AFS", parents), ["AFS"])

    def test_self_reference_is_ignored(self):
        self.assertEqual(etl.parse_type_parents("HTM=HTM"), {})

    def test_empty_map_means_no_nesting(self):
        self.assertEqual(etl.parse_type_parents(""), {})


class LimitsFileIsNotASliceTests(LimitsTestCase):
    def test_limits_file_is_recognised(self):
        self.assertTrue(etl.is_limits_file(
            Path("Состояние лимитов на дату 21_09_2026 - Результат.xlsx")))
        self.assertFalse(etl.is_limits_file(
            Path("Позиция за период   21.09.2026  -  21.09.2026   - SECURITIES.xlsx")))

    def test_limits_file_does_not_break_slice_selection(self):
        """В папке-дате лежат ТРИ файла; файл лимитов не должен считаться срезом."""
        from test_date_folders import slice_rows
        from test_portfolio_dynamics_etl import export_name, write_export

        folder = self.tmp / "data" / "2026-09-21"
        write_export(folder / export_name("21.09.2026"), slice_rows(30, 60), period_end="21.09.2026")
        write_export(folder / export_name("14.09.2026"), slice_rows(28, 55), period_end="14.09.2026")
        write_limits_file(folder / "Состояние лимитов на дату 21_09_2026 - Результат.xlsx",
                          DEFAULT_ROWS)

        t0, t7 = etl.split_slice_files(sorted(folder.glob("*.xlsx")), folder.name)

        self.assertIn("21.09.2026", t0.name)
        self.assertIn("14.09.2026", t7.name)


if __name__ == "__main__":
    unittest.main()
