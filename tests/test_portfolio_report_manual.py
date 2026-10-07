"""Тесты «Дополнительных портфелей» в «Отчёте по портфелям».

Проверяется: портфель из Excel-файла попадает в CSV с Open QTY (на дату и на
начало года — одно и то же, изменение ноль) и чистой стоимостью в рублях;
P&L, не введённый при запуске, берётся как вчера; DV01 и Yield у такого
портфеля не пишутся; код, который есть в выгрузке, файлом не перебивается.
"""
import argparse
import datetime as dt
import sys
import unittest
from pathlib import Path

import pandas as pd

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "tests"))

import config  # noqa: E402
from common import settings  # noqa: E402
from reports.portfolio_report import etl, manual, workbook  # noqa: E402
from reports.portfolio_report.report import PortfolioReport  # noqa: E402
# Модулем, а не классом: импортированный класс unittest прогнал бы ещё раз.
import test_portfolio_report as base  # noqa: E402

EXTRA = {"code": "OFZ_EXTRA", "name": "Внебиржевой ОФЗ", "type": "HTM",
         "volume": 150, "qty": 1_000, "duration": 4.2}


class ManualPortfolioReportTest(unittest.TestCase):
    setUp = base.PortfolioReportTest.setUp
    tearDown = base.PortfolioReportTest.tearDown

    def set_extra(self, *records):
        settings.set_value("portfolio_dynamics_manual_portfolios", list(records))
        config.reload()

    def run_report(self, path: Path, manual_pl=None) -> pd.DataFrame:
        args = argparse.Namespace(date=None, input=str(path), no_import=True, output=None,
                                  rgbi=None, ruonia=None, rwa=None, manual_pl=manual_pl)
        PortfolioReport().run(args)
        day = etl.positions.read_business_date(path)
        return pd.read_csv(etl.default_output_path(day), dtype=str, keep_default_na=False,
                           encoding="utf-8-sig")

    @staticmethod
    def values(flat: pd.DataFrame, code: str) -> dict:
        rows = flat[flat["axis_2"] == code]
        return {r.axis_3: (r.axis_1, float(r.value) if r.value else r.text_value)
                for r in rows.itertuples()}

    def test_manual_portfolio_rows(self):
        self.set_extra(EXTRA)
        flat = self.run_report(self.day1, manual_pl="OFZ_EXTRA=12,5")
        self.assertEqual(self.values(flat, "OFZ_EXTRA"), {
            "Open QTY": ("HTM", 1_000),
            "Open QTY на начало года": ("HTM", 1_000),
            "Изменение Open QTY": ("HTM", 0),
            "Чистая стоимость": ("HTM", 150_000_000),
            "Total Full PL with Funding": ("HTM", 12_500_000),
        })  # DV01 и Yield нет — в средневзвешенный Yield портфель не входит

    def test_pl_not_entered_is_taken_from_yesterday(self):
        self.set_extra(EXTRA)
        self.run_report(self.day1, manual_pl="OFZ_EXTRA=12.5")
        flat = self.run_report(self.day2)
        self.assertEqual(self.values(flat, "OFZ_EXTRA")["Total Full PL with Funding"][1],
                         12_500_000)
        history = manual.read_history()
        self.assertEqual(history[dt.date(2026, 9, 29)], {"OFZ_EXTRA": 12.5},
                         "значение «как вчера» записывается и на сегодняшнюю дату")

        # Ввели новое — оно важнее, повторный прогон без ввода его не теряет.
        self.run_report(self.day2, manual_pl="OFZ_EXTRA=13")
        flat = self.run_report(self.day2)
        self.assertEqual(self.values(flat, "OFZ_EXTRA")["Total Full PL with Funding"][1],
                         13_000_000)
        self.assertEqual(manual.read_history()[dt.date(2026, 9, 28)], {"OFZ_EXTRA": 12.5})

    def test_without_any_pl_there_is_no_pl_row(self):
        self.set_extra(dict(EXTRA, qty=None))
        values = self.values(self.run_report(self.day1), "OFZ_EXTRA")
        self.assertEqual(set(values), {"Чистая стоимость"})

    def test_comment_comes_from_the_file(self):
        """Комментарий доп. портфеля — из файла: правка в витрине перезаписывается,
        стёртый в файле — пропадает."""
        self.set_extra(dict(EXTRA, comment_report="Ведём вручную",
                            comment_dynamics="Это для динамики"))
        flat = self.run_report(self.day1, manual_pl="OFZ_EXTRA=1")
        self.assertEqual(self.values(flat, "OFZ_EXTRA")["Комментарий"], ("HTM", "Ведём вручную"))

        book = workbook.workbook_path(etl.default_output_path(dt.date(2026, 9, 28)))
        from openpyxl import load_workbook
        wb = load_workbook(book)
        ws = wb[workbook.SHEET]
        header = [c.value for c in ws[workbook.HEADER_ROW]]
        for row in ws.iter_rows(min_row=workbook.HEADER_ROW + 1):
            if row[0].value == "OFZ_EXTRA":
                row[header.index(workbook.COMMENT_HEADER)].value = "Правка в витрине"
        wb.save(book)

        flat = self.run_report(self.day2)
        self.assertEqual(self.values(flat, "OFZ_EXTRA")["Комментарий"], ("HTM", "Ведём вручную"))

        self.set_extra(EXTRA)
        flat = self.run_report(self.day2)
        self.assertNotIn("Комментарий", self.values(flat, "OFZ_EXTRA"))

    def test_code_from_the_export_is_not_overridden(self):
        self.set_extra(dict(EXTRA, code="AFS_TR_RUR", type="AFS"))
        with self.assertLogs("portfolio_report", level="WARNING") as captured:
            flat = self.run_report(self.day1)
        self.assertEqual(self.values(flat, "AFS_TR_RUR")["Чистая стоимость"][1], 400_000_000)
        self.assertTrue(any("AFS_TR_RUR" in line for line in captured.output))

    def test_pl_for_unknown_code_is_skipped(self):
        self.set_extra(EXTRA)
        flat = self.run_report(self.day1, manual_pl="OFZ_EXTRA=1; NOPE=2")
        self.assertNotIn("NOPE", set(flat["axis_2"]))

    def test_parse_entered(self):
        self.assertEqual(manual.parse_entered("A=12,5, B=3; c=1 000"),
                         {"A": 12.5, "B": 3.0, "C": 1000.0})
        with self.assertRaises(ValueError):
            manual.parse_entered("A")


if __name__ == "__main__":
    unittest.main()
