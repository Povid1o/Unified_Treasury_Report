"""Тесты формул выходного .xlsx «Динамики портфелей».

Формулы не проверяются обычными тестами: openpyxl их не вычисляет, а Excel
пересчитает только при открытии — то есть у человека, а не на прогоне. Поэтому
здесь два уровня.

1. Структурный (работает всегда): все ссылки формул ведут на существующие листы
   и именованные диапазоны, а сами функции — только уровня Excel 2007. Это
   ловит опечатку в имени листа и случайно заехавший XLOOKUP, из-за которого
   файл откроется с #ИМЯ? у половины витрины.
2. Вычислительный (если в окружении есть пакет formulas): книга считается
   целиком, и статусы листа checks сверяются с тем, что посчитал питон в
   workbook.evaluate_checks. Расхождение означает, что консоль пишет в лог одно,
   а файл показывает другое.
"""
import datetime as dt
import re
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "tests"))

from reports.portfolio_dynamics import workbook  # noqa: E402
from reports.portfolio_dynamics.etl import (  # noqa: E402
    DIM_COLUMNS, LIMIT_COLUMNS, SNAPSHOT_COLUMNS, TYPE_DAILY_COLUMNS,
    PortfolioDynamicsData,
)

try:
    import formulas  # noqa: F401
    HAS_FORMULAS = True
except ImportError:
    HAS_FORMULAS = False

# Функции, допустимые по контракту (CONTRACT.md, п. 6 и требования к файлу):
# только то, что понимает Excel 2007. Динамические массивы и новые функции
# запрещены — файл открывают в том числе в старых версиях.
ALLOWED_FUNCTIONS = {
    "IF", "IFERROR", "OR", "AND", "NOT",
    "SUM", "SUMIFS", "SUMPRODUCT",
    "COUNT", "COUNTA", "COUNTIF", "COUNTIFS",
    "INDEX", "MATCH", "MAX", "MIN",
}
BANNED_FUNCTIONS = {
    "XLOOKUP", "XMATCH", "FILTER", "UNIQUE", "SORT", "SORTBY", "SEQUENCE",
    "LET", "LAMBDA", "TEXTJOIN", "IFS", "MAXIFS", "MINIFS", "SWITCH",
}

_FUNCTION = re.compile(r"\b([A-Z][A-Z0-9\.]{1,})\s*\(")
_SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!")
# \b на конце обязателен: без него движок откатывается на «IFERRO» внутри
# «IFERROR(» и принимает хвост имени функции за именованный диапазон.
_NAMED_RANGE = re.compile(r"(?<![!'\w$])([A-Z][A-Z0-9_]{3,})\b(?!\s*\()")
_LITERALS = {"TRUE", "FALSE"}


def build_demo_data() -> PortfolioDynamicsData:
    """Небольшой, но полный набор: 2 типа, 4 портфеля, история за 3 даты."""
    business_date = dt.date(2026, 9, 18)
    dim = pd.DataFrame([
        ["AFS_OFZ", "ОФЗ, AFS", "AFS", True, True, 10],
        ["AFS_CORP", "Корпоративные, AFS", "AFS", True, True, 20],
        ["HTM_LONG", "Длинные, HTM", "HTM", True, True, 30],
        ["HTM_ALCO", "АЛКО, HTM", "HTM", True, True, 40],
    ], columns=DIM_COLUMNS)
    limits = pd.DataFrame([
        ["AFS", 100_000, 70_000, 90_000, 100_000, dt.date(2026, 1, 1), "Иванов И.И.", 45_000],
        # Второй тип — с пустым остатком: колонка необязательная, и книга
        # обязана собираться, когда система лимитов его не прислала.
        ["HTM", 300_000, 210_000, 270_000, 300_000, dt.date(2026, 1, 1), "Иванов И.И.", None],
    ], columns=LIMIT_COLUMNS)

    volumes = {"AFS_OFZ": 30_000, "AFS_CORP": 25_000, "HTM_LONG": 140_000, "HTM_ALCO": 60_000}
    by_type = {"AFS": 55_000, "HTM": 200_000}
    history = []
    for offset, factor in ((14, 0.96), (7, 0.98), (0, 1.0)):
        day = business_date - dt.timedelta(days=offset)
        for portfolio_type, total in by_type.items():
            history.append([day, portfolio_type, round(total * factor)])
    # Свежая дата должна совпасть со срезом до копейки, иначе CHK_16 поймает.
    for row in history:
        if row[0] == business_date:
            row[2] = by_type[row[1]]

    snapshot = pd.DataFrame([
        [business_date, code, volume, round(volume * 0.98), 3.0, 2.5, f"Заметка {code}"]
        for code, volume in volumes.items()
    ], columns=SNAPSHOT_COLUMNS)

    return PortfolioDynamicsData(
        business_date=business_date, lookback_days=7,
        dim_portfolio=dim, fact_limit=limits,
        fact_type_daily=pd.DataFrame(history, columns=TYPE_DAILY_COLUMNS),
        fact_portfolio_snapshot=snapshot,
    )


class WorkbookFormulaTestCase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.data = build_demo_data()
        cls.path = Path(cls._tmp.name) / "mock.xlsx"
        workbook.save_workbook(cls.data, cls.path)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def formulas_in_book(self):
        """(лист, координата, формула) по всем листам книги."""
        wb = load_workbook(self.path)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        yield sheet_name, cell.coordinate, cell.value


class StructureTests(WorkbookFormulaTestCase):
    def test_every_formula_references_an_existing_sheet(self):
        """Опечатка в имени листа даёт #ССЫЛКА! во всей витрине."""
        sheets = set(load_workbook(self.path).sheetnames)
        bad = []
        for sheet, coord, formula in self.formulas_in_book():
            for quoted, plain in _SHEET_REF.findall(formula):
                referenced = quoted or plain
                if referenced not in sheets:
                    bad.append(f"{sheet}!{coord}: лист {referenced!r}")
        self.assertEqual(bad, [], "ссылки на несуществующие листы")

    def test_every_named_range_used_is_defined(self):
        wb = load_workbook(self.path)
        defined = set(wb.defined_names)
        sheets = {s.upper() for s in wb.sheetnames}
        bad = []
        for sheet, coord, formula in self.formulas_in_book():
            # Имена листов и текстовые константы в формулах — не диапазоны.
            stripped = re.sub(r'"[^"]*"', "", formula)
            for token in _NAMED_RANGE.findall(stripped):
                if token in ALLOWED_FUNCTIONS or token in sheets or token in _LITERALS:
                    continue
                if token not in defined:
                    bad.append(f"{sheet}!{coord}: имя {token!r}")
        self.assertEqual(bad, [], "формулы ссылаются на незаведённые именованные диапазоны")

    def test_only_excel_2007_functions_are_used(self):
        """XLOOKUP/FILTER/UNIQUE в старом Excel дают #ИМЯ? — контракт их запрещает."""
        used = set()
        for _sheet, _coord, formula in self.formulas_in_book():
            used.update(_FUNCTION.findall(formula))
        self.assertEqual(used & BANNED_FUNCTIONS, set(), "запрещённые функции")
        self.assertEqual(used - ALLOWED_FUNCTIONS, set(),
                         "функции вне согласованного списка — проверьте совместимость с Excel 2007")

    def test_named_ranges_required_by_the_contract_exist(self):
        defined = set(load_workbook(self.path).defined_names)
        for name in ("BUSINESS_DATE", "LOOKBACK_DAYS", "PORTFOLIO_CODES", "PORTFOLIO_TYPES"):
            self.assertIn(name, defined)

    def test_every_check_row_has_an_id_and_a_status_formula(self):
        ws = load_workbook(self.path)["checks"]
        ids, statuses = [], 0
        for row in range(5, 5 + 24):
            ids.append(ws.cell(row=row, column=1).value)
            value = ws.cell(row=row, column=5).value
            if isinstance(value, str) and value.startswith("="):
                statuses += 1
        self.assertEqual(ids, [f"CHK_{i:02d}" for i in range(1, 25)])
        self.assertEqual(statuses, 24)
        self.assertIsNone(ws.cell(row=29, column=1).value, "проверок должно быть ровно 24")

    def test_no_worksheet_autofilter_next_to_a_table(self):
        """Два автофильтра на одном диапазоне — Excel открывает файл с
        предложением восстановить и выбрасывает таблицу целиком
        («Удалённое свойство: Таблица из части /xl/tables/tableN.xml»).
        Фильтр должен быть ровно один — тот, что заводит сама умная таблица."""
        wb = load_workbook(self.path)
        offenders = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if ws.tables and ws.auto_filter.ref:
                offenders.append(f"{sheet_name}: лист {ws.auto_filter.ref}, "
                                 f"таблиц {len(ws.tables)}")
        self.assertEqual(offenders, [], "на листе с таблицей стоит лишний автофильтр")

    def test_every_table_keeps_its_own_filter(self):
        """Убрав лишний фильтр, нельзя потерять и нужный: выпадающие списки в
        шапке даёт autoFilter самой таблицы."""
        wb = load_workbook(self.path)
        for sheet_name in ("dim_portfolio", "fact_limit", "fact_type_daily",
                           "fact_portfolio_snapshot", "dict"):
            table = list(wb[sheet_name].tables.values())[0]
            self.assertIsNotNone(table.autoFilter, sheet_name)
            self.assertEqual(table.autoFilter.ref, table.ref, sheet_name)

    def test_no_leftover_filter_database_names(self):
        """_xlnm._FilterDatabase появлялись вместе с лист-уровневым фильтром."""
        names = list(load_workbook(self.path).defined_names)
        self.assertEqual([n for n in names if "FilterDatabase" in n], [])

    def test_table_refs_match_the_data(self):
        """Ref таблицы шире или уже данных — Excel предлагает «восстановить» файл."""
        wb = load_workbook(self.path)
        expected = {
            "dim_portfolio": len(self.data.dim_portfolio),
            "fact_limit": len(self.data.fact_limit),
            "fact_type_daily": len(self.data.fact_type_daily),
            "fact_portfolio_snapshot": len(self.data.fact_portfolio_snapshot),
        }
        for sheet_name, rows in expected.items():
            ws = wb[sheet_name]
            self.assertEqual(len(ws.tables), 1, sheet_name)
            ref = list(ws.tables.values())[0].ref
            last_row = int(re.search(r"(\d+)$", ref).group(1))
            self.assertEqual(last_row, max(rows + 1, 2), f"{sheet_name}: ref={ref}")


@unittest.skipUnless(HAS_FORMULAS, "нужен пакет formulas: pip install formulas")
class EvaluationTests(WorkbookFormulaTestCase):
    """Вычисление книги целиком — запускается, если в окружении есть formulas."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        import logging
        logging.disable(logging.WARNING)
        import formulas as engine
        cls.solution = engine.ExcelModel().loads(str(cls.path)).finish().calculate()

    def value(self, sheet, ref):
        cell = self.solution.get(f"'[{self.path.name}]{sheet.upper()}'!{ref}")
        try:
            return cell.value[0, 0]
        except Exception:
            return None

    def test_no_cell_evaluates_to_an_excel_error(self):
        broken = []
        for key, cell in self.solution.items():
            if not isinstance(key, str) or "!" not in key:
                continue
            try:
                value = cell.value[0, 0]
            except Exception:
                continue
            if isinstance(value, str) and value.startswith("#"):
                broken.append(f"{key} -> {value}")
        self.assertEqual(broken, [], "ячейки с ошибкой Excel")

    def test_checks_sheet_agrees_with_the_python_evaluation(self):
        """Иначе консоль пишет в лог одно, а файл показывает другое."""
        python_side = {cid: status for cid, status, _v in workbook.evaluate_checks(self.data)}
        excel_side = {
            str(self.value("checks", f"A{row}")): str(self.value("checks", f"E{row}"))
            for row in range(5, 29)
        }
        self.assertEqual(excel_side, python_side)

    def test_all_checks_pass_on_valid_data(self):
        statuses = {str(self.value("checks", f"E{row}")) for row in range(5, 29)}
        self.assertEqual(statuses, {"OK"})

    def test_view_by_type_sums_and_traffic_light(self):
        rows = {self.value("view_by_type", f"A{r}"): r for r in (6, 7)}
        afs, htm = rows["AFS"], rows["HTM"]

        self.assertEqual(self.value("view_by_type", f"D{afs}"), 55_000)   # объём типа
        self.assertEqual(self.value("view_by_type", f"H{afs}"), 55_000)   # сумма по портфелям
        self.assertEqual(self.value("view_by_type", f"I{afs}"), 0)        # расхождение грейнов
        self.assertEqual(self.value("view_by_type", f"L{afs}"), "ЗЕЛЁНАЯ")
        self.assertEqual(self.value("view_by_type", f"M{afs}"), 45_000)   # свободный лимит

        # HTM: 200 000 между жёлтой (270 000)? нет — ниже, значит зелёная зона.
        self.assertEqual(self.value("view_by_type", f"D{htm}"), 200_000)
        self.assertEqual(self.value("view_by_type", f"L{htm}"), "ЗЕЛЁНАЯ")
        self.assertAlmostEqual(self.value("view_by_type", f"K{htm}"), 200_000 / 300_000, places=6)

    def test_view_monitor_deltas(self):
        self.assertEqual(self.value("view_monitor", "A6"), "AFS_OFZ")
        self.assertEqual(self.value("view_monitor", "D6"), 30_000)
        self.assertEqual(self.value("view_monitor", "E6"), 29_400)
        self.assertEqual(self.value("view_monitor", "F6"), 600)
        self.assertAlmostEqual(self.value("view_monitor", "G6"), 30_000 / 29_400 - 1, places=6)
        self.assertAlmostEqual(self.value("view_monitor", "J6"), 0.5, places=6)  # гэп дюрации
        self.assertEqual(self.value("view_monitor", "K6"), 100_000)              # лимит типа


if __name__ == "__main__":
    unittest.main()
