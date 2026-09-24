"""Тесты отчёта «Динамика портфелей».

Проверяется то, что ломается молча: слияние заметок, идемпотентность дозаписи
истории, неприкосновенность прошлых дат, сверка двух грейнов, разбор иерархии
и средневзвешенная (а не средняя) дюрация.
"""
import datetime as dt
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))

import config  # noqa: E402
from common import settings  # noqa: E402
from reports.portfolio_dynamics import etl, workbook  # noqa: E402

# Объёмы во входных файлах пишем в рублях, как в реальной выгрузке.
MLN = config.PORTFOLIO_DYNAMICS_VALUE_SCALE

# Колонки выгрузки: нужные четыре вперемешку с лишними, как в реальном файле.
EXPORT_HEADER = [
    "Тип актива ", "Дата погашения", "Имя актива ", "ISIN ",
    "Чистая стоимость позиции (нач.)", "Чистая стоимость позиции (кон.)",
    "Duration (нач.)", "Дюрация", "WAPP (кон.)",
]
COL_VALUE = EXPORT_HEADER.index("Чистая стоимость позиции (кон.)")
COL_DUR_START = EXPORT_HEADER.index("Duration (нач.)")
COL_DUR_END = EXPORT_HEADER.index("Дюрация")


def write_export(path: Path, rows, period_start="01.01.2026", period_end="01.09.2026") -> Path:
    """Синтетическая выгрузка позиций: шапка, заголовки, иерархия «Позиция: ...».

    rows — список (значение колонки «Тип актива», объём, дюрация нач., дюрация кон.);
    None в числовой позиции означает пустую ячейку.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Position"
    ws.cell(row=1, column=1,
            value=f"Позиция за период [{period_start}] - [{period_end}] - SECURITIES")
    ws.cell(row=2, column=1, value="Строка 2 (Игнорируем)")
    ws.cell(row=3, column=1, value="Строка 3 (Игнорируем)")
    ws.cell(row=4, column=1, value="Instrument")
    for j, title in enumerate(EXPORT_HEADER, start=1):
        ws.cell(row=5, column=j, value=title)

    for i, (asset_type, value, dur_start, dur_end) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=asset_type)
        if value is not None:
            ws.cell(row=i, column=COL_VALUE + 1, value=value)
        if dur_start is not None:
            ws.cell(row=i, column=COL_DUR_START + 1, value=dur_start)
        if dur_end is not None:
            ws.cell(row=i, column=COL_DUR_END + 1, value=dur_end)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_name(period_end: str) -> str:
    return f"Позиция за период [01.01.2026] - [{period_end}] - SECURITIES.xlsx"


# Два портфеля, по две бумаги в каждом. AFS: 30 + 10 млн, HTM: 60 + 20 млн.
T0_ROWS = [
    ("Позиция: AFS_TR_RUR", None, None, None),
    ("Bond", 30 * MLN, 3.0, 2.5),
    ("Bond", 10 * MLN, 1.0, 1.5),
    ("Позиция: HTM_ALCO", None, None, None),
    ("Bond", 60 * MLN, 6.0, 5.5),
    ("Bond", 20 * MLN, 2.0, 2.0),
]
T7_ROWS = [
    ("Позиция: AFS_TR_RUR", None, None, None),
    ("Bond", 28 * MLN, 3.0, 2.5),
    ("Bond", 10 * MLN, 1.0, 1.5),
    ("Позиция: HTM_ALCO", None, None, None),
    ("Bond", 55 * MLN, 6.0, 5.5),
    ("Bond", 20 * MLN, 2.0, 2.0),
]


class PortfolioDynamicsTestCase(unittest.TestCase):
    """Общая песочница: файлы T0/T-7 и выходная папка во временном каталоге."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        # Свои настройки, а не настройки разработчика: отчёт пишет файл
        # разметки типов рядом с settings.json, и тест не должен сорить в проекте.
        self._saved_env = os.environ.get(settings.SETTINGS_FILE_ENV)
        settings_file = self.tmp / "settings.json"
        settings_file.write_text(json.dumps({}), encoding="utf-8")
        os.environ[settings.SETTINGS_FILE_ENV] = str(settings_file)
        config.reload()
        self.t0_path = write_export(self.tmp / export_name("01.09.2026"), T0_ROWS)
        self.t7_path = write_export(
            self.tmp / export_name("25.08.2026"), T7_ROWS, period_end="25.08.2026"
        )
        self.out_dir = self.tmp / "out"
        self.out_dir.mkdir()

    def tearDown(self):
        if self._saved_env is None:
            os.environ.pop(settings.SETTINGS_FILE_ENV, None)
        else:
            os.environ[settings.SETTINGS_FILE_ENV] = self._saved_env
        config.reload()
        self._tmp.cleanup()

    def build(self, previous=None, bootstrap=False, t0_path=None):
        return etl.build_data(
            t0_path or self.t0_path, self.t7_path,
            previous_path=previous, bootstrap=bootstrap,
        )

    def save(self, data, name=None) -> Path:
        path = self.out_dir / (name or f"{etl.OUTPUT_FILENAME_PREFIX}{data.business_date}.xlsx")
        return workbook.save_workbook(data, path)

    def bootstrap_release(self) -> Path:
        """Первый выпуск + заполненные руками лимиты (их скрипт не заводит)."""
        data = self.build(bootstrap=True)
        for column, value in (("limit_amount", 1000), ("green_max_util", 700),
                              ("yellow_max_util", 900), ("red_max_util", 1000)):
            data.fact_limit[column] = value
        return self.save(data)


# ════════════════════════════════════════════════════════════════════════════
# Разбор входных файлов
# ════════════════════════════════════════════════════════════════════════════
class ParseSliceTests(PortfolioDynamicsTestCase):
    def test_hierarchy_splits_securities_between_portfolios(self):
        """Бумаги уходят в свой портфель, шапка до первого маркера и «Итого» — мимо."""
        path = write_export(self.tmp / export_name("02.09.2026"), [
            ("Мусор из шапки выгрузки", 999 * MLN, None, None),
            ("Позиция: AFS_TR_RUR", None, None, None),
            ("Bond", 30 * MLN, 3.0, 2.5),
            ("Bond", 10 * MLN, 1.0, 1.5),
            ("Итого", 40 * MLN, None, None),
            ("Позиция: HTM_ALCO", None, None, None),
            ("Bond", 60 * MLN, 6.0, 5.5),
            ("Всего по портфелю", 60 * MLN, None, None),
        ], period_end="02.09.2026")

        result = etl.parse_slice(path, "T0")
        volumes = dict(zip(result.frame["portfolio_code"], result.frame["volume"]))

        self.assertEqual(list(result.frame["portfolio_code"]), ["AFS_TR_RUR", "HTM_ALCO"])
        self.assertAlmostEqual(volumes["AFS_TR_RUR"], 40.0)
        self.assertAlmostEqual(volumes["HTM_ALCO"], 60.0)
        self.assertEqual(result.stats.securities, 3)
        self.assertEqual(result.stats.rows_without_portfolio, 1)  # строка до первого маркера
        self.assertEqual(result.stats.total_rows_skipped, 2)      # «Итого» и «Всего»

    def test_duration_is_weighted_by_value_not_averaged(self):
        """Две бумаги 30 и 10 млн с дюрациями 3.0 и 1.0 -> 2.5, а не 2.0."""
        result = etl.parse_slice(self.t0_path, "T0")
        row = result.frame.set_index("portfolio_code").loc["AFS_TR_RUR"]

        self.assertAlmostEqual(row["duration_current_yrs"], (3.0 * 30 + 1.0 * 10) / 40)
        self.assertNotAlmostEqual(row["duration_current_yrs"], (3.0 + 1.0) / 2)

    def test_end_duration_is_weighted_when_reading_is_enabled(self):
        settings.set_value("portfolio_dynamics_read_duration_end", True)
        config.reload()
        row = etl.parse_slice(self.t0_path, "T0").frame.set_index("portfolio_code").loc["AFS_TR_RUR"]
        self.assertAlmostEqual(row["duration_target_yrs"], (2.5 * 30 + 1.5 * 10) / 40)

    def test_end_duration_is_not_read_by_default(self):
        """Колонка «Дюрация» есть в выгрузке, но по умолчанию не читается."""
        self.assertFalse(config.PORTFOLIO_DYNAMICS_READ_DURATION_END)
        frame = etl.parse_slice(self.t0_path, "T0").frame
        self.assertTrue(frame["duration_target_yrs"].isna().all())
        self.assertTrue(frame["duration_current_yrs"].notna().all())

    def test_security_without_duration_keeps_its_volume_but_not_its_weight(self):
        settings.set_value("portfolio_dynamics_read_duration_end", True)
        config.reload()
        path = write_export(self.tmp / export_name("03.09.2026"), [
            ("Позиция: AFS_TR_RUR", None, None, None),
            ("Bond", 30 * MLN, 3.0, None),
            ("Bond", 10 * MLN, None, None),
        ], period_end="03.09.2026")

        row = etl.parse_slice(path, "T0").frame.iloc[0]
        self.assertAlmostEqual(row["volume"], 40.0)
        self.assertAlmostEqual(row["duration_current_yrs"], 3.0)
        self.assertIsNone(row["duration_target_yrs"])


class KuapDurationTests(PortfolioDynamicsTestCase):
    """«Дюрация цель» — дюрация, установленная КУАП, из настройки."""

    def set_kuap(self, raw):
        settings.set_value("portfolio_dynamics_kuap_durations", raw)
        config.reload()

    def target(self, data):
        frame = data.fact_portfolio_snapshot.set_index("portfolio_code")
        return frame["duration_target_yrs"]

    def test_kuap_duration_is_written_to_listed_portfolios_only(self):
        self.set_kuap("afs_tr_rur=3,5")
        target = self.target(self.build(bootstrap=True))
        self.assertEqual(target["AFS_TR_RUR"], 3.5)
        others = target.drop("AFS_TR_RUR")
        self.assertTrue(others.isna().all(), "без КУАП и без чтения выгрузки цель пустая")

    def test_kuap_duration_overrides_the_export(self):
        settings.set_value("portfolio_dynamics_read_duration_end", True)
        self.set_kuap("AFS_TR_RUR=7")
        target = self.target(self.build(bootstrap=True))
        self.assertEqual(target["AFS_TR_RUR"], 7.0)
        self.assertTrue(target.drop("AFS_TR_RUR").notna().all(), "остальные — из выгрузки")

    def test_unknown_code_is_reported_not_fatal(self):
        self.set_kuap("NO_SUCH_PORTFOLIO=2")
        with self.assertLogs("portfolio_dynamics", level="WARNING") as captured:
            self.build(bootstrap=True)
        self.assertTrue(any("NO_SUCH_PORTFOLIO" in line for line in captured.output))

    def test_setting_is_validated(self):
        for bad in ("AFS_TR_RUR", "AFS_TR_RUR=много", "AFS_TR_RUR=-1", "A=1, A=2"):
            with self.assertRaises(settings.SettingsError, msg=bad):
                settings.set_value("portfolio_dynamics_kuap_durations", bad)

    def test_several_portfolios_with_decimal_commas(self):
        self.assertEqual(
            settings.parse_durations("A=3,5, B=1.25,C=2"),
            {"A": 3.5, "B": 1.25, "C": 2.0},
        )

    def test_subtotal_in_marker_row_wins_and_mismatch_is_reported(self):
        path = write_export(self.tmp / export_name("04.09.2026"), [
            ("Позиция: AFS_TR_RUR", 100 * MLN, None, None),
            ("Bond", 30 * MLN, 3.0, 2.5),
        ], period_end="04.09.2026")

        result = etl.parse_slice(path, "T0")
        self.assertAlmostEqual(result.frame.iloc[0]["volume"], 100.0)
        self.assertEqual(len(result.stats.subtotal_mismatches), 1)

    def test_numbers_as_text_with_separators(self):
        self.assertAlmostEqual(etl.parse_number("1 234 567,89"), 1234567.89)
        self.assertAlmostEqual(etl.parse_number("30,000,000"), 30000000.0)
        self.assertAlmostEqual(etl.parse_number("1,234,567.89"), 1234567.89)
        self.assertAlmostEqual(etl.parse_number("3,14"), 3.14)
        self.assertIsNone(etl.parse_number("..."))
        self.assertIsNone(etl.parse_number("-"))
        self.assertIsNone(etl.parse_number(None))

    def test_filename_variants_of_the_export_are_all_recognised(self):
        """Выгрузка называет ФАЙЛ без скобок и с лишними пробелами, а ту же строку
        в ШАПКЕ ЛИСТА пишет со скобками — разбираться должны оба вида."""
        variants = {
            "Позиция за период   18.09.2026  -  18.09.2026   - SECURITIES.xlsx": dt.date(2026, 9, 18),
            "Позиция за период [01.01.2026] - [18.09.2026] - SECURITIES.xlsx": dt.date(2026, 9, 18),
            "Позиция за период 01.01.2026 - 18.09.2026 - SECURITIES.xlsx": dt.date(2026, 9, 18),
        }
        for name, expected in variants.items():
            with self.subTest(name=name):
                path = write_export(self.tmp / "варианты" / name, T0_ROWS)
                self.assertEqual(etl._business_date_from_name(path), expected)

    def test_unrelated_filenames_are_not_mistaken_for_exports(self):
        for name in ("Economic_OVP_Report_v2.xlsx", "ЧПД 2026 09 18.xlsx", "отчёт.xlsx"):
            with self.subTest(name=name):
                self.assertIsNone(etl._business_date_from_name(Path(name)))

    def test_business_date_is_second_date_of_the_period(self):
        self.assertEqual(etl.parse_slice(self.t0_path, "T0").business_date, dt.date(2026, 9, 1))
        self.assertEqual(etl.parse_slice(self.t7_path, "T-7").business_date, dt.date(2026, 8, 25))


# ════════════════════════════════════════════════════════════════════════════
# Инкрементальность: заметки, история, новые портфели
# ════════════════════════════════════════════════════════════════════════════
class OlderReleaseTests(PortfolioDynamicsTestCase):
    """Выпуск, сделанный ДО появления новой колонки, обязан читаться.

    Отчёт инкрементальный: предыдущий файл — вход для следующего. Значит,
    любая новая колонка обязана быть необязательной при чтении, иначе первый
    же запуск после обновления падает на вчерашнем файле, и починить это можно
    только руками.
    """

    def drop_column(self, path, sheet_name, column):
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb[sheet_name]
        for index in range(1, ws.max_column + 1):
            if ws.cell(row=1, column=index).value == column:
                ws.delete_cols(index)
                break
        else:
            self.fail("колонки %r нет в листе %r" % (column, sheet_name))
        wb.save(path)
        return path

    def test_release_without_the_remaining_column_still_loads(self):
        previous = self.drop_column(self.bootstrap_release(), "fact_limit", "limit_remaining")

        loaded = etl.load_previous_release(previous)

        self.assertIn("limit_remaining", loaded.fact_limit.columns)
        self.assertTrue(loaded.fact_limit["limit_remaining"].isna().all())

    def test_the_missing_column_is_mentioned_in_the_log(self):
        previous = self.drop_column(self.bootstrap_release(), "fact_limit", "limit_remaining")

        with self.assertLogs("portfolio_dynamics", level="INFO") as captured:
            etl.load_previous_release(previous)

        self.assertTrue(any("limit_remaining" in line for line in captured.output))

    def test_a_report_builds_on_such_a_release(self):
        """Читается — мало: на нём должен собираться следующий выпуск."""
        previous = self.drop_column(self.bootstrap_release(), "fact_limit", "limit_remaining")

        data = self.build(previous=previous)

        self.assertIn("limit_remaining", data.fact_limit.columns)
        self.assertFalse(data.fact_limit.empty)

    def test_a_genuinely_missing_column_is_still_an_error(self):
        """Послабление касается только новых колонок, а не любых пропаж."""
        previous = self.drop_column(self.bootstrap_release(), "fact_limit", "limit_amount")

        with self.assertRaises(etl.PortfolioDynamicsError):
            etl.load_previous_release(previous)


class IncrementalTests(PortfolioDynamicsTestCase):
    def test_notes_survive_a_rerun(self):
        """Заметки — единственная ручная колонка на машинном листе, merge обязателен."""
        previous = self.bootstrap_release()
        _write_notes(previous, {"AFS_TR_RUR": "Держим объём до КУАП", "HTM_ALCO": "Не наращиваем"})

        data = self.build(previous=previous)
        notes = dict(zip(data.fact_portfolio_snapshot["portfolio_code"],
                         data.fact_portfolio_snapshot["note_text"]))

        self.assertEqual(notes["AFS_TR_RUR"], "Держим объём до КУАП")
        self.assertEqual(notes["HTM_ALCO"], "Не наращиваем")
        self.assertEqual(data.notes_restored, 2)

    def test_second_run_for_the_same_date_replaces_history_rows(self):
        """Идемпотентность: повторный прогон за ту же дату T0 не задваивает строки."""
        previous = self.bootstrap_release()
        first = pd.read_excel(previous, sheet_name="fact_type_daily")

        second_data = self.build(previous=previous)
        second_path = self.save(second_data, name="rerun.xlsx")
        second = pd.read_excel(second_path, sheet_name="fact_type_daily")

        self.assertEqual(len(second), len(first))
        self.assertFalse(second.duplicated(["business_date", "portfolio_type"]).any())
        self.assertEqual(second_data.history_rows_replaced, len(first))

    def test_past_dates_are_never_rewritten(self):
        previous = self.bootstrap_release()
        _append_history(previous, [
            (dt.date(2026, 1, 5), "AFS", 11.0),
            (dt.date(2026, 1, 5), "HTM", 22.0),
        ])
        before = pd.read_excel(previous, sheet_name="fact_type_daily")
        before = before[before["business_date"] == pd.Timestamp(2026, 1, 5)]

        data = self.build(previous=previous)
        after = data.fact_type_daily
        after = after[after["business_date"] == dt.date(2026, 1, 5)]

        self.assertEqual(len(after), 2)
        pd.testing.assert_series_equal(
            before.sort_values("portfolio_type")["volume_amount"].reset_index(drop=True),
            after.sort_values("portfolio_type")["volume_amount"].reset_index(drop=True),
            check_dtype=False,
        )
        self.assertEqual(data.history_rows_carried, 2)

    def test_grains_agree_within_tolerance(self):
        """Сумма volume_t0 по типу = volume_amount этого типа за дату T0 (порог 0.5%)."""
        data = self.build(bootstrap=True)
        by_type = (
            data.fact_portfolio_snapshot
            .merge(data.dim_portfolio[["portfolio_code", "portfolio_type"]], on="portfolio_code")
            .groupby("portfolio_type")["volume_t0"].sum()
        )
        today = data.fact_type_daily[data.fact_type_daily["business_date"] == data.business_date]
        reference = today.set_index("portfolio_type")["volume_amount"]

        for portfolio_type, total in by_type.items():
            self.assertLess(abs(total / reference[portfolio_type] - 1),
                            config.PORTFOLIO_DYNAMICS_TOLERANCE)
        self.assertEqual(set(by_type.index), {"AFS", "HTM"})

    def test_new_portfolio_lands_in_dim_and_is_logged(self):
        previous = self.bootstrap_release()
        t0_with_new = write_export(self.tmp / "t0_new" / export_name("01.09.2026"),
                                   T0_ROWS + [("Позиция: HTM_LINKERS_SEC", None, None, None),
                                              ("Bond", 5 * MLN, 4.0, 4.0)])

        with self.assertLogs("portfolio_dynamics", level="WARNING") as captured:
            data = self.build(previous=previous, t0_path=t0_with_new)

        self.assertIn("HTM_LINKERS_SEC", list(data.dim_portfolio["portfolio_code"]))
        self.assertEqual(data.new_portfolio_codes, ["HTM_LINKERS_SEC"])
        self.assertTrue(any("HTM_LINKERS_SEC" in line for line in captured.output))
        # Тип угадан по префиксу, потому что HTM уже заведён в fact_limit.
        new_row = data.dim_portfolio.set_index("portfolio_code").loc["HTM_LINKERS_SEC"]
        self.assertEqual(new_row["portfolio_type"], "HTM")

    def test_unknown_prefix_leaves_type_empty(self):
        """Код, не сводящийся ни к одному известному типу, остаётся без типа:
        лучше пусто и предупреждение, чем выдуманный тип «ZZZ»."""
        previous = self.bootstrap_release()
        t0_with_new = write_export(self.tmp / "t0_unknown" / export_name("01.09.2026"),
                                   T0_ROWS + [("Позиция: ZZZ_STRANGE", None, None, None),
                                              ("Bond", 5 * MLN, 1.0, 1.0)])

        data = self.build(previous=previous, t0_path=t0_with_new)
        new_row = data.dim_portfolio.set_index("portfolio_code").loc["ZZZ_STRANGE"]
        self.assertIsNone(new_row["portfolio_type"])

    def test_known_type_in_the_code_is_recognised(self):
        previous = self.bootstrap_release()
        t0_with_new = write_export(self.tmp / "t0_known" / export_name("01.09.2026"),
                                   T0_ROWS + [("Позиция: TSS_FX", None, None, None),
                                              ("Bond", 5 * MLN, 1.0, 1.0)])

        data = self.build(previous=previous, t0_path=t0_with_new)
        self.assertEqual(
            data.dim_portfolio.set_index("portfolio_code").loc["TSS_FX", "portfolio_type"],
            "TSS")

    def test_limits_are_carried_over_untouched(self):
        previous = self.bootstrap_release()
        before = pd.read_excel(previous, sheet_name="fact_limit")

        data = self.build(previous=previous)
        pd.testing.assert_frame_equal(
            before.reset_index(drop=True),
            data.fact_limit.reset_index(drop=True)[before.columns],
            check_dtype=False,
        )

    def test_run_without_previous_release_refuses_to_guess(self):
        with self.assertRaisesRegex(etl.PortfolioDynamicsError, "--bootstrap"):
            self.build(previous=None, bootstrap=False)


# ════════════════════════════════════════════════════════════════════════════
# Выходной файл
# ════════════════════════════════════════════════════════════════════════════
class WorkbookTests(PortfolioDynamicsTestCase):
    def test_sheets_are_readable_by_pandas_without_arguments(self):
        path = self.bootstrap_release()
        for sheet in ("dim_portfolio", "fact_limit", "fact_type_daily", "fact_portfolio_snapshot", "dict"):
            frame = pd.read_excel(path, sheet_name=sheet)
            self.assertFalse(frame.empty, sheet)

        snapshot = pd.read_excel(path, sheet_name="fact_portfolio_snapshot")
        self.assertEqual(list(snapshot.columns), etl.SNAPSHOT_COLUMNS)
        self.assertEqual(snapshot["business_date"].nunique(), 1)

    def test_checks_pass_on_a_filled_second_release(self):
        """На втором выпуске с заполненными лимитами FAIL остаётся только по T-7-истории."""
        previous = self.bootstrap_release()
        data = self.build(previous=previous)
        failed = {cid for cid, status, _v in workbook.evaluate_checks(data) if status == "FAIL"}

        # CHK_05 требует строк истории на дату T0 минус сдвиг: их не будет, пока
        # история не накопится за нужное число прогонов.
        self.assertLessEqual(failed, {"CHK_05"})


def _write_notes(path: Path, notes: dict) -> None:
    """Имитирует ручное заполнение note_text казначейством прямо в файле."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["fact_portfolio_snapshot"]
    codes = {ws.cell(row=r, column=2).value: r for r in range(2, ws.max_row + 1)}
    for code, text in notes.items():
        ws.cell(row=codes[code], column=7, value=text)
    wb.save(path)


def _append_history(path: Path, rows) -> None:
    """Дописывает строки истории за прошлые даты — как это сделал бы прошлый прогон."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb["fact_type_daily"]
    for business_date, portfolio_type, volume in rows:
        ws.append([business_date, portfolio_type, volume])
        ws.cell(row=ws.max_row, column=1).number_format = "YYYY-MM-DD"
    wb.save(path)


if __name__ == "__main__":
    unittest.main()
