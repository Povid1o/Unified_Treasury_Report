"""Тесты плоского CSV «Динамики портфелей» (второй выход рядом с xlsx).

Проверяется то, что ломается молча: пустой T-7 не превращается в ноль и
ложный прирост, лимит типа не размножается по портфелям, в ежедневный файл
не попадает история (иначе при дозаписи в BI она задвоится), а xlsx от
появления второго выхода не меняется.
"""
import argparse
import datetime as dt
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR))
sys.path.insert(0, str(BASE_DIR / "tests"))

import config  # noqa: E402
from reports.portfolio_dynamics import etl, flat, workbook  # noqa: E402
from reports.portfolio_dynamics.report import PortfolioDynamicsReport  # noqa: E402
from test_portfolio_dynamics_etl import PortfolioDynamicsTestCase  # noqa: E402
from test_workbook_formulas import build_demo_data  # noqa: E402

DAY = "2026-09-18"


def _value(frame, code, metric):
    """Значение показателя портфеля (code) или типа (code = None). None — строки нет."""
    rows = frame[(frame["axis_2"] == (code or "")) & (frame["axis_3"] == metric)]
    if rows.empty:
        return None
    assert len(rows) == 1, rows
    row = rows.iloc[0]
    return row["text_value"] if row["text_value"] else row["value"]


class DailyFlatTests(unittest.TestCase):
    def setUp(self):
        self.data = build_demo_data()
        self.frame = flat.to_flat(self.data)

    def test_columns_follow_portfolio_report_plus_text_value(self):
        self.assertEqual(list(self.frame.columns), [
            "id", "date_", "axis_0", "axis_1", "axis_2", "axis_3", "axis_4",
            "value", "text_value", "nversionid"])
        self.assertEqual(list(self.frame["id"]), list(range(len(self.frame))))
        self.assertEqual(set(self.frame["axis_0"]), {"Динамика портфелей"})

    def test_only_the_report_date(self):
        """История в ежедневном файле задвоилась бы при дозаписи в BI."""
        self.assertEqual(set(self.frame["date_"]), {DAY})
        volumes = self.frame[self.frame["axis_3"] == flat.M_TYPE_VOLUME]
        self.assertEqual(sorted(volumes["axis_1"]), ["AFS", "HTM"])

    def test_portfolio_values_in_report_unit(self):
        # Демо в млн: AFS_OFZ 30 000 млн = 30 млрд, T-7 = 29 400 млн.
        self.assertAlmostEqual(_value(self.frame, "AFS_OFZ", flat.M_T0), 30.0)
        self.assertAlmostEqual(_value(self.frame, "AFS_OFZ", flat.M_T7), 29.4)
        self.assertAlmostEqual(_value(self.frame, "AFS_OFZ", flat.M_DELTA), 0.6)
        self.assertAlmostEqual(_value(self.frame, "AFS_OFZ", flat.M_DELTA_PCT),
                               (30_000 / 29_400 - 1) * 100)
        self.assertAlmostEqual(_value(self.frame, "AFS_OFZ", flat.M_DUR_GAP), 0.5)
        row = self.frame[(self.frame["axis_2"] == "AFS_OFZ")
                         & (self.frame["axis_3"] == flat.M_T0)].iloc[0]
        self.assertEqual((row["axis_1"], row["axis_4"]), ("AFS", "млрд RUB"))

    def test_type_level_rows(self):
        self.assertAlmostEqual(self._type("AFS", flat.M_TYPE_VOLUME), 55.0)
        self.assertAlmostEqual(self._type("HTM", flat.M_TYPE_LIMIT), 300.0)

    def _type(self, portfolio_type, metric):
        rows = self.frame[(self.frame["axis_1"] == portfolio_type) & (self.frame["axis_2"] == "")
                          & (self.frame["axis_3"] == metric)]
        self.assertEqual(len(rows), 1)
        return rows.iloc[0]["value"]

    def test_limit_is_not_repeated_per_portfolio(self):
        limits = self.frame[self.frame["axis_3"] == flat.M_TYPE_LIMIT]
        self.assertEqual(len(limits), 2)
        self.assertTrue((limits["axis_2"] == "").all())

    def test_texts_go_to_text_value(self):
        self.assertEqual(_value(self.frame, "HTM_ALCO", flat.M_NOTE), "Заметка HTM_ALCO")
        self.assertEqual(_value(self.frame, "HTM_ALCO", flat.M_NAME), "АЛКО, HTM")
        texts = self.frame[self.frame["axis_3"].isin([flat.M_NOTE, flat.M_NAME])]
        self.assertTrue(texts["value"].isna().all())

    def test_missing_values_are_skipped_not_zeroed(self):
        # У HTM_ALCO нет дюрации по КУАП — ни её, ни гэпа, остальное на месте.
        self.assertIsNone(_value(self.frame, "HTM_ALCO", flat.M_DUR_TARGET))
        self.assertIsNone(_value(self.frame, "HTM_ALCO", flat.M_DUR_GAP))
        self.assertIsNotNone(_value(self.frame, "HTM_ALCO", flat.M_T0))

    def test_empty_t7_gives_no_fake_growth(self):
        data = build_demo_data()
        data.fact_portfolio_snapshot.loc[
            data.fact_portfolio_snapshot["portfolio_code"] == "AFS_CORP", "volume_t7"] = None
        frame = flat.to_flat(data)
        self.assertAlmostEqual(_value(frame, "AFS_CORP", flat.M_T0), 25.0)
        for metric in (flat.M_T7, flat.M_DELTA, flat.M_DELTA_PCT):
            self.assertIsNone(_value(frame, "AFS_CORP", metric), metric)

    def test_no_empty_rows(self):
        empty = self.frame[self.frame["value"].isna() & (self.frame["text_value"] == "")]
        self.assertTrue(empty.empty, empty)


class HistoryFlatTests(unittest.TestCase):
    def test_until_is_inclusive_and_only_type_volumes(self):
        frame = flat.history_to_flat(build_demo_data().fact_type_daily, dt.date(2026, 9, 11))
        self.assertEqual(sorted(set(frame["date_"])), ["2026-09-04", "2026-09-11"])
        self.assertEqual(set(frame["axis_3"]), {flat.M_TYPE_VOLUME})
        self.assertTrue((frame["axis_2"] == "").all())

    def test_file_name_carries_the_range(self):
        with tempfile.TemporaryDirectory() as tmp:
            history = build_demo_data().fact_type_daily
            path = flat.save_history(history, dt.date(2026, 9, 17), Path(tmp))
            self.assertEqual(path.name, "dinamika_portfeley_istoriya_2026-09-04_2026-09-11.csv")
            self.assertIsNone(flat.save_history(history, dt.date(2026, 1, 1), Path(tmp)))

    def test_without_until_the_whole_history(self):
        frame = flat.history_to_flat(build_demo_data().fact_type_daily)
        self.assertEqual(sorted(set(frame["date_"])), ["2026-09-04", "2026-09-11", DAY])


class HistorySinceTests(unittest.TestCase):
    """По какую дату история уже лежит в CSV папки — и с какой даты дописывать."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.day = dt.date(2026, 9, 18)

    def tearDown(self):
        self._tmp.cleanup()

    def touch(self, *names):
        for name in names:
            (self.dir / name).write_text("x", encoding="utf-8")

    def test_no_csv_means_the_whole_history(self):
        self.touch("dinamika_portfeley_2026-09-17.xlsx")  # xlsx не в счёт
        self.assertIsNone(flat.covered_until(self.dir, self.day))
        self.assertEqual(flat.history_since(self.dir, self.day), dt.date.min)

    def test_yesterday_means_only_today(self):
        self.touch("dinamika_portfeley_2026-09-10.csv", "dinamika_portfeley_2026-09-17.csv")
        self.assertEqual(flat.history_since(self.dir, self.day), self.day)

    def test_gap_is_filled(self):
        self.touch("dinamika_portfeley_2026-09-11.csv")
        self.assertEqual(flat.history_since(self.dir, self.day), dt.date(2026, 9, 12))

    def test_own_date_and_later_files_do_not_count(self):
        """Повторный запуск за ту же дату собирает то же самое, что и первый."""
        self.touch("dinamika_portfeley_2026-09-18.csv", "dinamika_portfeley_2026-09-25.csv")
        self.assertIsNone(flat.covered_until(self.dir, self.day))

    def test_one_time_history_export_counts(self):
        self.touch("dinamika_portfeley_istoriya_2026-01-09_2026-09-16.csv")
        self.assertEqual(flat.history_since(self.dir, self.day), dt.date(2026, 9, 17))

    def test_other_names_are_ignored(self):
        self.touch("dinamika_portfeley_2026-09-17 (1).csv", "otchet_po_portfelyam_2026-09-17.csv",
                   "dinamika_portfeley_2026-13-45.csv")
        self.assertIsNone(flat.covered_until(self.dir, self.day))

    def test_to_flat_takes_type_volumes_from_the_given_date(self):
        data = build_demo_data()
        volumes = lambda frame: sorted(set(frame[frame["axis_3"] == flat.M_TYPE_VOLUME]["date_"]))
        self.assertEqual(volumes(flat.to_flat(data)), [DAY])
        self.assertEqual(volumes(flat.to_flat(data, dt.date.min)), ["2026-09-04", "2026-09-11", DAY])
        self.assertEqual(volumes(flat.to_flat(data, dt.date(2026, 9, 5))), ["2026-09-11", DAY])
        # Портфели — только на отчётную дату, сколько бы истории ни дописалось.
        frame = flat.to_flat(data, dt.date.min)
        self.assertEqual(set(frame[frame["axis_2"] != ""]["date_"]), {DAY})


class DailyCsvPicksUpHistoryTests(PortfolioDynamicsTestCase):
    """Сквозной прогон: первый CSV получает историю из xlsx, следующий — только свой день."""

    def setUp(self):
        super().setUp()
        config.PORTFOLIO_DYNAMICS_OUTPUT_DIR = self.out_dir  # отсюда берётся предыдущий выпуск
        self.release = self.out_dir / "dinamika_portfeley_2026-09-18.xlsx"
        workbook.save_workbook(build_demo_data(), self.release)  # история 04.09, 11.09, 18.09

    def run_for(self, period_end):
        day = dt.datetime.strptime(period_end, "%d.%m.%Y").date().isoformat()
        from test_portfolio_dynamics_etl import T0_ROWS, export_name, write_export
        t0 = write_export(self.tmp / export_name(period_end), T0_ROWS, period_end=period_end)
        args = argparse.Namespace(
            t0_input=str(t0), t7_input=str(self.t7_path), folder=None, date=None,
            no_import=True, t0_date=None, t7_date=None, previous=None, bootstrap=False,
            history=None, diagnose=False, history_csv=None, from_xlsx=None,
            output=str(self.out_dir / f"dinamika_portfeley_{day}.xlsx"))
        PortfolioDynamicsReport().run(args)
        frame = pd.read_csv(self.out_dir / f"dinamika_portfeley_{day}.csv",
                            encoding="utf-8-sig", keep_default_na=False)
        return sorted(set(frame[frame["axis_3"] == flat.M_TYPE_VOLUME]["date_"]))

    def test_first_csv_gets_history_then_only_its_day(self):
        self.assertEqual(self.run_for("22.09.2026"),
                         ["2026-09-04", "2026-09-11", "2026-09-18", "2026-09-22"])
        self.assertEqual(self.run_for("23.09.2026"), ["2026-09-23"])
        # Повтор за 23.09 — то же самое, а не вся история заново.
        self.assertEqual(self.run_for("23.09.2026"), ["2026-09-23"])

    def test_days_after_the_last_csv_are_filled(self):
        (self.out_dir / "dinamika_portfeley_2026-09-11.csv").write_text("x", encoding="utf-8")
        self.assertEqual(self.run_for("22.09.2026"), ["2026-09-18", "2026-09-22"])


class RunWritesBothFormatsTests(PortfolioDynamicsTestCase):
    def _args(self, **extra):
        values = dict(t0_input=str(self.t0_path), t7_input=str(self.t7_path), folder=None,
                      date=None, no_import=True, t0_date=None, t7_date=None, previous=None,
                      bootstrap=True, history=None, diagnose=False, history_csv=None,
                      from_xlsx=None,
                      output=str(self.out_dir / "dinamika_portfeley_2026-09-01.xlsx"))
        values.update(extra)
        return argparse.Namespace(**values)

    def test_xlsx_and_csv_side_by_side(self):
        PortfolioDynamicsReport().run(self._args())
        xlsx = self.out_dir / "dinamika_portfeley_2026-09-01.xlsx"
        csv = self.out_dir / "dinamika_portfeley_2026-09-01.csv"
        self.assertTrue(xlsx.exists())
        self.assertIn("view_monitor", load_workbook(xlsx, read_only=True).sheetnames)
        frame = pd.read_csv(csv, encoding="utf-8-sig", keep_default_na=False)
        self.assertEqual(set(frame["axis_2"]) - {""}, {"AFS_TR_RUR", "HTM_ALCO"})
        # CSV не мешает найти предыдущий выпуск — им остаётся xlsx.
        self.assertEqual(etl.find_previous_release(self.out_dir), xlsx)
        self.assertEqual(list(self.out_dir.glob("*istoriya*")), [])

    def test_history_csv_flag(self):
        PortfolioDynamicsReport().run(self._args(history_csv="2026-09-01"))
        found = list(self.out_dir.glob("dinamika_portfeley_istoriya_*.csv"))
        self.assertEqual([p.name for p in found],
                         ["dinamika_portfeley_istoriya_2026-09-01_2026-09-01.csv"])

    def test_bad_history_date_fails_before_reading(self):
        with self.assertRaises(etl.PortfolioDynamicsError):
            PortfolioDynamicsReport().run(self._args(history_csv="01.09.2026"))
        self.assertEqual(list(self.out_dir.iterdir()), [])


class HistoryFromXlsxTests(PortfolioDynamicsTestCase):
    """--from-xlsx: история из одного готового выпуска, без срезов и без записи отчёта."""

    def setUp(self):
        super().setUp()
        config.PORTFOLIO_DYNAMICS_OUTPUT_DIR = self.out_dir
        self.release = self.out_dir / "dinamika_portfeley_2026-09-18.xlsx"
        workbook.save_workbook(build_demo_data(), self.release)

    def _run(self, **extra):
        values = dict(diagnose=False, history_csv=None, from_xlsx="", output=None)
        values.update(extra)
        PortfolioDynamicsReport().run(argparse.Namespace(**values))

    def _history_files(self):
        return sorted(p.name for p in self.out_dir.glob("dinamika_portfeley_istoriya_*.csv"))

    def test_whole_history_from_the_freshest_release(self):
        self._run()
        self.assertEqual(self._history_files(),
                         ["dinamika_portfeley_istoriya_2026-09-04_2026-09-18.csv"])
        # Ничего, кроме файла истории: ни нового xlsx, ни ежедневного CSV.
        self.assertEqual(sorted(p.name for p in self.out_dir.iterdir()),
                         [self.release.name] + self._history_files())

    def test_until_date_and_units_survive_the_round_trip(self):
        self._run(history_csv="2026-09-11", from_xlsx=str(self.release))
        path = self.out_dir / "dinamika_portfeley_istoriya_2026-09-04_2026-09-11.csv"
        frame = pd.read_csv(path, encoding="utf-8-sig", keep_default_na=False)
        # Демо в млн: 98% от 55 000 млн AFS = 53.9 млрд — как в xlsx, без сдвига в 1000 раз.
        row = frame[(frame["date_"] == "2026-09-11") & (frame["axis_1"] == "AFS")].iloc[0]
        self.assertAlmostEqual(float(row["value"]), 53.9)
        self.assertEqual(row["axis_4"], "млрд RUB")

    def test_path_without_extension_is_found(self):
        self._run(from_xlsx=str(self.release.with_suffix("")))
        self.assertEqual(len(self._history_files()), 1)

    def test_missing_file_is_an_error(self):
        with self.assertRaises(etl.PortfolioDynamicsError):
            self._run(from_xlsx=str(self.out_dir / "нет такого.xlsx"))


if __name__ == "__main__":
    unittest.main()
