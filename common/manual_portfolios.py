"""Excel-файл «Дополнительных портфелей» и его синхронизация с настройками.

Зачем. Портфели, которых нет в выгрузке позиций, раньше заводились только
экраном консоли — по полю за раз. Здесь тот же список лежит в обычной
Excel-книге (по умолчанию additional_portfolios.xlsx рядом с settings.json):
строка на портфель, правится за минуту. Файл общий для «Динамики портфелей»
и «Отчёта по портфелям»; комментарии к этим портфелям — тоже здесь, свой
для каждого отчёта.

Кто главный. Файл. Отчёты при каждом запуске читают его (sync), и если он
отличается от настроек — переписывают settings.json под него. Экран консоли
правит тот же файл: settings.set_value пишет сначала в Excel, потом в json.
Поэтому консоль и файл всегда показывают одно и то же.

Файла нет (первый запуск, новая установка) — он создаётся из того, что уже
задано в настройках: заведённые через консоль портфели не теряются.

Модуль без rich и без импорта config — как и common/settings.py: его
вызывают и отчёты, и экран настроек, и сам settings.set_value.
"""
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import settings

SHEET = "Портфели"
HELP_SHEET = "Как заполнять"

# (поле записи, заголовок колонки, ширина, формат числа)
COLUMNS = [
    ("code", "Портфель", 22, None),
    ("type", "Тип", 12, None),
    ("volume", "Балансовая стоимость, млн RUB", 18, "#,##0.00"),
    ("qty", "Open QTY, шт", 16, "#,##0"),
    ("duration", "Дюрация, лет", 12, "0.00"),
    ("comment_report", "Комментарий в портфелях", 45, None),
    ("comment_dynamics", "Комментарий в динамике", 45, None),
]
HEADERS = {field: title for field, title, _w, _f in COLUMNS}
TEXT_FIELDS = ("comment_report", "comment_dynamics")

# Колонка узнаётся по началу заголовка: «Портфель», «Код портфеля» и т.п.
# Порядок важен: комментарии и «Open QTY» проверяются раньше общих слов.
# «Название» в файл не пишется, но если его добавить колонкой — прочитается.
_HEADER_PREFIXES = [
    ("comment_report", ("комментарий в порт", "комментарий портф", "комментарий (порт")),
    ("comment_dynamics", ("комментарий в дин", "комментарий динам", "комментарий (дин")),
    ("code", ("портфель", "код")),
    ("name", ("назван",)),
    ("type", ("тип",)),
    ("qty", ("open qty", "qty", "количество", "кол-во", "штук")),
    ("volume", ("балансов", "объём", "объем", "стоимость", "чистая стоимость")),
    ("duration", ("дюрац", "duration")),
]

HELP_LINES = [
    "Портфели, которых нет в выгрузке «Позиция за период». Файл общий для «Динамики "
    "портфелей» и «Отчёта по портфелям». Строка — портфель, пустые строки пропускаются.",
    "",
    "Портфель — код, под которым портфель будет в обоих отчётах. Обязательно.",
    "Тип — AFS, HTM, HTM_KUAP или TSS. Обязательно.",
    "Балансовая стоимость, млн RUB — обязательно: объём в «Динамике», «Чистая стоимость» "
    "в «Отчёте по портфелям».",
    "Open QTY, шт — для «Отчёта по портфелям»; изменение с начала года — ноль.",
    "Дюрация, лет — для «Динамики»; пусто — проверка CHK_14 покажет FAIL.",
    "Комментарии — свой для каждого отчёта. Для этих портфелей пишутся только здесь: "
    "правка в витрине отчёта перезапишется значением из файла.",
    "",
    "P&L с начала года отчёт спрашивает при запуске (Enter — как вчера).",
    "Сохраните файл перед запуском отчёта.",
]

_HEADER_FILL = PatternFill("solid", fgColor="BF8F00")
_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
# Строк с выпадающим списком типов и заливкой ввода — с запасом под новые.
_INPUT_ROWS = 200


@dataclass
class SyncResult:
    records: List[dict]
    path: Optional[Path]
    created: bool = False   # файла не было — создан из настроек
    updated: bool = False   # файл отличался — settings.json переписан под него


def file_path() -> Optional[Path]:
    """Путь к Excel-файлу; None — настройка очищена, файл не ведётся."""
    value = settings.get(settings.MANUAL_PORTFOLIOS_FILE_KEY)
    return Path(value) if str(value).strip() else None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _header_field(title: str) -> Optional[str]:
    name = title.casefold()
    for field, prefixes in _HEADER_PREFIXES:
        if name.startswith(prefixes):
            return field
    return None


def read(path: Path) -> List[dict]:
    """Excel -> проверенный список записей (как settings.get отдаёт его)."""
    path = Path(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # битый файл не должен молча обнулять список
        raise settings.SettingsError(
            f"Не удалось прочитать файл дополнительных портфелей {path}: {exc}") from exc
    try:
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    columns: Dict[str, int] = {}
    header_index = None
    for index, row in enumerate(rows[:10]):
        found = {}
        for col, value in enumerate(row):
            field = _header_field(_text(value))
            if field and field not in found:
                found[field] = col
        if "code" in found:
            columns, header_index = found, index
            break
    if header_index is None:
        raise settings.SettingsError(
            f"{path.name}: на листе «{ws.title}» нет строки заголовков с колонкой "
            f"«{HEADERS['code']}». Ожидаются колонки: {', '.join(HEADERS.values())}.")
    for field in ("type", "volume"):
        if field not in columns:
            raise settings.SettingsError(
                f"{path.name}: нет колонки «{HEADERS[field]}».")

    raw = []
    for number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        record = {field: (row[col] if col < len(row) else None)
                  for field, col in columns.items()}
        if not any(_text(value) for value in record.values()):
            continue
        if not _text(record.get("code")):
            raise settings.SettingsError(
                f"{path.name}: в строке {number} заполнены значения, но не указан код портфеля.")
        raw.append(record)

    try:
        return settings.parse_value(settings.SETTINGS_BY_KEY[settings.MANUAL_PORTFOLIOS_KEY], raw)
    except settings.SettingsError as exc:
        raise settings.SettingsError(f"{path.name}: {exc}") from exc


def write(path: Optional[Path], records: List[dict]) -> Optional[Path]:
    """Пишет список в Excel целиком, через временный файл."""
    if path is None:
        return None
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    for col, (_field, title, width, _fmt) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 32

    last_row = max(len(records) + 1, _INPUT_ROWS)
    for r in range(2, last_row + 1):
        record = records[r - 2] if r - 2 < len(records) else {}
        for col, (field, _title, _width, fmt) in enumerate(COLUMNS, start=1):
            value = record.get(field)
            cell = ws.cell(row=r, column=col, value=value)
            cell.fill = _INPUT_FILL
            if fmt:
                cell.number_format = fmt
            if field in TEXT_FIELDS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    type_col = get_column_letter([f for f, *_ in COLUMNS].index("type") + 1)
    validation = DataValidation(type="list", formula1='"' + ",".join(settings.KNOWN_TYPES_HINT) + '"',
                                allow_blank=True)
    validation.error = "Тип — один из: " + ", ".join(settings.KNOWN_TYPES_HINT)
    validation.showErrorMessage = False  # предупреждение не мешает ввести новый тип
    ws.add_data_validation(validation)
    validation.add(f"{type_col}2:{type_col}{last_row}")
    ws.freeze_panes = "B2"

    help_ws = wb.create_sheet(HELP_SHEET)
    help_ws.column_dimensions["A"].width = 120
    for r, line in enumerate(HELP_LINES, start=1):
        cell = help_ws.cell(row=r, column=1, value=line or None)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    help_ws["A1"].font = Font(bold=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(str(tmp), str(path))
    except PermissionError as exc:
        try:
            tmp.unlink()
        except OSError:
            pass
        raise settings.SettingsError(
            f"Нет доступа для записи в {path} — файл открыт в Excel? Закройте его и "
            f"повторите: {exc}") from exc
    return path


def sync() -> SyncResult:
    """Сверяет Excel-файл с настройками. Файл главный; его нет — создаётся."""
    path = file_path()
    stored = settings.get(settings.MANUAL_PORTFOLIOS_KEY)
    if path is None:
        return SyncResult(records=stored, path=None)
    if not path.is_file():
        write(path, stored)
        return SyncResult(records=stored, path=path, created=True)
    records = read(path)
    if records != stored:
        settings.set_value(settings.MANUAL_PORTFOLIOS_KEY, records, sync_file=False)
        return SyncResult(records=records, path=path, updated=True)
    return SyncResult(records=records, path=path)
